"""
Comprehensive automated verification suite for Unified MCA & GST Hub.
Tests all functional requirements:
- Phone normalization & rejection guards
- MCA CSV & XLSX ingestion with repeated entityId director retention
- GST CSV & XLSX ingestion with trade vs legal name separation
- Idempotent re-imports
- Search attribution & exact lookups
- Formula-injection-safe exports
- Audit logging & quality issue creation
"""

import os
import io
import csv
import datetime
import tempfile
import openpyxl
import xlwt
from decimal import Decimal
from django.test import TestCase, Client
from django.contrib.auth.models import User
from django.urls import reverse
from unittest.mock import patch, PropertyMock

from django.core.files.uploadedfile import SimpleUploadedFile

from hub.models import (
    MCACompany, MCADirectorEntry, GSTRecord,
    ImportBatch, StagedRow, QualityIssue,
    AuditEvent, ExportJob, SystemSetting,
    BusinessLink
)
from hub.cleaners import (
    clean_mobile, clean_email, clean_date, clean_timestamp, clean_capital,
    clean_pincode, clean_pan, clean_gstin, clean_cin, clean_din,
    clean_name, clean_address, sanitize_spreadsheet_value
)
from hub.importers import (
    stage_and_validate_batch, execute_import_publication,
    detect_dataset_type, generate_suggested_mapping,
    validate_mapping_requirements, read_excel_rows,
    get_excel_sheets, auto_detect_header_row, read_file_headers
)
from hub.services import perform_unified_search
from hub.exporters import export_mca_dataset, export_gst_dataset, export_cleaned_master_data
from hub.forms import DatasetUploadForm


class MobileCleanerTestCase(TestCase):
    """Verifies strict phone cleaning rules according to R&D and user specification."""

    def test_clean_mobile_valid_variations(self):
        # Leading backticks
        num, status, err = clean_mobile("``9774465515")
        self.assertIn(status, ('valid', 'cleaned'))
        self.assertEqual(num, "9774465515")

        # Leading apostrophes
        num, status, err = clean_mobile("'9876543210")
        self.assertIn(status, ('valid', 'cleaned'))
        self.assertEqual(num, "9876543210")

        # Recognized +91 prefix
        num, status, err = clean_mobile("+91 9876543210")
        self.assertIn(status, ('valid', 'cleaned'))
        self.assertEqual(num, "9876543210")

        # Recognized 0091 prefix
        num, status, err = clean_mobile("0091-9876543210")
        self.assertIn(status, ('valid', 'cleaned'))
        self.assertEqual(num, "9876543210")

        # Bare 91 prefix when exactly 12 digits
        num, status, err = clean_mobile("919876543210")
        self.assertIn(status, ('valid', 'cleaned'))
        self.assertEqual(num, "9876543210")

        # Domestic leading 0 when exactly 11 digits
        num, status, err = clean_mobile("09876543210")
        self.assertIn(status, ('valid', 'cleaned'))
        self.assertEqual(num, "9876543210")

        # Parentheses and spaces
        num, status, err = clean_mobile("(98765) 43210")
        self.assertIn(status, ('valid', 'cleaned'))
        self.assertEqual(num, "9876543210")

    def test_clean_mobile_critical_guards(self):
        # Do not strip 91 if the number is already valid 10 digits
        num, status, err = clean_mobile("9123456789")
        self.assertEqual(status, 'valid')
        self.assertEqual(num, "9123456789")

        # Reject short numbers
        num, status, err = clean_mobile("9876543")
        self.assertEqual(status, 'invalid')
        self.assertIsNone(num)

        # Reject multiple numbers
        num, status, err = clean_mobile("9876543210 / 9123456789")
        self.assertEqual(status, 'multiple_numbers_review')
        self.assertIsNone(num)

        # Reject scientific notation
        num, status, err = clean_mobile("9.876543E+09")
        self.assertEqual(status, 'invalid')
        self.assertIsNone(num)

        # Strict domestic prefix check (must start with 6-9)
        num, status, err = clean_mobile("5123456789", strict_prefix=True)
        self.assertEqual(status, 'invalid')
        self.assertIsNone(num)

        # Blank numbers return missing
        num, status, err = clean_mobile("")
        self.assertEqual(status, 'missing')
        self.assertIsNone(num)

    def test_user_required_phone_examples(self):
        """
        Tests the exact test suite examples given in user requirements:
        '9876543210          → 9876543210
        ``9876543210         → 9876543210
        +91 98765 43210      → 9876543210
        0091-9876543210      → 9876543210
        09876543210          → 9876543210
        9876543210.0         → 9876543210
        9876543210 / 9123456789 → multiple_numbers_review
        1234567890           → invalid
        blank/null           → missing
        """
        # 1. '9876543210 → 9876543210
        num, status, _ = clean_mobile("'9876543210")
        self.assertEqual(num, "9876543210")
        self.assertEqual(status, "cleaned")

        # 2. ``9876543210 → 9876543210
        num, status, _ = clean_mobile("``9876543210")
        self.assertEqual(num, "9876543210")
        self.assertEqual(status, "cleaned")

        # 3. +91 98765 43210 → 9876543210
        num, status, _ = clean_mobile("+91 98765 43210")
        self.assertEqual(num, "9876543210")
        self.assertEqual(status, "cleaned")

        # 4. 0091-9876543210 → 9876543210
        num, status, _ = clean_mobile("0091-9876543210")
        self.assertEqual(num, "9876543210")
        self.assertEqual(status, "cleaned")

        # 5. 09876543210 → 9876543210
        num, status, _ = clean_mobile("09876543210")
        self.assertEqual(num, "9876543210")
        self.assertEqual(status, "cleaned")

        # 6. 9876543210.0 → 9876543210
        num, status, _ = clean_mobile("9876543210.0")
        self.assertEqual(num, "9876543210")
        self.assertEqual(status, "cleaned")

        # 7. 9876543210 / 9123456789 → multiple_numbers_review
        num, status, _ = clean_mobile("9876543210 / 9123456789")
        self.assertIsNone(num)
        self.assertEqual(status, "multiple_numbers_review")

        # 8. 1234567890 → invalid (starts with 1, not 6-9)
        num, status, _ = clean_mobile("1234567890")
        self.assertIsNone(num)
        self.assertEqual(status, "invalid")

        # 9. blank/null → missing
        num_null, status_null, _ = clean_mobile(None)
        self.assertIsNone(num_null)
        self.assertEqual(status_null, "missing")

        num_blank, status_blank, _ = clean_mobile("")
        self.assertIsNone(num_blank)
        self.assertEqual(status_blank, "missing")

        num_spaces, status_spaces, _ = clean_mobile("   ")
        self.assertIsNone(num_spaces)
        self.assertEqual(status_spaces, "missing")


class DateCleanerTestCase(TestCase):
    """Verifies explicit date parsing for MCA and GST formats."""

    def test_mca_date_parsing(self):
        d, ok, err = clean_date("09/24/2026", preferred_format="%m/%d/%Y")
        self.assertTrue(ok)
        self.assertEqual(d, datetime.date(2026, 9, 24))

    def test_gst_apostrophe_prefixed_date(self):
        d, ok, err = clean_date("'2026-09-25", preferred_format="%Y-%m-%d")
        self.assertTrue(ok)
        self.assertEqual(d, datetime.date(2026, 9, 25))


class ImporterEndToEndTestCase(TestCase):
    """Verifies end-to-end ingestion on sample datasets."""

    def setUp(self):
        self.admin = User.objects.create_superuser('testadmin', 'test@example.com', 'pass123')
        self.mca_file_csv = 'sample_data/sample_mca.csv'
        self.gst_file_xlsx = 'sample_data/sample_gst.xlsx'

    def test_mca_ingestion_and_director_attachment(self):
        """Tests that repeated entity IDs attach all related directors and do not discard records."""
        batch = ImportBatch.objects.create(
            batch_id="TEST-MCA-001",
            dataset_type="MCA",
            file_name="sample_mca.csv",
            file_path=self.mca_file_csv,
            file_checksum="dummy123",
            file_format="CSV",
            header_row_index=0,
            status="UPLOADED",
            created_by=self.admin,
            mapping_config={
                'columns': {
                    'entityId': 'entityId',
                    'entityType': 'entityType',
                    'name': 'name',
                    'state': 'state',
                    'district': 'district',
                    'roc': 'roc',
                    'nicCode': 'nicCode',
                    'nicLabel': 'nicLabel',
                    'classOfCompany': 'classOfCompany',
                    'dateOfIncorporation': 'dateOfIncorporation',
                    'paidUpCapital': 'paidUpCapital',
                    'authorizedCapital': 'authorizedCapital',
                    'email': 'email',
                    'registeredAddress': 'registeredAddress',
                    'directorName': 'directorName',
                    'directorEmail': 'directorEmail',
                    'directorMobile': 'directorMobile',
                    'upload_date': 'upload_date'
                },
                'date_format': '%m/%d/%Y',
                'delimiter': ',',
                'encoding': 'utf-8'
            }
        )

        # Stage and validate
        stage_summary = stage_and_validate_batch(batch, strict_phone=True)
        self.assertEqual(batch.total_input_rows, 2688)
        self.assertEqual(batch.status, 'PREVIEWED')

        # Publish
        res = execute_import_publication(batch, user=self.admin)
        self.assertEqual(res['status'], 'COMPLETED')

        # In sample MCA data: exactly 1,160 unique companies/LLPs
        company_count = MCACompany.objects.count()
        self.assertEqual(company_count, 1160)

        # In sample MCA data: exactly 2,688 director rows
        director_count = MCADirectorEntry.objects.count()
        self.assertEqual(director_count, 2688)

        # Verify company with multiple directors (e.g. U56102TR2026PTC014875)
        vaheen = MCACompany.objects.get(entity_id="U56102TR2026PTC014875")
        self.assertEqual(vaheen.directors.count(), 2)
        dir_names = [d.director_name for d in vaheen.directors.all()]
        self.assertIn("SHYAMA GHOSH", dir_names)
        self.assertIn("SOURAV CHAKRABORTY", dir_names)

        # Verify quality issues created for flagged phones (exactly 8 flagged from benchmark!)
        issues = QualityIssue.objects.filter(batch=batch, error_code__in=['FLAGGED_MOBILE', 'INVALID_PHONE', 'MULTIPLE_NUMBERS'])
        self.assertEqual(issues.count(), 8)

        # Verify AuditEvent recorded
        audit = AuditEvent.objects.filter(object_id=batch.batch_id, action='IMPORT').first()
        self.assertIsNotNone(audit)

    def test_gst_ingestion_and_name_separation(self):
        """Tests that GST records preserve legal_name and trade_name separately and parse properly."""
        batch = ImportBatch.objects.create(
            batch_id="TEST-GST-001",
            dataset_type="GST",
            file_name="sample_gst.xlsx",
            file_path=self.gst_file_xlsx,
            file_checksum="dummygst123",
            file_format="XLSX",
            sheet_name="Processed Data",
            header_row_index=0,
            status="UPLOADED",
            created_by=self.admin,
            mapping_config={
                'columns': {
                    'gst_number': 'gst_number',
                    'legal_name': 'legal_name',
                    'trade_name': 'trade_name',
                    'nature_of_business': 'nature_of_business',
                    'type': 'type',
                    'registration_date': 'registration_date',
                    'gst_type': 'gst_type',
                    'status': 'status',
                    'jurisdiction': 'jurisdiction',
                    'state_jurisdiction': 'state_jurisdiction',
                    'pan': 'pan',
                    'mobile': 'mobile',
                    'email': 'email',
                    'address': 'address',
                    'Pincode': 'Pincode',
                    'state': 'state',
                    'city': 'city',
                    'insert_time': 'insert_time'
                },
                'date_format': '%Y-%m-%d'
            }
        )

        stage_summary = stage_and_validate_batch(batch, strict_phone=True)
        self.assertEqual(batch.total_input_rows, 3034)

        execute_import_publication(batch, user=self.admin)
        self.assertEqual(GSTRecord.objects.count(), 3034)

        # Verify individual registration (03GJVPS4703Q1ZA)
        rec = GSTRecord.objects.get(gst_number="03GJVPS4703Q1ZA")
        self.assertEqual(rec.legal_name, "JASKARAN SINGH")
        self.assertEqual(rec.trade_name, "ORBIT CAD ENTERPRISES")
        self.assertNotEqual(rec.legal_name, rec.trade_name)
        self.assertEqual(rec.registration_date, datetime.date(2026, 9, 25))
        self.assertEqual(rec.status, "ACT")
        self.assertEqual(rec.canonical_mobile, "9807400025")
        self.assertIn(rec.contact_validity, ("valid", "cleaned", "VALID"))

        # Verify exactly 1 flagged mobile in GST sample as per R&D benchmark
        flagged_issues = QualityIssue.objects.filter(batch=batch, error_code__in=['FLAGGED_MOBILE', 'INVALID_PHONE', 'MULTIPLE_NUMBERS'])
        self.assertEqual(flagged_issues.count(), 1)


class SearchAndExportTestCase(TestCase):
    """Verifies unified search attribution and formula-safe export."""

    def setUp(self):
        self.admin = User.objects.create_superuser('expadmin', 'exp@example.com', 'pass123')
        # Create sample MCA Company & Director
        self.company = MCACompany.objects.create(
            entity_id="U12345MH2026PTC999999",
            entity_type="company",
            name="APEX INNOVATIONS PRIVATE LIMITED",
            state="Maharashtra",
            city="Mumbai"
        )
        self.director = MCADirectorEntry.objects.create(
            company=self.company,
            director_name="VIKRAM MALHOTRA",
            canonical_mobile="9820011223",
            contact_validity="VALID"
        )
        # Create sample GST Record with SAME mobile number
        self.gst = GSTRecord.objects.create(
            gst_number="27AAAPL1234M1Z5",
            legal_name="VIKRAM MALHOTRA",
            trade_name="APEX TRADING CO",
            canonical_mobile="9820011223",
            contact_validity="VALID",
            state="Maharashtra"
        )

    def test_search_attribution_by_phone(self):
        """Searching by phone returns both entities without automatically merging them."""
        results = perform_unified_search("9820011223")
        self.assertEqual(len(results['mca_directors']), 1)
        self.assertEqual(len(results['gst_records']), 1)
        self.assertEqual(results['mca_directors'][0].director_name, "VIKRAM MALHOTRA")
        self.assertEqual(results['gst_records'][0].trade_name, "APEX TRADING CO")

    def test_formula_injection_prevention(self):
        """Ensures formula characters are escaped in exported files."""
        self.assertEqual(sanitize_spreadsheet_value("=1+1"), "'=1+1")
        self.assertEqual(sanitize_spreadsheet_value("@SUM(A1:A10)"), "'@SUM(A1:A10)")
        self.assertEqual(sanitize_spreadsheet_value("+919876543210"), "'+919876543210")
        self.assertEqual(sanitize_spreadsheet_value("Standard Text"), "Standard Text")


class AdminChangePasswordTestCase(TestCase):
    """
    Verifies all requirements for Admin Change Password feature:
    - Only authenticated admin users can access
    - Unauthorized access attempts blocked
    - Current password verification
    - Confirmation password match
    - Strong password policy enforcement (8+ chars, upper, lower, digit, special)
    - Successful change invalidates session & updates hash securely
    - Audit event generation
    """

    def setUp(self):
        self.client = Client()
        self.admin_user = User.objects.create_superuser('admin', 'admin@knowletive.com', 'AdminPass123!')
        self.normal_user = User.objects.create_user('staffuser', 'staff@knowletive.com', 'StaffPass123!', is_staff=True, is_superuser=False)

    def test_unauthenticated_access_redirects_to_login(self):
        """Unauthenticated requests must be redirected to login."""
        response = self.client.get(reverse('hub:settings'))
        self.assertEqual(response.status_code, 302)
        self.assertIn('/login/', response.url)

        response = self.client.post(reverse('hub:admin_change_password'), {})
        self.assertEqual(response.status_code, 302)
        self.assertIn('/login/', response.url)

    def test_non_admin_cannot_change_password(self):
        """Non-admin / normal users cannot access admin change password endpoint."""
        self.client.login(username='staffuser', password='StaffPass123!')
        
        # In settings page, change password card should not be shown
        response = self.client.get(reverse('hub:settings'))
        self.assertEqual(response.status_code, 200)
        self.assertNotContains(response, 'id="changePasswordSection"')

        # Attempting POST to change-password view should be denied
        post_data = {
            'current_password': 'StaffPass123!',
            'new_password': 'NewPassword123!',
            'confirm_new_password': 'NewPassword123!'
        }
        post_response = self.client.post(reverse('hub:admin_change_password'), post_data)
        self.assertEqual(post_response.status_code, 302)
        self.assertRedirects(post_response, reverse('hub:dashboard'))
        
        # Verify password did not change
        self.normal_user.refresh_from_db()
        self.assertTrue(self.normal_user.check_password('StaffPass123!'))

    def test_incorrect_current_password(self):
        """Fails when current password is wrong."""
        self.client.login(username='admin', password='AdminPass123!')
        post_data = {
            'current_password': 'WrongPassword!',
            'new_password': 'NewPassword123@',
            'confirm_new_password': 'NewPassword123@'
        }
        response = self.client.post(reverse('hub:admin_change_password'), post_data)
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "The current password entered is incorrect.")

    def test_mismatched_passwords(self):
        """Fails when new password and confirmation password do not match."""
        self.client.login(username='admin', password='AdminPass123!')
        post_data = {
            'current_password': 'AdminPass123!',
            'new_password': 'NewPassword123@',
            'confirm_new_password': 'DifferentPassword123@'
        }
        response = self.client.post(reverse('hub:admin_change_password'), post_data)
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "The new password and confirmation password do not match.")

    def test_weak_password_enforcement(self):
        """Fails when new password violates any strong password rules."""
        self.client.login(username='admin', password='AdminPass123!')

        # Too short (< 8)
        resp = self.client.post(reverse('hub:admin_change_password'), {
            'current_password': 'AdminPass123!',
            'new_password': 'Aa1!',
            'confirm_new_password': 'Aa1!'
        })
        self.assertContains(resp, "at least 8 characters")

        # Missing uppercase
        resp = self.client.post(reverse('hub:admin_change_password'), {
            'current_password': 'AdminPass123!',
            'new_password': 'lowercase123!',
            'confirm_new_password': 'lowercase123!'
        })
        self.assertContains(resp, "at least one uppercase letter (A-Z)")

        # Missing lowercase
        resp = self.client.post(reverse('hub:admin_change_password'), {
            'current_password': 'AdminPass123!',
            'new_password': 'UPPERCASE123!',
            'confirm_new_password': 'UPPERCASE123!'
        })
        self.assertContains(resp, "at least one lowercase letter (a-z)")

        # Missing number
        resp = self.client.post(reverse('hub:admin_change_password'), {
            'current_password': 'AdminPass123!',
            'new_password': 'NoNumbersHere!',
            'confirm_new_password': 'NoNumbersHere!'
        })
        self.assertContains(resp, "at least one number (0-9)")

        # Missing special character
        resp = self.client.post(reverse('hub:admin_change_password'), {
            'current_password': 'AdminPass123!',
            'new_password': 'NoSpecialChar123',
            'confirm_new_password': 'NoSpecialChar123'
        })
        self.assertContains(resp, "at least one special character")

    def test_successful_password_change_and_relogin(self):
        """Successfully changes password, invalidates session, logs audit event, and allows login with new password."""
        self.client.login(username='admin', password='AdminPass123!')
        
        post_data = {
            'current_password': 'AdminPass123!',
            'new_password': 'BrandNewSecurePassword123#',
            'confirm_new_password': 'BrandNewSecurePassword123#'
        }
        response = self.client.post(reverse('hub:admin_change_password'), post_data)
        
        # Must redirect to login page after invalidating session
        self.assertRedirects(response, reverse('hub:login'))

        # Check DB user password updated securely
        self.admin_user.refresh_from_db()
        self.assertTrue(self.admin_user.check_password('BrandNewSecurePassword123#'))
        self.assertFalse(self.admin_user.check_password('AdminPass123!'))

        # Check audit event logged
        audit = AuditEvent.objects.filter(model_name='AdminUser', action='UPDATE').first()
        self.assertIsNotNone(audit)
        self.assertIn("Password changed for admin", audit.object_repr)

        # Old password cannot login
        failed_login = self.client.login(username='admin', password='AdminPass123!')
        self.assertFalse(failed_login)

        # New password succeeds
        success_login = self.client.login(username='admin', password='BrandNewSecurePassword123#')
        self.assertTrue(success_login)


class PhoneManualCorrectionTestCase(TestCase):
    """
    Verifies requirement 15 and 16:
    - Admin can manually correct an invalid phone number.
    - Validate the corrected number again before saving it.
    """

    def setUp(self):
        self.client = Client()
        self.admin = User.objects.create_superuser('phoneadmin', 'phone@test.com', 'AdminPass123!')
        self.company = MCACompany.objects.create(
            entity_id="U99999MH2026PTC111111",
            entity_type="company",
            name="PHONE TEST CORP"
        )
        self.director = MCADirectorEntry.objects.create(
            company=self.company,
            director_name="ANIL KUMAR",
            raw_mobile="1234567890",
            canonical_mobile=None,
            contact_validity="invalid"
        )
        self.issue = QualityIssue.objects.create(
            dataset_type='MCA',
            record_id=self.director.id,
            record_type='DIRECTOR',
            entity_identifier="U99999MH2026PTC111111 / ANIL KUMAR",
            field_name='directorMobile',
            raw_value="1234567890",
            error_code='INVALID_PHONE',
            reason_description="Phone number must start with 6, 7, 8, or 9"
        )

    def test_manual_phone_correction_invalid_fails(self):
        """Entering an invalid corrected phone number is rejected."""
        self.client.login(username='phoneadmin', password='AdminPass123!')
        response = self.client.post(reverse('hub:quality_resolve', args=[self.issue.id]), {
            'note': 'Attempting invalid correction',
            'corrected_phone': '55555'  # Invalid length & prefix
        })
        self.assertEqual(response.status_code, 302)
        
        # Verify issue was NOT resolved
        self.issue.refresh_from_db()
        self.assertFalse(self.issue.is_resolved)

        # Verify director was NOT updated
        self.director.refresh_from_db()
        self.assertEqual(self.director.contact_validity, 'invalid')
        self.assertIsNone(self.director.canonical_mobile)

    def test_manual_phone_correction_valid_succeeds(self):
        """Entering a valid phone number normalizes, updates the record, and resolves the issue."""
        self.client.login(username='phoneadmin', password='AdminPass123!')
        response = self.client.post(reverse('hub:quality_resolve', args=[self.issue.id]), {
            'note': 'Verified directly via corporate registry MCA card',
            'corrected_phone': '+91 98765 43210'
        })
        self.assertEqual(response.status_code, 302)

        # Verify issue is resolved
        self.issue.refresh_from_db()
        self.assertTrue(self.issue.is_resolved)
        self.assertEqual(self.issue.resolved_by, self.admin)

        # Verify director record was updated and validated
        self.director.refresh_from_db()
        self.assertEqual(self.director.canonical_mobile, '9876543210')
        self.assertEqual(self.director.cleaned_phone, '9876543210')
        self.assertIn(self.director.contact_validity, ('valid', 'cleaned'))
        self.assertEqual(self.director.original_phone, '+91 98765 43210')


class RecordDeletionTestCase(TestCase):
    """
    Automated verification for Delete button and atomic deletion endpoints.
    Verifies:
    - Deleting an MCA company with multiple directors (cascades and prevents orphans)
    - Deleting an MCA company without directors
    - Deleting a GST record
    - Canceling the confirmation dialog (markup verification)
    - Unauthorized delete attempts (regular users, anonymous users)
    - Failed/already-deleted requests
    - Refreshing statistics and pagination counts
    - Audit event generation
    """

    def setUp(self):
        self.admin = User.objects.create_superuser(
            username='deladmin', email='deladmin@test.com', password='AdminPass123!'
        )
        self.regular_user = User.objects.create_user(
            username='regularuser', email='user@test.com', password='UserPass123!'
        )

        # MCA Company 1: with directors
        self.mca_with_dirs = MCACompany.objects.create(
            entity_id="U11111MH2026PTC100001",
            entity_type="company",
            name="TECHSOLVE PRIVATE LIMITED",
            state="Maharashtra",
            email="info@techsolve.com"
        )
        self.dir1 = MCADirectorEntry.objects.create(
            company=self.mca_with_dirs,
            director_name="RAHUL SHARMA",
            canonical_mobile="9876543210",
            raw_mobile="+91 9876543210",
            contact_validity="valid"
        )
        self.dir2 = MCADirectorEntry.objects.create(
            company=self.mca_with_dirs,
            director_name="PRIYA VERMA",
            canonical_mobile="9123456780",
            raw_mobile="9123456780",
            contact_validity="valid"
        )

        # MCA Company 2: without directors
        self.mca_no_dirs = MCACompany.objects.create(
            entity_id="U22222DL2026PTC200002",
            entity_type="company",
            name="APEX VENTURES LIMITED",
            state="Delhi",
            email="contact@apex.com"
        )

        # GST Record
        self.gst_rec = GSTRecord.objects.create(
            gst_number="27AAAAA0000A1Z5",
            legal_name="MAHARASHTRA TRADING CO",
            trade_name="MTC SUPPLIES",
            canonical_mobile="9988776655",
            raw_mobile="9988776655",
            contact_validity="valid"
        )

    def test_delete_mca_with_directors_success(self):
        """Admin can delete an MCA company and all its related directors are also deleted."""
        self.client.login(username='deladmin', password='AdminPass123!')
        
        # Verify initial presence
        self.assertTrue(MCACompany.objects.filter(id=self.mca_with_dirs.id).exists())
        self.assertEqual(MCADirectorEntry.objects.filter(company=self.mca_with_dirs).count(), 2)

        url = reverse('hub:mca_delete', args=[self.mca_with_dirs.id])
        response = self.client.post(url, HTTP_X_REQUESTED_WITH='XMLHttpRequest')
        
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertTrue(data['success'])
        self.assertIn('deleted successfully', data['message'])
        self.assertIn('stats', data)
        self.assertIn('total_mca', data['stats'])

        # Verify parent company is gone
        self.assertFalse(MCACompany.objects.filter(id=self.mca_with_dirs.id).exists())
        # Verify directors are cascade-deleted to prevent orphaned records
        self.assertFalse(MCADirectorEntry.objects.filter(company_id=self.mca_with_dirs.id).exists())

        # Verify audit event
        audit = AuditEvent.objects.filter(model_name='MCACompany', action='DELETE').first()
        self.assertIsNotNone(audit)
        self.assertEqual(audit.user, self.admin)
        self.assertIn('directors_deleted', audit.changes)
        self.assertEqual(audit.changes['directors_deleted'], 2)

    def test_delete_mca_without_directors_success(self):
        """Admin can delete an MCA company that has no directors recorded."""
        self.client.login(username='deladmin', password='AdminPass123!')
        
        url = reverse('hub:mca_delete', args=[self.mca_no_dirs.id])
        response = self.client.post(url, HTTP_X_REQUESTED_WITH='XMLHttpRequest')
        
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertTrue(data['success'])

        # Company deleted
        self.assertFalse(MCACompany.objects.filter(id=self.mca_no_dirs.id).exists())

    def test_delete_gst_record_success(self):
        """Admin can delete a GST record safely."""
        self.client.login(username='deladmin', password='AdminPass123!')
        
        url = reverse('hub:gst_delete', args=[self.gst_rec.id])
        response = self.client.post(url, HTTP_X_REQUESTED_WITH='XMLHttpRequest')
        
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertTrue(data['success'])
        self.assertIn('stats', data)
        self.assertIn('total_gst', data['stats'])

        # Verify GST record is deleted
        self.assertFalse(GSTRecord.objects.filter(id=self.gst_rec.id).exists())

        # Verify audit event
        audit = AuditEvent.objects.filter(model_name='GSTRecord', action='DELETE').first()
        self.assertIsNotNone(audit)
        self.assertEqual(audit.user, self.admin)

    def test_unauthorized_delete_attempt_regular_user(self):
        """Non-superuser users cannot delete records (returns 403 Forbidden)."""
        self.client.login(username='regularuser', password='UserPass123!')
        
        url = reverse('hub:mca_delete', args=[self.mca_with_dirs.id])
        response = self.client.post(url, HTTP_X_REQUESTED_WITH='XMLHttpRequest')
        
        self.assertEqual(response.status_code, 403)
        data = response.json()
        self.assertFalse(data['success'])
        self.assertIn('Access denied', data['message'])

        # Record must still exist
        self.assertTrue(MCACompany.objects.filter(id=self.mca_with_dirs.id).exists())

    def test_unauthorized_delete_attempt_anonymous(self):
        """Unauthenticated user is redirected to login."""
        url = reverse('hub:mca_delete', args=[self.mca_with_dirs.id])
        response = self.client.post(url)
        self.assertEqual(response.status_code, 302)
        self.assertIn('/login/', response.url)

    def test_delete_already_deleted_or_nonexistent_record(self):
        """Attempting to delete an already-deleted or non-existent record returns 404 with clear message."""
        self.client.login(username='deladmin', password='AdminPass123!')
        
        url = reverse('hub:mca_delete', args=[999999])
        response = self.client.post(url, HTTP_X_REQUESTED_WITH='XMLHttpRequest')
        
        self.assertEqual(response.status_code, 404)
        data = response.json()
        self.assertFalse(data['success'])
        self.assertTrue(data.get('already_deleted'))
        self.assertIn('does not exist or has already been deleted', data['message'])

    def test_confirmation_dialog_and_cancel_button_rendered_in_templates(self):
        """Verify confirmation modal, Cancel button, and Delete button exist in both MCA and GST pages for admins."""
        self.client.login(username='deladmin', password='AdminPass123!')
        
        # MCA template check
        res_mca = self.client.get(reverse('hub:mca_list'))
        html_mca = res_mca.content.decode('utf-8')
        self.assertIn('deleteConfirmModal', html_mca)
        self.assertIn('btn-delete-record', html_mca)
        self.assertIn('Delete record', html_mca)
        self.assertIn('modalCancelBtn', html_mca)
        self.assertIn('Cancel', html_mca)
        self.assertIn('modalConfirmDeleteBtn', html_mca)
        self.assertIn('Warning:', html_mca)

        # GST template check
        res_gst = self.client.get(reverse('hub:gst_list'))
        html_gst = res_gst.content.decode('utf-8')
        self.assertIn('deleteConfirmModal', html_gst)
        self.assertIn('btn-delete-record', html_gst)
        self.assertIn('Delete record', html_gst)
        self.assertIn('modalCancelBtn', html_gst)
        self.assertIn('Cancel', html_gst)
        self.assertIn('modalConfirmDeleteBtn', html_gst)
        self.assertIn('Warning:', html_gst)

    def test_regular_user_cannot_see_delete_button_in_ui(self):
        """Regular non-superuser users do not see the delete button in the UI."""
        self.client.login(username='regularuser', password='UserPass123!')
        
        res_mca = self.client.get(reverse('hub:mca_list'))
        html_mca = res_mca.content.decode('utf-8')
        # Ensure the actual action button with class is not rendered in the table
        self.assertNotIn('class="btn btn-outline-danger btn-delete-record"', html_mca)

        res_gst = self.client.get(reverse('hub:gst_list'))
        html_gst = res_gst.content.decode('utf-8')
        self.assertNotIn('class="btn btn-outline-danger btn-delete-record"', html_gst)

    def test_statistics_and_pagination_count_refresh(self):
        """Verify updated count returned in delete payload accurately reflects remaining records."""
        self.client.login(username='deladmin', password='AdminPass123!')
        
        initial_mca_count = MCACompany.objects.filter(is_archived=False).count()
        initial_gst_count = GSTRecord.objects.filter(is_archived=False).count()

        # Delete MCA
        res = self.client.post(reverse('hub:mca_delete', args=[self.mca_with_dirs.id]), HTTP_X_REQUESTED_WITH='XMLHttpRequest')
        data = res.json()
        self.assertEqual(data['stats']['total_mca'], initial_mca_count - 1)

        # Delete GST
        res_gst = self.client.post(reverse('hub:gst_delete', args=[self.gst_rec.id]), HTTP_X_REQUESTED_WITH='XMLHttpRequest')
        data_gst = res_gst.json()
        self.assertEqual(data_gst['stats']['total_gst'], initial_gst_count - 1)


class GSTSearchFilterTestCase(TestCase):
    """
    Automated verification for GST Search Filter.
    Tests:
    - Full GSTIN search
    - Partial GSTIN search
    - Legal-name search
    - Trade-name search
    - PAN search
    - Mobile-number search (canonical and raw)
    - City and state search
    - Address and jurisdiction search
    - Case-insensitive search
    - Clear-search behavior (empty query returns all records)
    - Search with pagination
    - Search combined with existing filters (status, phone_status)
    - No-result state ('No matching GST records found')
    - Template UI presence: placeholder, clear button, loading spinner, and search input
    """

    def setUp(self):
        self.user = User.objects.create_user(
            username='gstuser', email='gstuser@test.com', password='GstPass123!'
        )
        self.client.login(username='gstuser', password='GstPass123!')

        # Create diverse GST test records
        self.r1 = GSTRecord.objects.create(
            gst_number="27AAACH1234F1Z1",
            legal_name="CHITRA INFOTECH PRIVATE LIMITED",
            trade_name="CHITRA CLOUD SYSTEMS",
            pan="AAACH1234F",
            canonical_mobile="9822012345",
            raw_mobile="+91 9822012345",
            email="support@chitraconnect.in",
            address="Plot 45, MIDC Industrial Area, Andheri East",
            city="Mumbai",
            state="Maharashtra",
            jurisdiction="Range IV Division VI Mumbai",
            status="ACT",
            contact_validity="valid"
        )
        self.r2 = GSTRecord.objects.create(
            gst_number="07BBBDE5678G2Z2",
            legal_name="DELHI DATA LOGISTICS ENTERPRISE",
            trade_name="DELHI EXPRESS COURIER",
            pan="BBBDE5678G",
            canonical_mobile="9811098765",
            raw_mobile="9811098765",
            email="billing@delhiexpress.com",
            address="Connaught Place, Block B",
            city="New Delhi",
            state="Delhi",
            jurisdiction="Ward 10 Delhi",
            status="CNL",
            contact_validity="cleaned"
        )
        self.r3 = GSTRecord.objects.create(
            gst_number="29CCCKA9012H3Z3",
            legal_name="KARNATAKA AGRO CHEMICALS LLP",
            trade_name="AGRO CHEM SOLUTIONS",
            pan="CCCKA9012H",
            canonical_mobile=None,
            raw_mobile="080-2345678",
            email="sales@karnatakaagro.org",
            address="Peenya 2nd Stage",
            city="Bengaluru",
            state="Karnataka",
            jurisdiction="Division Bangalore West",
            status="ACT",
            contact_validity="invalid"
        )

    def test_full_gstin_search(self):
        """Exact GSTIN query returns the specific record."""
        response = self.client.get(reverse('hub:gst_list'), {'q': '27AAACH1234F1Z1'})
        self.assertEqual(response.status_code, 200)
        results = list(response.context['page_obj'])
        self.assertEqual(len(results), 1)
        self.assertEqual(results[0].gst_number, '27AAACH1234F1Z1')

    def test_partial_gstin_search(self):
        """Partial substring of GSTIN matches correctly."""
        response = self.client.get(reverse('hub:gst_list'), {'q': 'AAACH1234'})
        self.assertEqual(response.status_code, 200)
        results = list(response.context['page_obj'])
        self.assertEqual(len(results), 1)
        self.assertEqual(results[0].gst_number, '27AAACH1234F1Z1')

    def test_legal_name_search(self):
        """Search matches legal name substring."""
        response = self.client.get(reverse('hub:gst_list'), {'q': 'CHITRA INFOTECH'})
        self.assertEqual(response.status_code, 200)
        results = list(response.context['page_obj'])
        self.assertEqual(len(results), 1)
        self.assertEqual(results[0].legal_name, 'CHITRA INFOTECH PRIVATE LIMITED')

    def test_trade_name_search(self):
        """Search matches trade name substring."""
        response = self.client.get(reverse('hub:gst_list'), {'q': 'EXPRESS COURIER'})
        self.assertEqual(response.status_code, 200)
        results = list(response.context['page_obj'])
        self.assertEqual(len(results), 1)
        self.assertEqual(results[0].trade_name, 'DELHI EXPRESS COURIER')

    def test_pan_search(self):
        """Search matches PAN number."""
        response = self.client.get(reverse('hub:gst_list'), {'q': 'CCCKA9012H'})
        self.assertEqual(response.status_code, 200)
        results = list(response.context['page_obj'])
        self.assertEqual(len(results), 1)
        self.assertEqual(results[0].pan, 'CCCKA9012H')

    def test_mobile_number_search(self):
        """Search matches canonical mobile and raw mobile values."""
        # Canonical mobile search
        response = self.client.get(reverse('hub:gst_list'), {'q': '9822012345'})
        results = list(response.context['page_obj'])
        self.assertEqual(len(results), 1)
        self.assertEqual(results[0].canonical_mobile, '9822012345')

        # Raw mobile search (e.g., partial landline or formatted value)
        response_raw = self.client.get(reverse('hub:gst_list'), {'q': '2345678'})
        results_raw = list(response_raw.context['page_obj'])
        self.assertEqual(len(results_raw), 1)
        self.assertEqual(results_raw[0].gst_number, '29CCCKA9012H3Z3')

    def test_city_and_state_search(self):
        """Search matches city or state."""
        response_city = self.client.get(reverse('hub:gst_list'), {'q': 'Bengaluru'})
        self.assertEqual(len(response_city.context['page_obj']), 1)
        self.assertEqual(response_city.context['page_obj'][0].city, 'Bengaluru')

        response_state = self.client.get(reverse('hub:gst_list'), {'q': 'Maharashtra'})
        self.assertEqual(len(response_state.context['page_obj']), 1)
        self.assertEqual(response_state.context['page_obj'][0].state, 'Maharashtra')

    def test_case_insensitive_search(self):
        """Search is case-insensitive (uppercase, lowercase, mixed case)."""
        response_lower = self.client.get(reverse('hub:gst_list'), {'q': 'chitra'})
        self.assertEqual(len(response_lower.context['page_obj']), 1)

        response_upper = self.client.get(reverse('hub:gst_list'), {'q': 'CHITRA'})
        self.assertEqual(len(response_upper.context['page_obj']), 1)

        response_mixed = self.client.get(reverse('hub:gst_list'), {'q': 'cHiTrA'})
        self.assertEqual(len(response_mixed.context['page_obj']), 1)

    def test_clear_search_displays_all_records(self):
        """Clearing or omitting the search query displays all records."""
        response = self.client.get(reverse('hub:gst_list'), {'q': ''})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context['page_obj'].paginator.count, 3)

    def test_search_combined_with_existing_filters(self):
        """Search combines with status or contact validity filters."""
        # Both r1 and r3 have status 'ACT', but only r1 matches 'chitra'
        response = self.client.get(reverse('hub:gst_list'), {
            'q': 'chitra',
            'status': 'ACT'
        })
        self.assertEqual(len(response.context['page_obj']), 1)
        self.assertEqual(response.context['page_obj'][0].gst_number, '27AAACH1234F1Z1')

        # 'ACT' status with search that doesn't match returns empty
        response_empty = self.client.get(reverse('hub:gst_list'), {
            'q': 'delhi',
            'status': 'ACT'
        })
        self.assertEqual(len(response_empty.context['page_obj']), 0)

        # Combine search with phone_status filter
        response_phone = self.client.get(reverse('hub:gst_list'), {
            'q': 'CHITRA',
            'phone_status': 'VALID'
        })
        self.assertEqual(len(response_phone.context['page_obj']), 1)

    def test_no_result_state_message(self):
        """When query matches nothing, 'No matching GST records found' is rendered."""
        response = self.client.get(reverse('hub:gst_list'), {'q': 'NON_EXISTENT_QUERY_XYZ'})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(len(response.context['page_obj']), 0)
        html = response.content.decode('utf-8')
        self.assertIn('No matching GST records found', html)

    def test_template_search_input_and_placeholder_elements(self):
        """Verify input, requested placeholder, loading spinner, and clear button render in HTML."""
        response = self.client.get(reverse('hub:gst_list'))
        html = response.content.decode('utf-8')
        expected_placeholder = 'Search GSTIN, legal name, trade name, PAN, mobile, email, city, or state'
        self.assertIn(expected_placeholder, html)
        self.assertIn('id="gstSearchInput"', html)
        self.assertIn('id="gstClearBtn"', html)
        self.assertIn('id="searchLoadingSpinner"', html)
        self.assertIn('debounced-search', html)


class MCANavbarAndActionsTestCase(TestCase):
    """
    Verifies all requirements for the MCA module sub-navbar:
    - Proper routing for each MCA section (Dashboard, Companies, Directors, Upload, Cleaning, Search, Reports, Settings)
    - Active section highlighting
    - Action toolbar buttons (+ Company, Export JSON, Export Filtered/All Excel, Import JSON, Clear All)
    - Permission enforcement (Admin vs non-admin for destructive actions like Clear All)
    - Search across CIN, company name, director name, phone, email, and location
    - Directors page displays linked companies
    - Companies page displays linked directors
    - No broken or dead links
    """

    def setUp(self):
        self.client = Client()
        self.admin = User.objects.create_superuser('admin_mca', 'admin@example.com', 'AdminPass123!')
        self.staff = User.objects.create_user('staff_mca', 'staff@example.com', 'StaffPass123!')

        self.company = MCACompany.objects.create(
            entity_id="U72200MH2021PTC123456",
            entity_type="company",
            name="TECH MATRIX GLOBAL PRIVATE LIMITED",
            state="Maharashtra",
            city="Mumbai",
            roc="RoC-Mumbai",
            email="info@techmatrix.com",
            registered_address="101 Cyber Heights, Bandra Kurla Complex, Mumbai",
            paid_up_capital=Decimal('500000.00'),
            date_of_incorporation=datetime.date(2021, 5, 20)
        )

        self.director = MCADirectorEntry.objects.create(
            company=self.company,
            director_name="Aarav Sharma",
            canonical_mobile="9876543210",
            raw_mobile="+91 9876543210",
            contact_validity="valid",
            director_email="aarav@techmatrix.com",
            designation="Director"
        )

    def test_mca_dashboard_routing_and_kpis(self):
        """MCA Dashboard renders with active section indicator and authoritative KPIs."""
        self.client.force_login(self.staff)
        res = self.client.get(reverse('hub:mca_dashboard'))
        self.assertEqual(res.status_code, 200)
        self.assertContains(res, "MCA Dashboard")
        self.assertContains(res, "Total MCA Entities")
        self.assertEqual(res.context['mca_active_section'], 'dashboard')
        self.assertEqual(res.context['total_entities'], 1)

    def test_mca_companies_view_and_linked_directors(self):
        """Companies page renders, includes subnavbar with active highlighting, and lists company with linked directors."""
        self.client.force_login(self.staff)
        res = self.client.get(reverse('hub:mca_list'))
        self.assertEqual(res.status_code, 200)
        self.assertContains(res, "TECH MATRIX GLOBAL PRIVATE LIMITED")
        self.assertContains(res, "U72200MH2021PTC123456")
        self.assertContains(res, "Aarav Sharma")
        # Subnavbar buttons present
        self.assertContains(res, "Import Excel / CSV")
        self.assertContains(res, "Import JSON")
        self.assertContains(res, "Export JSON")
        self.assertContains(res, "Export Filtered Excel")
        self.assertContains(res, "Export All Excel")
        self.assertContains(res, "+ Company")
        self.assertContains(res, "Clear All")

    def test_mca_directors_routing_and_linked_company(self):
        """Directors page renders director details and linked company name/CIN."""
        self.client.force_login(self.staff)
        res = self.client.get(reverse('hub:mca_directors'))
        self.assertEqual(res.status_code, 200)
        self.assertContains(res, "Aarav Sharma")
        self.assertContains(res, "9876543210")
        self.assertContains(res, "TECH MATRIX GLOBAL PRIVATE LIMITED")
        self.assertEqual(res.context['mca_active_section'], 'directors')

    def test_mca_upload_view(self):
        """MCA-dedicated upload page renders correctly."""
        self.client.force_login(self.staff)
        res = self.client.get(reverse('hub:mca_upload'))
        self.assertEqual(res.status_code, 200)
        self.assertContains(res, "Upload Authoritative MCA Dataset")
        self.assertEqual(res.context['mca_active_section'], 'upload')

    def test_mca_cleaning_view(self):
        """MCA data cleaning view displays exception inbox filtered for MCA."""
        self.client.force_login(self.staff)
        res = self.client.get(reverse('hub:mca_cleaning'))
        self.assertEqual(res.status_code, 200)
        self.assertContains(res, "MCA Data Cleaning & Exception Inbox")
        self.assertEqual(res.context['mca_active_section'], 'cleaning')

    def test_mca_search_view_comprehensive(self):
        """MCA search searches across company name, CIN, director, phone, address, and city."""
        self.client.force_login(self.staff)
        # Search by company name
        res_comp = self.client.get(reverse('hub:mca_search'), {'q': 'TECH MATRIX'})
        self.assertEqual(len(res_comp.context['page_obj']), 1)

        # Search by CIN
        res_cin = self.client.get(reverse('hub:mca_search'), {'q': 'U72200MH2021PTC123456'})
        self.assertEqual(len(res_cin.context['page_obj']), 1)

        # Search by director name
        res_dir = self.client.get(reverse('hub:mca_search'), {'q': 'Aarav'})
        self.assertEqual(len(res_dir.context['page_obj']), 1)

        # Search by phone number
        res_phone = self.client.get(reverse('hub:mca_search'), {'q': '9876543210'})
        self.assertEqual(len(res_phone.context['page_obj']), 1)

        # Search by location
        res_loc = self.client.get(reverse('hub:mca_search'), {'q': 'Mumbai'})
        self.assertEqual(len(res_loc.context['page_obj']), 1)

        # Non matching
        res_none = self.client.get(reverse('hub:mca_search'), {'q': 'NON_EXISTENT_MCA_XYZ'})
        self.assertEqual(len(res_none.context['page_obj']), 0)
        self.assertContains(res_none, "No matching MCA records found")

    def test_mca_reports_view(self):
        """MCA reports/export page renders."""
        self.client.force_login(self.staff)
        res = self.client.get(reverse('hub:mca_reports'))
        self.assertEqual(res.status_code, 200)
        self.assertContains(res, "MCA Reports & Audited Data Export")
        self.assertEqual(res.context['mca_active_section'], 'reports')

    def test_mca_settings_view(self):
        """MCA settings page displays user profile, password change form for admin, and logout confirmation modal."""
        self.client.force_login(self.admin)
        res = self.client.get(reverse('hub:mca_settings'))
        self.assertEqual(res.status_code, 200)
        self.assertContains(res, "MCA Module & Profile Settings")
        self.assertContains(res, "admin_mca")
        self.assertContains(res, "Change Password")
        self.assertContains(res, "Confirm Sign Out")

    def test_mca_create_company_action(self):
        """+ Company action creates an MCA company with optional initial director."""
        self.client.force_login(self.staff)
        post_data = {
            'name': 'NEW HORIZON TECHNOLOGIES LLP',
            'entity_type': 'llp',
            'entity_id': 'AAA-9999',
            'state': 'Karnataka',
            'city': 'Bengaluru',
            'email': 'hello@newhorizon.in',
            'director_name': 'Priya Nair',
            'director_mobile': '+91 9845012345',
            'director_email': 'priya@newhorizon.in'
        }
        res = self.client.post(reverse('hub:mca_create_company'), post_data, follow=True)
        self.assertEqual(res.status_code, 200)
        self.assertTrue(MCACompany.objects.filter(entity_id='AAA-9999').exists())
        created_comp = MCACompany.objects.get(entity_id='AAA-9999')
        self.assertEqual(created_comp.name, 'NEW HORIZON TECHNOLOGIES LLP')
        self.assertEqual(created_comp.directors.count(), 1)
        self.assertEqual(created_comp.directors.first().canonical_mobile, '9845012345')

    def test_mca_export_json_action(self):
        """Export JSON returns formatted JSON attachment with MCA records and directors."""
        self.client.force_login(self.staff)
        res = self.client.get(reverse('hub:mca_export_json'))
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res['Content-Type'], 'application/json')
        self.assertIn('attachment', res['Content-Disposition'])
        import json
        data = json.loads(res.content.decode('utf-8'))
        self.assertIsInstance(data, list)
        self.assertEqual(len(data), 1)
        self.assertEqual(data[0]['entity_id'], 'U72200MH2021PTC123456')
        self.assertEqual(data[0]['directors'][0]['director_name'], 'Aarav Sharma')

    def test_mca_export_excel_action(self):
        """Export Excel returns valid XLSX attachment."""
        self.client.force_login(self.staff)
        res = self.client.get(reverse('hub:mca_export_excel') + '?scope=all')
        self.assertEqual(res.status_code, 200)
        self.assertIn('application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', res['Content-Type'])
        self.assertIn('attachment', res['Content-Disposition'])

    def test_mca_clear_all_permissions_and_confirmation(self):
        """Clear All is restricted to superuser and requires exact confirmation code 'CLEAR MCA'."""
        # 1. Non-admin cannot clear
        self.client.force_login(self.staff)
        res_denied = self.client.post(reverse('hub:mca_clear_all'), {'confirmation_code': 'CLEAR MCA'}, follow=True)
        self.assertEqual(res_denied.status_code, 200)
        self.assertTrue(MCACompany.objects.filter(id=self.company.id).exists())

        # 2. Admin with wrong confirmation code fails
        self.client.force_login(self.admin)
        res_wrong_code = self.client.post(reverse('hub:mca_clear_all'), {'confirmation_code': 'WRONG CODE'}, follow=True)
        self.assertEqual(res_wrong_code.status_code, 200)
        self.assertTrue(MCACompany.objects.filter(id=self.company.id).exists())

        # 3. Admin with exact confirmation code succeeds
        res_success = self.client.post(reverse('hub:mca_clear_all'), {'confirmation_code': 'CLEAR MCA'}, follow=True)
        self.assertEqual(res_success.status_code, 200)
        self.assertEqual(MCACompany.objects.count(), 0)
        self.assertEqual(MCADirectorEntry.objects.count(), 0)


class GSTToolbarActionsTestCase(TestCase):
    """
    Verifies all requirements for the GST module toolbar action buttons:
    - Exact label, icon, and visual layout matching reference image:
      [+ Import Excel / CSV], [+ New Row], [Restore JSON], [Backup JSON], [Export CSV], [Export Excel]
    - Completely distinct from MCA navbar (no MCA buttons or options)
    - Retains existing Import GST File and Export GST buttons
    - Correct routing and functionality for CSV, Excel, and JSON exports and imports
    - + New Row modal and creation with phone cleaning
    - Preserves GST search, pagination, view, edit, delete
    """

    def setUp(self):
        self.client = Client()
        self.admin = User.objects.create_superuser('admin_gst', 'admin_gst@example.com', 'AdminPass123!')
        self.staff = User.objects.create_user('staff_gst', 'staff_gst@example.com', 'StaffPass123!')

        self.gst = GSTRecord.objects.create(
            gst_number="27ABCDE1234F1Z5",
            legal_name="RELIABLE TEXTILES PRIVATE LIMITED",
            trade_name="RELIABLE WEAVES",
            pan="ABCDE1234F",
            canonical_mobile="9876543210",
            raw_mobile="+91 9876543210",
            contact_validity="valid",
            email="info@reliabletextiles.com",
            status="ACT",
            state="Maharashtra",
            city="Mumbai"
        )

    def test_gst_list_renders_toolbar_without_duplicates(self):
        """GST list page displays exact buttons shown in reference image and no duplicate buttons."""
        self.client.force_login(self.staff)
        res = self.client.get(reverse('hub:gst_list'))
        self.assertEqual(res.status_code, 200)

        # Confirm duplicate set from header row is removed
        self.assertNotContains(res, "Import GST File")

        # Action toolbar buttons from reference image are rendered exactly once
        self.assertContains(res, "+ Import Excel / CSV")
        self.assertContains(res, "+ New Row")
        self.assertContains(res, "Restore JSON")
        self.assertContains(res, "Backup JSON")
        self.assertContains(res, "Export CSV")
        self.assertContains(res, "Export Excel")

        # Confirm NO MCA buttons appear in GST toolbar
        self.assertNotContains(res, "+ Company")
        self.assertNotContains(res, "MCA Dashboard")
        self.assertNotContains(res, "Export Filtered Excel")
        self.assertNotContains(res, "Export All Excel")

    def test_gst_new_row_action(self):
        """+ New Row button creates a new GST record with phone cleaning."""
        self.client.force_login(self.staff)
        post_data = {
            'gst_number': '29XYZAB5678C1Z2',
            'trade_name': 'APEX DIGITAL SERVICES',
            'legal_name': 'APEX DIGITAL PRIVATE LIMITED',
            'pan': 'XYZAB5678C',
            'mobile': '+91 9811223344',
            'email': 'contact@apexdigital.in',
            'status': 'ACT',
            'state': 'Karnataka',
            'city': 'Bengaluru'
        }
        res = self.client.post(reverse('hub:gst_create'), post_data, follow=True)
        self.assertEqual(res.status_code, 200)
        self.assertTrue(GSTRecord.objects.filter(gst_number='29XYZAB5678C1Z2').exists())
        created_gst = GSTRecord.objects.get(gst_number='29XYZAB5678C1Z2')
        self.assertEqual(created_gst.trade_name, 'APEX DIGITAL SERVICES')
        self.assertEqual(created_gst.canonical_mobile, '9811223344')

    def test_gst_backup_json_action(self):
        """Backup JSON returns formatted JSON attachment for GST records."""
        self.client.force_login(self.staff)
        res = self.client.get(reverse('hub:gst_export_json'))
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res['Content-Type'], 'application/json')
        self.assertIn('attachment', res['Content-Disposition'])
        import json
        data = json.loads(res.content.decode('utf-8'))
        self.assertIsInstance(data, list)
        self.assertEqual(len(data), 1)
        self.assertEqual(data[0]['gst_number'], '27ABCDE1234F1Z5')

    def test_gst_export_csv_action(self):
        """Export CSV returns valid CSV attachment for GST records."""
        self.client.force_login(self.staff)
        res = self.client.get(reverse('hub:gst_export_csv') + '?scope=filtered')
        self.assertEqual(res.status_code, 200)
        self.assertIn('text/csv', res['Content-Type'])
        self.assertIn('attachment', res['Content-Disposition'])
        self.assertIn(b'27ABCDE1234F1Z5', res.content)

    def test_gst_export_excel_action(self):
        """Export Excel returns valid XLSX attachment for GST dataset."""
        self.client.force_login(self.staff)
        res = self.client.get(reverse('hub:gst_export_excel') + '?scope=filtered')
        self.assertEqual(res.status_code, 200)
        self.assertIn('application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', res['Content-Type'])
        self.assertIn('attachment', res['Content-Disposition'])


class ColumnMappingWorkflowTestCase(TestCase):
    """Verifies end-to-end column mapping, alias matching, validation, and duplicate prevention."""

    def setUp(self):
        self.user = User.objects.create_user(username='qa_auditor', password='qapassword123', is_staff=True)

    def test_alias_matching_mca_and_gst(self):
        """Aliases like 'Company Name', 'CIN No', 'Director Name', 'DIN No', 'GST No/GSTIN', etc. auto-match."""
        mca_headers = ['Company Name', 'CIN No', 'Director Name', 'DIN No', 'Contact Number', 'Incorporation Date']
        mca_map = generate_suggested_mapping(mca_headers, 'MCA')
        self.assertEqual(mca_map['Company Name'], 'name')
        self.assertEqual(mca_map['CIN No'], 'entityId')
        self.assertEqual(mca_map['Director Name'], 'directorName')
        self.assertEqual(mca_map['DIN No'], 'din')
        self.assertEqual(mca_map['Contact Number'], 'directorMobile')
        self.assertEqual(mca_map['Incorporation Date'], 'dateOfIncorporation')

        gst_headers = ['GST No/GSTIN', 'Legal Name', 'Trade Name', 'Mobile/Phone/Contact', 'PAN No']
        gst_map = generate_suggested_mapping(gst_headers, 'GST')
        self.assertEqual(gst_map['GST No/GSTIN'], 'gst_number')
        self.assertEqual(gst_map['Legal Name'], 'legal_name')
        self.assertEqual(gst_map['Trade Name'], 'trade_name')
        self.assertEqual(gst_map['Mobile/Phone/Contact'], 'mobile')
        self.assertEqual(gst_map['PAN No'], 'pan')

    def test_duplicate_target_mapping_prevention(self):
        """Mapping two source headers to the same target field is rejected."""
        invalid_mapping = {
            'Col_A': 'entityId',
            'Col_B': 'entityId',
            'Col_C': 'name'
        }
        is_valid, errors, warnings = validate_mapping_requirements(invalid_mapping, 'MCA')
        self.assertFalse(is_valid)
        self.assertTrue(any('Multiple source columns cannot map to the same target field' in e for e in errors))

    def test_missing_mandatory_field_rejected(self):
        """Missing mandatory identity fields (entityId for MCA, gst_number for GST) fails validation."""
        mca_without_cin = {'Company_Name': 'name', 'Phone': 'directorMobile'}
        is_valid, errors, warnings = validate_mapping_requirements(mca_without_cin, 'MCA')
        self.assertFalse(is_valid)
        self.assertTrue(any('Mandatory field \'entityId\'' in e for e in errors))

        gst_without_gstin = {'Trade_Name': 'trade_name', 'Mobile': 'mobile'}
        is_valid, errors, warnings = validate_mapping_requirements(gst_without_gstin, 'GST')
        self.assertFalse(is_valid)
        self.assertTrue(any('Mandatory field \'gst_number\'' in e for e in errors))

    def test_warnings_for_unmapped_optional_fields(self):
        """Key optional fields missing from mapping produce informative warnings without blocking."""
        valid_minimal_mca = {'CIN': 'entityId'}
        is_valid, errors, warnings = validate_mapping_requirements(valid_minimal_mca, 'MCA')
        self.assertTrue(is_valid)
        self.assertTrue(len(warnings) > 0)


class QualityInboxClearAllTestCase(TestCase):
    """Verifies Quality Inbox Clear All functionality, admin permission, empty state, and audit logging."""

    def setUp(self):
        self.admin = User.objects.create_superuser(username='qa_super', password='adminpass123', email='admin@dmt.com')
        self.staff = User.objects.create_user(username='qa_staff', password='staffpass123', is_staff=True, is_superuser=False)

        # Create sample quality issues
        QualityIssue.objects.create(
            dataset_type='MCA',
            record_type='DIRECTOR',
            entity_identifier='U12345MH2026PTC123456 / John Doe',
            field_name='directorMobile',
            raw_value='12345',
            error_code='INVALID_PHONE',
            reason_description='Phone number is shorter than 10 digits'
        )
        QualityIssue.objects.create(
            dataset_type='GST',
            record_type='GST',
            entity_identifier='27ABCDE1234F1Z5',
            field_name='mobile',
            raw_value='9876543210 / 9123456789',
            error_code='MULTIPLE_NUMBERS',
            reason_description='Multiple phone numbers detected'
        )

        # Create other core data to verify isolation
        self.company = MCACompany.objects.create(
            entity_type='company',
            entity_id='U99999MH2026PTC999999',
            name='TEST INTEGRITY COMPANY',
            state='Maharashtra'
        )
        self.director = MCADirectorEntry.objects.create(
            company=self.company,
            director_name='Alice Smith',
            raw_mobile='9822334455',
            canonical_mobile='9822334455',
            contact_validity='valid'
        )
        self.gst = GSTRecord.objects.create(
            gst_number='29ABCDE9999F1Z9',
            trade_name='TEST INTEGRITY GST',
            legal_name='TEST INTEGRITY GST LLP'
        )

    def test_unauthorized_user_cannot_clear_quality_inbox(self):
        """Regular staff user cannot clear quality inbox (HTTP 403 / redirect with error)."""
        self.client.force_login(self.staff)
        res = self.client.post(reverse('hub:quality_clear_all'))
        self.assertEqual(res.status_code, 302)
        # Issues must still exist in DB
        self.assertEqual(QualityIssue.objects.count(), 2)

        # AJAX request receives 403
        ajax_res = self.client.post(reverse('hub:quality_clear_all'), HTTP_X_REQUESTED_WITH='XMLHttpRequest')
        self.assertEqual(ajax_res.status_code, 403)
        self.assertEqual(QualityIssue.objects.count(), 2)

    def test_admin_can_clear_all_quality_issues(self):
        """Admin can clear all quality issues and action is logged in AuditEvent."""
        self.client.force_login(self.admin)
        res = self.client.post(reverse('hub:quality_clear_all'))
        self.assertEqual(res.status_code, 302)
        self.assertEqual(QualityIssue.objects.count(), 0)

        # Verify audit log record
        audit = AuditEvent.objects.filter(model_name='QualityIssue', action='DELETE').first()
        self.assertIsNotNone(audit)
        self.assertEqual(audit.user, self.admin)
        self.assertIn('Cleared all 2 items from Quality Inbox', audit.object_repr)
        self.assertEqual(audit.changes.get('cleared_count'), 2)

        # Verify other application data is completely intact
        self.assertTrue(MCACompany.objects.filter(id=self.company.id).exists())
        self.assertTrue(MCADirectorEntry.objects.filter(id=self.director.id).exists())
        self.assertTrue(GSTRecord.objects.filter(id=self.gst.id).exists())

    def test_clear_all_when_empty_state(self):
        """When quality inbox is empty, clear-all safely handles it without error."""
        QualityIssue.objects.all().delete()
        self.client.force_login(self.admin)
        res = self.client.post(reverse('hub:quality_clear_all'))
        self.assertEqual(res.status_code, 302)
        self.assertEqual(QualityIssue.objects.count(), 0)

    def test_ui_renders_clear_all_button_for_admin(self):
        """Admin sees active Clear All button; regular staff does not see Clear All button."""
        # Staff user
        self.client.force_login(self.staff)
        res = self.client.get(reverse('hub:quality_inbox'))
        self.assertEqual(res.status_code, 200)
        self.assertNotContains(res, 'id="openClearQualityModalBtn"')

        # Admin user
        self.client.force_login(self.admin)
        admin_res = self.client.get(reverse('hub:quality_inbox'))
        self.assertEqual(admin_res.status_code, 200)
        self.assertContains(admin_res, 'Clear all quality inbox items?')
        self.assertContains(admin_res, 'Clear All')


class GSTExcelImportAndDataCleaningAuditTestCase(TestCase):
    """
    Complete audit test suite for Data Cleaning Workflow & GST Excel Import.
    Covers:
    - Strict Indian mobile phone cleaning rules & blank/null invalid handling
    - Emails, PAN, GSTIN, CIN, DIN, names, addresses, dates cleaning
    - Valid .xlsx GST import with trade/legal name separation & telephone normalization
    - Valid .xls (BIFF8) GST import via xlrd engine
    - Header detection bypassing title banner and blank rows
    - Multi-sheet Excel support and sheet selection
    - Cross-module dataset rejection (MCA file uploaded to GST workflow)
    - Duplicate GSTIN detection in file and safe DB updates
    - Preview actions: Clean Data (re-run), Reset, Download Cleaned File, Error Report
    - Database transaction rollback ensuring zero partial/corrupted records on failure
    - Table refresh & flash counts (imported, updated, unchanged, duplicate, rejected)
    """

    def setUp(self):
        self.client = Client()
        self.admin = User.objects.create_superuser('gstsuper', 'gstsuper@example.com', 'SuperPass123!')
        self.client.force_login(self.admin)
        self.temp_files = []

    def tearDown(self):
        for path in self.temp_files:
            if os.path.exists(path):
                try:
                    os.remove(path)
                except Exception:
                    pass

    def test_strict_mobile_cleaning_rules(self):
        """
        Clean mobile numbers using required rule:
        - Convert value to text
        - Trim whitespace
        - Remove apostrophes, backticks, spaces, brackets, hyphens, dots, non-numeric chars
        - Remove leading country code +91, 0091, 91 (when 12 digits), or leading 0 (when 11 digits)
        - Keep only valid 10-digit Indian mobile number (starting with 6, 7, 8, 9)
        - Store invalid or ambiguous numbers as blank/null (None) and show in validation report
        """
        # Valid 10-digit unchanged
        num, status, err = clean_mobile("9876543210")
        self.assertEqual(num, "9876543210")
        self.assertEqual(status, "valid")
        self.assertIsNone(err)

        # Leading apostrophe & backticks
        num, status, err = clean_mobile("'9876543210")
        self.assertEqual(num, "9876543210")
        self.assertEqual(status, "cleaned")

        num, status, err = clean_mobile("``9876543210")
        self.assertEqual(num, "9876543210")
        self.assertEqual(status, "cleaned")

        # Country code +91 with spaces
        num, status, err = clean_mobile("+91 98765 43210")
        self.assertEqual(num, "9876543210")
        self.assertEqual(status, "cleaned")

        # 0091 prefix with hyphens
        num, status, err = clean_mobile("0091-9876543210")
        self.assertEqual(num, "9876543210")
        self.assertEqual(status, "cleaned")

        # 12-digit with 91 prefix
        num, status, err = clean_mobile("919876543210")
        self.assertEqual(num, "9876543210")
        self.assertEqual(status, "cleaned")

        # 11-digit with leading 0
        num, status, err = clean_mobile("09876543210")
        self.assertEqual(num, "9876543210")
        self.assertEqual(status, "cleaned")

        # Float format from Excel .0
        num, status, err = clean_mobile("9876543210.0")
        self.assertEqual(num, "9876543210")
        self.assertEqual(status, "cleaned")

        # Brackets and dots
        num, status, err = clean_mobile("(98765).43210")
        self.assertEqual(num, "9876543210")
        self.assertEqual(status, "cleaned")

        # Invalid: Multiple numbers -> stored as None, multiple_numbers_review
        num, status, err = clean_mobile("9876543210 / 9123456789")
        self.assertIsNone(num)
        self.assertEqual(status, "multiple_numbers_review")
        self.assertIsNotNone(err)

        # Invalid prefix (must start with 6-9 in India) -> stored as None
        num, status, err = clean_mobile("1234567890")
        self.assertIsNone(num)
        self.assertEqual(status, "invalid")
        self.assertIn("must start with 6, 7, 8, or 9", err)

        # Invalid length (short) -> stored as None
        num, status, err = clean_mobile("98765")
        self.assertIsNone(num)
        self.assertEqual(status, "invalid")

        # Scientific notation -> stored as None
        num, status, err = clean_mobile("9.876543E+09")
        self.assertIsNone(num)
        self.assertEqual(status, "invalid")

        # Blank / None -> stored as None, missing
        num, status, err = clean_mobile(None)
        self.assertIsNone(num)
        self.assertEqual(status, "missing")

        num, status, err = clean_mobile("   ")
        self.assertIsNone(num)
        self.assertEqual(status, "missing")

    def test_cleaners_suite_fields(self):
        """Validates cleaning of emails, PAN, GSTIN, CIN, DIN, names, addresses, and dates without corruption."""
        # Email
        email, status, err = clean_email("Test.User+Service@Example.COM")
        self.assertEqual(email, "test.user+service@example.com")
        self.assertEqual(status, "VALID")

        bad_email, status, err = clean_email("invalid_email_format")
        self.assertIsNone(bad_email)
        self.assertEqual(status, "FLAGGED")

        # PAN
        pan, ok, err = clean_pan("abcde1234f")
        self.assertEqual(pan, "ABCDE1234F")
        self.assertTrue(ok)

        bad_pan, ok, err = clean_pan("ABC1234567")
        self.assertFalse(ok)

        # GSTIN
        gstin, ok, err = clean_gstin("27abcde1234f1z5")
        self.assertEqual(gstin, "27ABCDE1234F1Z5")
        self.assertTrue(ok)

        # CIN & LLPIN
        cin, ok, err = clean_cin("U12345MH2026PTC123456")
        self.assertEqual(cin, "U12345MH2026PTC123456")
        self.assertTrue(ok)

        llpin, ok, err = clean_cin("AAB-1234")
        self.assertEqual(llpin, "AAB-1234")
        self.assertTrue(ok)

        # DIN
        din, ok, err = clean_din("123456")
        self.assertEqual(din, "00123456") # 8-digit zero-padded
        self.assertTrue(ok)

        # Name
        name, ok, err = clean_name(" 'Acme   Technologies   Pvt   Ltd' ")
        self.assertEqual(name, "Acme Technologies Pvt Ltd")
        self.assertTrue(ok)

        # Address
        addr, ok, err = clean_address(' "Plot  No.  12,   Sector  4" ')
        self.assertEqual(addr, "Plot No. 12, Sector 4")
        self.assertTrue(ok)

        # Date
        dt, ok, err = clean_date("2026-09-25")
        self.assertEqual(dt, datetime.date(2026, 9, 25))
        self.assertTrue(ok)

    def test_valid_gst_xlsx_import(self):
        """Creates and imports a valid GST .xlsx file and verifies DB records, name separation, and phone cleaning."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "GST Data"
        headers = ["GSTIN", "Legal Name", "Trade Name", "PAN", "Mobile Number", "Email", "Address", "City", "State", "PIN", "Status", "Registration Date"]
        ws.append(headers)
        ws.append([
            "27ABCDE1234F1Z5",
            "ALPHA COMMERCIAL SERVICES PRIVATE LIMITED",
            "ALPHA LOGISTICS",
            "ABCDE1234F",
            "+91 98765-43210",
            "alpha@example.com",
            "101 Business Hub, Andheri",
            "Mumbai",
            "Maharashtra",
            "400053",
            "Active",
            "2026-09-25"
        ])
        temp_file = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
        temp_path = temp_file.name
        temp_file.close()
        wb.save(temp_path)
        self.temp_files.append(temp_path)

        batch = ImportBatch.objects.create(
            batch_id="TEST-AUDIT-GST-XLSX",
            dataset_type="GST",
            file_name="audit_gst.xlsx",
            file_path=temp_path,
            file_checksum="chk_audit_xlsx",
            file_format="XLSX",
            sheet_name="GST Data",
            header_row_index=0,
            status="UPLOADED",
            created_by=self.admin,
            mapping_config={
                'columns': {
                    'GSTIN': 'gst_number',
                    'Legal Name': 'legal_name',
                    'Trade Name': 'trade_name',
                    'PAN': 'pan',
                    'Mobile Number': 'mobile',
                    'Email': 'email',
                    'Address': 'address',
                    'City': 'city',
                    'State': 'state',
                    'PIN': 'Pincode',
                    'Status': 'status',
                    'Registration Date': 'registration_date'
                }
            }
        )

        stage_summary = stage_and_validate_batch(batch, strict_phone=True)
        self.assertEqual(batch.total_input_rows, 1)
        self.assertEqual(batch.status, 'PREVIEWED')

        pub_res = execute_import_publication(batch, user=self.admin)
        self.assertEqual(pub_res['status'], 'COMPLETED')
        self.assertEqual(pub_res['imported_count'], 1)

        # Verify DB record
        record = GSTRecord.objects.get(gst_number="27ABCDE1234F1Z5")
        self.assertEqual(record.legal_name, "ALPHA COMMERCIAL SERVICES PRIVATE LIMITED")
        self.assertEqual(record.trade_name, "ALPHA LOGISTICS")
        self.assertEqual(record.canonical_mobile, "9876543210")
        self.assertEqual(record.raw_mobile, "+91 98765-43210")
        self.assertEqual(record.pan, "ABCDE1234F")
        self.assertEqual(record.email, "alpha@example.com")
        self.assertEqual(record.city, "Mumbai")
        self.assertEqual(record.state, "Maharashtra")
        self.assertEqual(record.pincode, "400053")
        self.assertEqual(record.registration_date, datetime.date(2026, 9, 25))

    def test_valid_gst_xls_import(self):
        """Creates and imports a valid legacy .xls (BIFF8) file via xlwt and xlrd engine."""
        wb = xlwt.Workbook()
        ws = wb.add_sheet("GST Records")
        headers = ["GSTIN", "Legal Name", "Trade Name", "Mobile Number", "Email"]
        for col_idx, h in enumerate(headers):
            ws.write(0, col_idx, h)
        row_vals = ["29XYZAB9876C1Z1", "BETA RETAILS LLP", "BETA MART", "09876543210", "beta@retail.com"]
        for col_idx, val in enumerate(row_vals):
            ws.write(1, col_idx, val)

        temp_file = tempfile.NamedTemporaryFile(suffix=".xls", delete=False)
        temp_path = temp_file.name
        temp_file.close()
        wb.save(temp_path)
        self.temp_files.append(temp_path)

        # Verify read_excel_rows parses .xls correctly
        rows = read_excel_rows(temp_path, sheet_name="GST Records")
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[0], headers)
        self.assertEqual(rows[1][0], "29XYZAB9876C1Z1")

        batch = ImportBatch.objects.create(
            batch_id="TEST-AUDIT-GST-XLS",
            dataset_type="GST",
            file_name="audit_gst.xls",
            file_path=temp_path,
            file_checksum="chk_audit_xls",
            file_format="XLS",
            sheet_name="GST Records",
            header_row_index=0,
            status="UPLOADED",
            created_by=self.admin,
            mapping_config={
                'columns': {
                    'GSTIN': 'gst_number',
                    'Legal Name': 'legal_name',
                    'Trade Name': 'trade_name',
                    'Mobile Number': 'mobile',
                    'Email': 'email'
                }
            }
        )

        stage_and_validate_batch(batch, strict_phone=True)
        self.assertEqual(batch.total_input_rows, 1)

        pub_res = execute_import_publication(batch, user=self.admin)
        self.assertEqual(pub_res['status'], 'COMPLETED')

        record = GSTRecord.objects.get(gst_number="29XYZAB9876C1Z1")
        self.assertEqual(record.legal_name, "BETA RETAILS LLP")
        self.assertEqual(record.trade_name, "BETA MART")
        self.assertEqual(record.canonical_mobile, "9876543210") # 11-digit leading 0 stripped

    def test_header_row_auto_detection_with_title_and_blank_rows(self):
        """Auto-detects header row even if title banners or blank rows appear first."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "GST Data"
        ws.append(["DEPARTMENT OF REVENUE - GST PORTAL EXPORT REPORT"])  # Row 0: Title banner
        ws.append(["", "", "", ""])                                      # Row 1: Blank row
        ws.append(["GST Number", "Legal Name", "Trade Name", "Contact Number", "Email ID"]) # Row 2: Actual header
        ws.append(["27ABCDE1234F1Z5", "Acme Corp", "Acme Retail", "9876543210", "acme@test.com"]) # Row 3: Data

        temp_file = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
        temp_path = temp_file.name
        temp_file.close()
        wb.save(temp_path)
        self.temp_files.append(temp_path)

        detected_index = auto_detect_header_row(temp_path, sheet_name="GST Data")
        self.assertEqual(detected_index, 2)

    def test_multiple_excel_sheets_support(self):
        """Supports files with multiple sheets and allows selecting the correct sheet."""
        wb = openpyxl.Workbook()
        ws1 = wb.active
        ws1.title = "Summary Information"
        ws1.append(["Note", "This sheet contains summary only"])

        ws2 = wb.create_sheet(title="Taxpayers")
        ws2.append(["GSTIN", "Legal Name", "Trade Name", "Mobile"])
        ws2.append(["27ABCDE1234F1Z5", "DELTA ENTERPRISES", "DELTA TRADERS", "9876543210"])

        temp_file = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
        temp_path = temp_file.name
        temp_file.close()
        wb.save(temp_path)
        self.temp_files.append(temp_path)

        sheets = get_excel_sheets(temp_path)
        self.assertEqual(sheets, ["Summary Information", "Taxpayers"])

        # Stage with Taxpayers sheet
        batch = ImportBatch.objects.create(
            batch_id="TEST-MULTISHEET-GST",
            dataset_type="GST",
            file_name="multisheet.xlsx",
            file_path=temp_path,
            file_checksum="chk_multisheet",
            file_format="XLSX",
            sheet_name="Taxpayers",
            header_row_index=0,
            status="UPLOADED",
            created_by=self.admin,
            mapping_config={
                'columns': {
                    'GSTIN': 'gst_number',
                    'Legal Name': 'legal_name',
                    'Trade Name': 'trade_name',
                    'Mobile': 'mobile'
                }
            }
        )
        stage_and_validate_batch(batch, strict_phone=True)
        self.assertEqual(batch.total_input_rows, 1)

    def test_mca_file_rejected_in_gst_workflow(self):
        """Rejects MCA files uploaded in the GST workflow with a clear validation message."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["entityId", "entityType", "name", "roc", "paidUpCapital", "authorizedCapital", "registeredAddress"])
        ws.append(["U12345MH2026PTC123456", "company", "ACME PRIVATE LIMITED", "ROC Mumbai", "100000", "500000", "Mumbai"])

        temp_file = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
        temp_path = temp_file.name
        temp_file.close()
        wb.save(temp_path)
        self.temp_files.append(temp_path)

        batch = ImportBatch.objects.create(
            batch_id="TEST-REJECT-MCA-IN-GST",
            dataset_type="GST",
            file_name="mca_sample.xlsx",
            file_path=temp_path,
            file_checksum="chk_mca_in_gst",
            file_format="XLSX",
            header_row_index=0,
            status="UPLOADED",
            created_by=self.admin
        )

        response = self.client.post(reverse('hub:import_map', args=[batch.batch_id]), {
            'target_field_entityId': 'ignore',
            'target_field_name': 'legal_name'
        }, follow=True)

        self.assertEqual(response.status_code, 200)
        content = response.content.decode('utf-8')
        self.assertIn("MCA dataset detected", content)

    def test_duplicate_gstin_detection_and_safe_update(self):
        """Detects duplicate GSTINs in the uploaded file and updates existing records safely."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["GSTIN", "Legal Name", "Trade Name", "Mobile Number"])
        ws.append(["27AAAAA1111A1Z1", "ORIGINAL LEGAL NAME", "ORIGINAL TRADE", "9876543210"])
        ws.append(["27AAAAA1111A1Z1", "ORIGINAL LEGAL NAME", "REVISED TRADE NAME", "9876543210"])

        temp_file = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
        temp_path = temp_file.name
        temp_file.close()
        wb.save(temp_path)
        self.temp_files.append(temp_path)

        batch = ImportBatch.objects.create(
            batch_id="TEST-DUP-GSTIN",
            dataset_type="GST",
            file_name="dup_gst.xlsx",
            file_path=temp_path,
            file_checksum="chk_dup_gst",
            file_format="XLSX",
            header_row_index=0,
            status="UPLOADED",
            created_by=self.admin,
            mapping_config={
                'columns': {
                    'GSTIN': 'gst_number',
                    'Legal Name': 'legal_name',
                    'Trade Name': 'trade_name',
                    'Mobile Number': 'mobile'
                }
            }
        )

        stage_and_validate_batch(batch, strict_phone=True)
        # Duplicate GSTIN counted as conflict
        self.assertEqual(batch.conflict_count, 1)

        staged_rows = StagedRow.objects.filter(batch=batch).order_by('row_number')
        self.assertEqual(staged_rows[0].proposed_action, 'CREATE')
        self.assertEqual(staged_rows[1].proposed_action, 'UPDATE')
        self.assertTrue(any('Duplicate GSTIN' in msg for msg in staged_rows[1].validation_messages))

        execute_import_publication(batch, user=self.admin)
        self.assertEqual(GSTRecord.objects.filter(gst_number="27AAAAA1111A1Z1").count(), 1)
        record = GSTRecord.objects.get(gst_number="27AAAAA1111A1Z1")
        self.assertEqual(record.trade_name, "REVISED TRADE NAME")

    def test_clean_preview_actions_endpoints(self):
        """Verifies Clean Data (re-run), Reset, Download Cleaned File, and Error Report actions."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["GSTIN", "Legal Name", "Mobile"])
        ws.append(["27BBBBB2222B1Z2", "OMEGA SERVICES", "+91 9988776655"])
        ws.append(["27CCCCC3333C1Z3", "SIGMA TECH", "12345"]) # Invalid mobile

        temp_file = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
        temp_path = temp_file.name
        temp_file.close()
        wb.save(temp_path)
        self.temp_files.append(temp_path)

        batch = ImportBatch.objects.create(
            batch_id="TEST-ACTIONS-GST",
            dataset_type="GST",
            file_name="actions_test.xlsx",
            file_path=temp_path,
            file_checksum="chk_actions_gst",
            file_format="XLSX",
            header_row_index=0,
            status="UPLOADED",
            created_by=self.admin,
            mapping_config={
                'columns': {
                    'GSTIN': 'gst_number',
                    'Legal Name': 'legal_name',
                    'Mobile': 'mobile'
                }
            }
        )

        stage_and_validate_batch(batch, strict_phone=True)
        self.assertEqual(batch.status, 'PREVIEWED')

        # 1. Preview page contains all action buttons and displays before/after
        prev_res = self.client.get(reverse('hub:import_preview', args=[batch.batch_id]))
        self.assertEqual(prev_res.status_code, 200)
        prev_content = prev_res.content.decode('utf-8')
        self.assertIn("Clean Data", prev_content)
        self.assertIn("Download Cleaned File", prev_content)
        self.assertIn("Reset", prev_content)
        self.assertIn("Apply & Save to Database", prev_content)
        self.assertIn("+91 9988776655", prev_content) # Raw phone
        self.assertIn("9988776655", prev_content)      # Cleaned phone

        # 2. Re-clean endpoint
        reclean_res = self.client.post(reverse('hub:import_reclean', args=[batch.batch_id]), follow=True)
        self.assertEqual(reclean_res.status_code, 200)
        batch.refresh_from_db()
        self.assertEqual(batch.status, 'PREVIEWED')

        # 3. Download Cleaned File
        cleaned_csv_res = self.client.get(reverse('hub:download_cleaned_file', args=[batch.batch_id]))
        self.assertEqual(cleaned_csv_res.status_code, 200)
        self.assertEqual(cleaned_csv_res['Content-Type'], 'text/csv; charset=utf-8')
        csv_content = cleaned_csv_res.content.decode('utf-8')
        self.assertIn("9988776655", csv_content)
        # Original file on disk must be completely untouched
        self.assertTrue(os.path.exists(temp_path))

        # 4. Download Error / Cleaning Report
        report_res = self.client.get(reverse('hub:download_cleaning_report', args=[batch.batch_id]))
        self.assertEqual(report_res.status_code, 200)
        report_content = report_res.content.decode('utf-8')
        self.assertIn("Row Number", report_content)
        self.assertIn("Raw Mobile", report_content)

        # 5. Reset endpoint
        reset_res = self.client.post(reverse('hub:import_reset', args=[batch.batch_id]), follow=True)
        self.assertEqual(reset_res.status_code, 200)
        batch.refresh_from_db()
        self.assertEqual(batch.status, 'UPLOADED')
        self.assertEqual(StagedRow.objects.filter(batch=batch).count(), 0)

    def test_atomic_transaction_rollback_on_failure(self):
        """Verifies database transaction rollback ensuring no partial or corrupted records persist on failure."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["GSTIN", "Legal Name", "Mobile Number"])
        ws.append(["27ROLLBACK1111A1Z1", "ROLLBACK CORP 1", "9876543210"])
        ws.append(["27ROLLBACK2222B1Z2", "ROLLBACK CORP 2", "9876543210"])

        temp_file = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
        temp_path = temp_file.name
        temp_file.close()
        wb.save(temp_path)
        self.temp_files.append(temp_path)

        batch = ImportBatch.objects.create(
            batch_id="TEST-ROLLBACK-GST",
            dataset_type="GST",
            file_name="rollback_gst.xlsx",
            file_path=temp_path,
            file_checksum="chk_rollback_gst",
            file_format="XLSX",
            header_row_index=0,
            status="UPLOADED",
            created_by=self.admin,
            mapping_config={
                'columns': {
                    'GSTIN': 'gst_number',
                    'Legal Name': 'legal_name',
                    'Mobile Number': 'mobile'
                }
            }
        )

        stage_and_validate_batch(batch, strict_phone=True)

        # Simulate exception during execution
        with patch.object(GSTRecord.objects, 'filter', side_effect=Exception("Database connection severed")):
            with self.assertRaises(Exception):
                execute_import_publication(batch, user=self.admin)

        self.assertEqual(GSTRecord.objects.filter(gst_number__startswith="27ROLLBACK").count(), 0)

    def test_gst_import_redirect_and_detailed_flash_message(self):
        """After successful import, redirects to GST table and shows flash message with detailed counts."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["GSTIN", "Legal Name", "Trade Name"])
        ws.append(["27FLASH1111A1Z1", "FLASH CORP", "FLASH TRADERS"])

        temp_file = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
        temp_path = temp_file.name
        temp_file.close()
        wb.save(temp_path)
        self.temp_files.append(temp_path)

        batch = ImportBatch.objects.create(
            batch_id="TEST-FLASH-GST",
            dataset_type="GST",
            file_name="flash_gst.xlsx",
            file_path=temp_path,
            file_checksum="chk_flash_gst",
            file_format="XLSX",
            header_row_index=0,
            status="UPLOADED",
            created_by=self.admin,
            mapping_config={
                'columns': {
                    'GSTIN': 'gst_number',
                    'Legal Name': 'legal_name',
                    'Trade Name': 'trade_name'
                }
            }
        )

        stage_and_validate_batch(batch, strict_phone=True)

        # POST to import_execute endpoint
        exec_res = self.client.post(reverse('hub:import_execute', args=[batch.batch_id]), follow=True)
        self.assertEqual(exec_res.status_code, 200)
        # Asserts redirected to gst_list
        self.assertEqual(exec_res.redirect_chain[0][0], reverse('hub:gst_list'))
        flash_text = exec_res.content.decode('utf-8')
        self.assertIn("GST import completed successfully", flash_text)
        self.assertIn("1 new imported", flash_text)

    def test_empty_excel_file(self):
        """Empty Excel file is handled gracefully without crashing."""
        wb = openpyxl.Workbook()
        ws = wb.active  # empty
        temp_file = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
        temp_path = temp_file.name
        temp_file.close()
        wb.save(temp_path)
        self.temp_files.append(temp_path)

        headers = read_file_headers(temp_path, file_format="XLSX")
        self.assertEqual(headers, [])

    def test_corrupted_excel_file(self):
        """Corrupted file with .xlsx extension returns empty and fails safely."""
        temp_file = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
        temp_file.write(b"NOT AN EXCEL FILE JUST CORRUPT BYTES 1234567890")
        temp_path = temp_file.name
        temp_file.close()
        self.temp_files.append(temp_path)

        headers = read_file_headers(temp_path, file_format="XLSX")
        self.assertEqual(headers, [])

    def test_wrong_file_type_form_validation(self):
        """Non-spreadsheet files are rejected during upload form validation."""
        fake_file = SimpleUploadedFile("malicious.pdf", b"%PDF-1.4 ...", content_type="application/pdf")
        form = DatasetUploadForm(data={
            'dataset_type': 'GST',
            'header_row': 1,
            'csv_delimiter': ',',
            'csv_encoding': 'utf-8'
        }, files={'data_file': fake_file})
        self.assertFalse(form.is_valid())
        self.assertIn('data_file', form.errors)

    def test_missing_required_columns_validation(self):
        """Import cannot proceed if mandatory identity field (GSTIN) is missing from column mapping."""
        mapping = {
            'Legal Name': 'legal_name',
            'Trade Name': 'trade_name',
            'Mobile': 'mobile'
        }
        is_valid, errors, warnings = validate_mapping_requirements(mapping, 'GST')
        self.assertFalse(is_valid)
        self.assertTrue(any('missing from column mapping' in e for e in errors))

    def test_renamed_columns_mapping_and_import(self):
        """User can map non-standard or renamed columns manually and import successfully."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Taxpayer ID", "Enterprise Legal Title", "Trade Handle", "WhatsApp Contact"])
        ws.append(["27RENAMED1111A1Z1", "RENAMED LEGAL NAME", "RENAMED TRADE NAME", "9876543210"])

        temp_file = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
        temp_path = temp_file.name
        temp_file.close()
        wb.save(temp_path)
        self.temp_files.append(temp_path)

        batch = ImportBatch.objects.create(
            batch_id="TEST-RENAMED-GST",
            dataset_type="GST",
            file_name="renamed_gst.xlsx",
            file_path=temp_path,
            file_checksum="chk_renamed_gst",
            file_format="XLSX",
            header_row_index=0,
            status="UPLOADED",
            created_by=self.admin,
            mapping_config={
                'columns': {
                    'Taxpayer ID': 'gst_number',
                    'Enterprise Legal Title': 'legal_name',
                    'Trade Handle': 'trade_name',
                    'WhatsApp Contact': 'mobile_number'
                }
            }
        )

        stage_and_validate_batch(batch, strict_phone=True)
        self.assertEqual(batch.proposed_new_records, 1)

        execute_import_publication(batch, user=self.admin)
        record = GSTRecord.objects.get(gst_number="27RENAMED1111A1Z1")
        self.assertEqual(record.legal_name, "RENAMED LEGAL NAME")
        self.assertEqual(record.trade_name, "RENAMED TRADE NAME")
        self.assertEqual(record.canonical_mobile, "9876543210")

    def test_invalid_gstin_pan_email_phone_values_flagged(self):
        """Invalid GSTIN, PAN, Email, and Phone values are flagged with warnings and cleaned phone stored as null."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["GSTIN", "Legal Name", "PAN", "Email", "Mobile", "Registration Date"])
        ws.append([
            "INVALIDGST123",            # Invalid GSTIN format
            "TEST BAD RECORD",
            "12345BADPAN",              # Invalid PAN
            "not_an_email",             # Invalid Email
            "5123456789",               # Starts with 5 (invalid)
            "invalid-date"              # Invalid Date
        ])

        temp_file = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
        temp_path = temp_file.name
        temp_file.close()
        wb.save(temp_path)
        self.temp_files.append(temp_path)

        batch = ImportBatch.objects.create(
            batch_id="TEST-INVALID-VALS",
            dataset_type="GST",
            file_name="invalid_vals.xlsx",
            file_path=temp_path,
            file_checksum="chk_invalid_vals",
            file_format="XLSX",
            header_row_index=0,
            status="UPLOADED",
            created_by=self.admin,
            mapping_config={
                'columns': {
                    'GSTIN': 'gst_number',
                    'Legal Name': 'legal_name',
                    'PAN': 'pan',
                    'Email': 'email',
                    'Mobile': 'mobile',
                    'Registration Date': 'registration_date'
                }
            }
        )

        stage_and_validate_batch(batch, strict_phone=True)
        staged = StagedRow.objects.get(batch=batch, row_number=1)
        self.assertEqual(staged.validation_status, 'WARNING')
        # Cleaned phone must be None / null
        self.assertIsNone(staged.cleaned_data.get('canonical_mobile'))
        # Validation messages must contain warnings
        msgs = " ".join(staged.validation_messages)
        self.assertIn("Mobile flagged", msgs)
        self.assertIn("Email", msgs)

    def test_repeated_import_of_same_file_idempotency(self):
        """Repeated import of identical file produces 0 duplicates and marks records unchanged."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["GSTIN", "Legal Name", "Trade Name", "Mobile"])
        ws.append(["27IDEMP1111A1Z1", "IDEMPOTENT CORP", "IDEMPOTENT TRADERS", "9876543210"])

        temp_file = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
        temp_path = temp_file.name
        temp_file.close()
        wb.save(temp_path)
        self.temp_files.append(temp_path)

        # Batch 1
        batch1 = ImportBatch.objects.create(
            batch_id="TEST-IDEMP-1",
            dataset_type="GST",
            file_name="idemp.xlsx",
            file_path=temp_path,
            file_checksum="chk_idemp",
            file_format="XLSX",
            header_row_index=0,
            status="UPLOADED",
            created_by=self.admin,
            mapping_config={
                'columns': {
                    'GSTIN': 'gst_number',
                    'Legal Name': 'legal_name',
                    'Trade Name': 'trade_name',
                    'Mobile': 'mobile'
                }
            }
        )
        stage_and_validate_batch(batch1, strict_phone=True)
        execute_import_publication(batch1, user=self.admin)
        self.assertEqual(GSTRecord.objects.filter(gst_number="27IDEMP1111A1Z1").count(), 1)

        # Batch 2 (same file)
        batch2 = ImportBatch.objects.create(
            batch_id="TEST-IDEMP-2",
            dataset_type="GST",
            file_name="idemp.xlsx",
            file_path=temp_path,
            file_checksum="chk_idemp",
            file_format="XLSX",
            header_row_index=0,
            status="UPLOADED",
            created_by=self.admin,
            mapping_config=batch1.mapping_config
        )
        stage_and_validate_batch(batch2, strict_phone=True)
        # Should detect as UNCHANGED
        self.assertEqual(batch2.proposed_new_records, 0)
        self.assertEqual(batch2.unchanged_records, 1)

        execute_import_publication(batch2, user=self.admin)
        # Count still 1
        self.assertEqual(GSTRecord.objects.filter(gst_number="27IDEMP1111A1Z1").count(), 1)

    def test_large_excel_file_safe_import(self):
        """Tests ingestion and staging of 150 rows without timeout or database errors."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["GSTIN", "Legal Name", "Trade Name", "Mobile", "City", "State"])
        for i in range(1, 151):
            gstin = f"27LARGE{i:04d}A1Z{i % 9}"
            ws.append([gstin, f"LARGE CORP {i}", f"LARGE TRADERS {i}", f"98765{i:05d}", "Mumbai", "Maharashtra"])

        temp_file = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
        temp_path = temp_file.name
        temp_file.close()
        wb.save(temp_path)
        self.temp_files.append(temp_path)

        batch = ImportBatch.objects.create(
            batch_id="TEST-LARGE-BATCH",
            dataset_type="GST",
            file_name="large_batch.xlsx",
            file_path=temp_path,
            file_checksum="chk_large",
            file_format="XLSX",
            header_row_index=0,
            status="UPLOADED",
            created_by=self.admin,
            mapping_config={
                'columns': {
                    'GSTIN': 'gst_number',
                    'Legal Name': 'legal_name',
                    'Trade Name': 'trade_name',
                    'Mobile': 'mobile',
                    'City': 'city',
                    'State': 'state'
                }
            }
        )

        stage_and_validate_batch(batch, strict_phone=True)
        self.assertEqual(batch.total_input_rows, 150)
        self.assertEqual(batch.proposed_new_records, 150)

        execute_import_publication(batch, user=self.admin)
        self.assertEqual(GSTRecord.objects.filter(gst_number__startswith="27LARGE").count(), 150)

    def test_unauthorized_user_access_blocked(self):
        """Unauthenticated requests to import endpoints are redirected to login."""
        self.client.logout()
        res = self.client.get(reverse('hub:import_upload'))
        self.assertEqual(res.status_code, 302)
        self.assertIn('/login/', res.url)

        res = self.client.get(reverse('hub:gst_list'))
        self.assertEqual(res.status_code, 302)
        self.assertIn('/login/', res.url)

    def test_gst_record_edit_delete_search_pagination(self):
        """Tests GST record editing, soft/hard deletion, search query, and pagination."""
        rec = GSTRecord.objects.create(
            gst_number="27SEARCH1111A1Z1",
            legal_name="SEARCHABLE LEGAL CORP",
            trade_name="SEARCHABLE TRADERS",
            canonical_mobile="9876543210",
            state="Maharashtra"
        )

        # 1. Search
        search_res = perform_unified_search("SEARCHABLE", dataset_filter='GST')
        self.assertTrue(any(r.gst_number == "27SEARCH1111A1Z1" for r in search_res['gst_records']))

        # 2. Edit
        edit_res = self.client.post(reverse('hub:gst_edit', args=[rec.id]), {
            'gst_number': '27SEARCH1111A1Z1',
            'legal_name': 'SEARCHABLE LEGAL CORP UPDATED',
            'trade_name': 'SEARCHABLE TRADERS',
            'state': 'Maharashtra',
            'raw_mobile': '9876543210'
        }, follow=True)
        self.assertEqual(edit_res.status_code, 200)
        rec.refresh_from_db()
        self.assertEqual(rec.legal_name, 'SEARCHABLE LEGAL CORP UPDATED')

        # 3. Delete
        del_res = self.client.post(reverse('hub:gst_delete', args=[rec.id]), follow=True)
        self.assertEqual(del_res.status_code, 200)
        self.assertFalse(GSTRecord.objects.filter(id=rec.id).exists())


class ClearAllGSTDataWorkflowTestCase(TestCase):
    """
    Exhaustive QA test suite for the 'Clear All GST Data' button and workflow.
    Covers:
    - Button presence in GST toolbar beside search, import, and export
    - Button absence in MCA module
    - Authorization: visible only to superuser admin, hidden from staff/unauthorized
    - Disabled state when 0 GST records exist
    - Active state when 1 or more GST records exist
    - Single button rendered (no duplicate buttons)
    - Confirmation dialog contents (title, warning text, cancel and clear buttons)
    - Cancel action leaves data unchanged
    - Direct GET URL access leaves data unchanged and warns user
    - Execution clears GST records, GST import batches, and GST quality issues
    - Execution preserves MCA companies, MCA directors, users, audit logs, and settings
    - Transaction atomicity: rollback on unexpected error leaves all records intact
    - AuditEvent logged with admin name, timestamp, deleted count, and module name
    - Table refresh, reset of search/filters/pagination, and display of empty state
    - Zero, one, and many records cleared
    - AJAX/JSON response handling
    """

    def setUp(self):
        self.admin = User.objects.create_superuser(
            username='gst_super_admin',
            email='admin@dmt.com',
            password='AdminPassword123!'
        )
        self.staff_user = User.objects.create_user(
            username='regular_staff',
            email='staff@dmt.com',
            password='StaffPassword123!',
            is_staff=True,
            is_superuser=False
        )

        # Baseline MCA data to verify strict isolation
        self.mca_company = MCACompany.objects.create(
            entity_id='U72200MH2026PTC123456',
            entity_type='company',
            name='ISOLATED MCA ENTERPRISES PVT LTD',
            state='Maharashtra'
        )
        self.mca_director = MCADirectorEntry.objects.create(
            company=self.mca_company,
            director_name='Rohan Sharma',
            raw_mobile='9820011223',
            canonical_mobile='9820011223',
            contact_validity='valid'
        )

        # Baseline System Setting
        self.system_setting = SystemSetting.objects.create(
            key='strict_phone_validation',
            value='true',
            description='Enforce 10-digit mobile phone format'
        )

    def test_button_visibility_permissions(self):
        """Admin sees Clear All GST Data button; regular staff and anonymous do not."""
        # 1. Anonymous redirected to login
        anon_res = self.client.get(reverse('hub:gst_list'))
        self.assertEqual(anon_res.status_code, 302)

        # 2. Staff user sees GST page but NO Clear All GST Data button
        self.client.force_login(self.staff_user)
        staff_res = self.client.get(reverse('hub:gst_list'))
        self.assertEqual(staff_res.status_code, 200)
        self.assertNotContains(staff_res, 'Clear All GST Data')
        self.assertNotContains(staff_res, 'id="openClearGstModalBtn"')
        self.assertNotContains(staff_res, 'id="clearGstModal"')

        # 3. Superuser admin sees Clear All GST Data button and modal
        self.client.force_login(self.admin)
        admin_res = self.client.get(reverse('hub:gst_list'))
        self.assertEqual(admin_res.status_code, 200)
        self.assertContains(admin_res, 'Clear All GST Data')
        self.assertContains(admin_res, 'id="openClearGstModalBtn"')
        self.assertContains(admin_res, 'id="clearGstModal"')

    def test_exactly_one_clear_button_rendered(self):
        """Exactly one 'Clear All GST Data' button is rendered and no duplicate button appears."""
        GSTRecord.objects.create(
            gst_number='27AABCS1429B1Z1',
            legal_name='UNIQUE GST ENTERPRISE'
        )
        self.client.force_login(self.admin)
        res = self.client.get(reverse('hub:gst_list'))
        self.assertEqual(res.status_code, 200)
        html = res.content.decode('utf-8')

        # Count occurrences of button element with Clear All GST Data
        button_count = html.count('Clear All GST Data')
        # In navbar button + modal submit button = 2 text occurrences (toolbar button + modal button)
        toolbar_btn_count = html.count('id="openClearGstModalBtn"')
        self.assertEqual(toolbar_btn_count, 1, "Exactly one Clear All GST Data button must be in the toolbar.")

    def test_button_disabled_when_zero_records(self):
        """Button is disabled when no GST records exist in the database."""
        GSTRecord.objects.all().delete()
        self.client.force_login(self.admin)
        res = self.client.get(reverse('hub:gst_list'))
        self.assertEqual(res.status_code, 200)
        html = res.content.decode('utf-8')
        self.assertIn('disabled', html)
        self.assertIn('No GST records exist to clear', html)

    def test_button_enabled_when_records_exist(self):
        """Button is active and opens confirmation modal when GST records exist."""
        GSTRecord.objects.create(
            gst_number='27AABCS1429B1Z2',
            legal_name='ACTIVE GST ENTERPRISE'
        )
        self.client.force_login(self.admin)
        res = self.client.get(reverse('hub:gst_list'))
        self.assertEqual(res.status_code, 200)
        html = res.content.decode('utf-8')
        self.assertIn('data-bs-target="#clearGstModal"', html)
        self.assertIn('Clear all GST data from the database', html)

    def test_button_absent_in_mca_module(self):
        """Do not add or display this button in the MCA module."""
        self.client.force_login(self.admin)
        for url_name in ['hub:mca_list', 'hub:mca_dashboard', 'hub:mca_search']:
            res = self.client.get(reverse(url_name))
            self.assertEqual(res.status_code, 200)
            self.assertNotContains(res, 'Clear All GST Data')
            self.assertNotContains(res, 'clearGstModal')

    def test_confirmation_dialog_content(self):
        """Dialog must show exact required title, warning text, and Cancel / Clear All GST Data buttons."""
        self.client.force_login(self.admin)
        res = self.client.get(reverse('hub:gst_list'))
        self.assertEqual(res.status_code, 200)
        html = res.content.decode('utf-8')

        self.assertIn('Clear all GST data?', html)
        self.assertIn('This will permanently remove all GST records from the database. MCA data will not be affected. This action cannot be undone.', html)
        self.assertIn('Cancel', html)
        self.assertIn('Clear All GST Data', html)
        self.assertIn('id="confirmClearGstBtn"', html)
        self.assertIn('id="cancelClearGstBtn"', html)

    def test_direct_get_request_safely_redirects_without_deletion(self):
        """Direct GET URL access does not delete anything and warns user."""
        GSTRecord.objects.create(gst_number='27GET0000001Z1', legal_name='GET SAFEGUARD CORP')
        self.client.force_login(self.admin)
        res = self.client.get(reverse('hub:gst_clear_all'), follow=True)
        self.assertEqual(res.status_code, 200)
        self.assertTrue(GSTRecord.objects.filter(gst_number='27GET0000001Z1').exists())

    def test_unauthorized_user_cannot_clear_gst_data(self):
        """Regular staff user receives 403 / redirect error and data remains untouched."""
        GSTRecord.objects.create(gst_number='27SEC0000001Z1', legal_name='SECURITY TEST CORP')
        self.client.force_login(self.staff_user)

        res = self.client.post(reverse('hub:gst_clear_all'), follow=True)
        self.assertEqual(res.status_code, 200)
        self.assertTrue(GSTRecord.objects.filter(gst_number='27SEC0000001Z1').exists())

        ajax_res = self.client.post(reverse('hub:gst_clear_all'), HTTP_X_REQUESTED_WITH='XMLHttpRequest')
        self.assertEqual(ajax_res.status_code, 403)
        self.assertTrue(GSTRecord.objects.filter(gst_number='27SEC0000001Z1').exists())

    def test_clear_all_gst_records_and_related_data_complete_flow(self):
        """
        Clears all GST records and related GST import batches & quality issues,
        while strictly preserving MCA companies, directors, settings, and users.
        """
        # Create multiple GST records
        r1 = GSTRecord.objects.create(gst_number='27MULTI00001Z1', legal_name='GST MULTI ONE')
        r2 = GSTRecord.objects.create(gst_number='27MULTI00002Z2', legal_name='GST MULTI TWO')
        r3 = GSTRecord.objects.create(gst_number='27MULTI00003Z3', legal_name='GST MULTI THREE')

        # Create GST import batch & staged row
        batch = ImportBatch.objects.create(
            batch_id='BATCH_GST_TEST_CLEAR',
            dataset_type='GST',
            file_name='gst_data.xlsx',
            file_path='/tmp/gst_data.xlsx',
            file_checksum='abc123hash',
            file_format='XLSX',
            status='COMPLETED',
            total_input_rows=3,
            imported_records=3
        )
        StagedRow.objects.create(
            batch=batch,
            row_number=1,
            raw_data={'gst_number': '27MULTI00001Z1'},
            cleaned_data={'gst_number': '27MULTI00001Z1'}
        )

        # Create GST Quality issue
        QualityIssue.objects.create(
            dataset_type='GST',
            record_type='GST',
            entity_identifier='27MULTI00001Z1',
            field_name='mobile',
            raw_value='123',
            error_code='INVALID_PHONE',
            reason_description='Invalid mobile'
        )

        # Create BusinessLink
        BusinessLink.objects.create(
            company=self.mca_company,
            gst_record=r1,
            status='APPROVED'
        )

        # Perform Clear All action
        self.client.force_login(self.admin)
        res = self.client.post(reverse('hub:gst_clear_all'), follow=True)
        self.assertEqual(res.status_code, 200)

        # 1. Verify GST records cleared
        self.assertEqual(GSTRecord.objects.count(), 0)
        self.assertEqual(ImportBatch.objects.filter(dataset_type='GST').count(), 0)
        self.assertEqual(StagedRow.objects.filter(batch__batch_id='BATCH_GST_TEST_CLEAR').count(), 0)
        self.assertEqual(QualityIssue.objects.filter(dataset_type='GST').count(), 0)
        self.assertEqual(BusinessLink.objects.count(), 0)

        # 2. Verify MCA and System integrity preserved
        self.assertEqual(MCACompany.objects.count(), 1)
        self.assertEqual(MCADirectorEntry.objects.count(), 1)
        self.assertEqual(SystemSetting.objects.count(), 1)
        self.assertTrue(User.objects.filter(username='gst_super_admin').exists())

        # 3. Verify audit log entry
        audit = AuditEvent.objects.filter(model_name='GSTRecord', action='DELETE').first()
        self.assertIsNotNone(audit)
        self.assertEqual(audit.user, self.admin)
        self.assertEqual(audit.dataset_type, 'GST')
        self.assertIn('Cleared all GST data', audit.object_repr)
        self.assertEqual(audit.changes.get('deleted_count'), 3)
        self.assertEqual(audit.changes.get('module_name'), 'GST')
        self.assertEqual(audit.changes.get('admin_username'), 'gst_super_admin')

        # 4. Verify table refresh & empty state displayed in response
        html = res.content.decode('utf-8')
        self.assertIn('Successfully cleared all 3 GST record(s)', html)
        self.assertIn('No matching GST records found', html)
        self.assertIn('No GST registrations exist in the database', html)

    def test_ajax_api_clear_all_endpoint(self):
        """AJAX requests receive structured JSON with count and redirect_url."""
        GSTRecord.objects.create(gst_number='27AJAX000001Z1', legal_name='AJAX TEST CORP')
        self.client.force_login(self.admin)

        res = self.client.post(
            reverse('hub:gst_clear_all'),
            HTTP_X_REQUESTED_WITH='XMLHttpRequest',
            HTTP_ACCEPT='application/json'
        )
        self.assertEqual(res.status_code, 200)
        data = res.json()
        self.assertTrue(data.get('success'))
        self.assertEqual(data.get('cleared_count'), 1)
        self.assertEqual(data.get('redirect_url'), reverse('hub:gst_list'))
        self.assertEqual(GSTRecord.objects.count(), 0)

    def test_clear_when_zero_records_handled_gracefully(self):
        """Attempting to clear when 0 records exist handles safely without errors."""
        GSTRecord.objects.all().delete()
        self.client.force_login(self.admin)

        res = self.client.post(reverse('hub:gst_clear_all'), follow=True)
        self.assertEqual(res.status_code, 200)
        self.assertIn('There are no GST records in the database to clear', res.content.decode('utf-8'))

    def test_atomic_rollback_on_database_error(self):
        """If an error occurs during deletion, transaction rolls back and data remains intact."""
        GSTRecord.objects.create(gst_number='27ERR0000001Z1', legal_name='ERROR ROLLBACK CORP')
        self.client.force_login(self.admin)

        from unittest.mock import patch
        with patch.object(GSTRecord.objects, 'all', side_effect=RuntimeError("Simulated DB Disk Failure")):
            res = self.client.post(reverse('hub:gst_clear_all'), follow=True)
            self.assertEqual(res.status_code, 200)
            self.assertIn('Failed to clear GST records', res.content.decode('utf-8'))

        # Record must still exist
        self.assertEqual(GSTRecord.objects.count(), 1)


class GSTExcelImportWorkflowTestCase(TestCase):
    """
    Comprehensive regression test suite for GST Excel Import workflow.
    Validates end-to-end import flow across all 13 required test cases:
    1. Valid .xlsx GST file
    2. Valid .xls GST file
    3. Empty file (0 bytes)
    4. Corrupted file
    5. Wrong file type (.pdf, .txt)
    6. MCA file uploaded in GST module (schema isolation rejection)
    7. Missing required GST columns (mandatory gst_number/gstin)
    8. Duplicate GSTIN values (in-file and against existing DB)
    9. Invalid GSTIN, PAN, email, mobile, and date cleaning
    10. Large Excel file / 50MB file size limit
    11. Network / API failure / atomic database transaction rollback
    12. Unauthorized user permissions
    13. Repeated upload of the same file (idempotent safe updates)
    """

    def setUp(self):
        self.admin = User.objects.create_superuser(
            username='gst_importer_admin',
            email='admin@gstimport.in',
            password='AdminPassword123!'
        )
        self.user = User.objects.create_user(
            username='gst_staff_user',
            email='staff@gstimport.in',
            password='StaffPassword123!'
        )
        self.client.force_login(self.admin)

    def _create_xlsx_file(self, headers, rows, filename="test_gst.xlsx", sheet_name="GST Registrations"):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = sheet_name
        ws.append(headers)
        for r in rows:
            ws.append(r)
        bio = io.BytesIO()
        wb.save(bio)
        bio.seek(0)
        return SimpleUploadedFile(
            filename,
            bio.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )

    def _create_xls_file(self, headers, rows, filename="test_gst.xls", sheet_name="GST Registrations"):
        wb = xlwt.Workbook(encoding='utf-8')
        ws = wb.add_sheet(sheet_name)
        for col_idx, h in enumerate(headers):
            ws.write(0, col_idx, h)
        for row_idx, r in enumerate(rows, start=1):
            for col_idx, val in enumerate(r):
                ws.write(row_idx, col_idx, val)
        bio = io.BytesIO()
        wb.save(bio)
        bio.seek(0)
        return SimpleUploadedFile(
            filename,
            bio.getvalue(),
            content_type="application/vnd.ms-excel"
        )

    def test_01_valid_xlsx_gst_file_upload_and_import(self):
        """Case 1: Upload and import valid .xlsx GST file end-to-end."""
        headers = [
            'gst_number', 'legal_name', 'trade_name', 'pan', 'mobile',
            'email', 'address', 'city', 'state', 'pincode', 'status', 'registration_date'
        ]
        rows = [
            [
                '27AABCS1429B1Z1', 'Apex Solutions Private Limited', 'Apex Tech',
                'AABCS1429B', '9820011223', 'contact@apex.in', '101 Tech Park',
                'Mumbai', 'Maharashtra', '400051', 'ACT', '2026-09-25'
            ],
            [
                '27BAPPL9876C1Z4', 'Bharat Logistics Limited', 'Bharat Express',
                'BAPPL9876C', '+91 9876543210', 'info@bharat.in', '404 Cargo Way',
                'Pune', 'Maharashtra', '411001', 'ACT', '2026-09-26'
            ]
        ]
        uploaded_file = self._create_xlsx_file(headers, rows)

        # 1. Upload to dedicated GST endpoint
        upload_res = self.client.post(reverse('hub:gst_upload'), {
            'data_file': uploaded_file,
            'sheet_name': 'GST Registrations'
        }, follow=False)

        self.assertEqual(upload_res.status_code, 302)
        batch = ImportBatch.objects.filter(dataset_type='GST').order_by('-created_at').first()
        self.assertIsNotNone(batch)
        self.assertEqual(batch.file_format, 'XLSX')
        self.assertEqual(batch.status, 'UPLOADED')
        self.assertIn(reverse('hub:import_map', kwargs={'batch_id': batch.batch_id}), upload_res.url)

        # 2. Confirm column mapping
        map_post_data = {
            'date_format': '%Y-%m-%d'
        }
        for h in headers:
            map_post_data[f"map_{h}"] = h
        map_res = self.client.post(
            reverse('hub:import_map', kwargs={'batch_id': batch.batch_id}),
            map_post_data,
            follow=False
        )
        self.assertEqual(map_res.status_code, 302)
        batch.refresh_from_db()
        self.assertEqual(batch.total_input_rows, 2)
        self.assertEqual(batch.proposed_new_records, 2)

        # 3. Apply and execute publication
        exec_res = self.client.post(
            reverse('hub:import_execute', kwargs={'batch_id': batch.batch_id}),
            follow=True
        )
        self.assertEqual(exec_res.status_code, 200)
        self.assertIn(reverse('hub:gst_list'), exec_res.redirect_chain[0][0])

        # 4. Verify records in database
        self.assertEqual(GSTRecord.objects.count(), 2)
        apex = GSTRecord.objects.get(gst_number='27AABCS1429B1Z1')
        self.assertEqual(apex.legal_name, 'Apex Solutions Private Limited')
        self.assertEqual(apex.trade_name, 'Apex Tech')
        self.assertEqual(apex.canonical_mobile, '9820011223')
        self.assertEqual(apex.contact_validity, 'valid')

        bharat = GSTRecord.objects.get(gst_number='27BAPPL9876C1Z4')
        self.assertEqual(bharat.canonical_mobile, '9876543210')
        self.assertEqual(bharat.contact_validity, 'cleaned')

    def test_02_valid_xls_gst_file_upload_and_import(self):
        """Case 2: Upload and import valid legacy .xls GST file end-to-end."""
        headers = [
            'gst_number', 'legal_name', 'trade_name', 'pan', 'mobile',
            'email', 'address', 'city', 'state', 'pincode', 'status', 'registration_date'
        ]
        rows = [
            [
                '27LEGAC1234F1Z9', 'Legacy Enterprise Private Limited', 'Legacy Retail',
                'LEGAC1234F', '9123456780', 'contact@legacy.in', '12 Old Mill Road',
                'Nagpur', 'Maharashtra', '440001', 'ACT', '2026-09-20'
            ]
        ]
        uploaded_file = self._create_xls_file(headers, rows)

        # 1. Upload to dedicated GST endpoint
        upload_res = self.client.post(reverse('hub:gst_upload'), {
            'data_file': uploaded_file
        }, follow=False)
        self.assertEqual(upload_res.status_code, 302)

        batch = ImportBatch.objects.filter(dataset_type='GST').order_by('-created_at').first()
        self.assertIsNotNone(batch)
        self.assertEqual(batch.file_format, 'XLS')

        # 2. Confirm mapping
        map_post_data = {'date_format': '%Y-%m-%d'}
        for h in headers:
            map_post_data[f"map_{h}"] = h
        self.client.post(reverse('hub:import_map', kwargs={'batch_id': batch.batch_id}), map_post_data)

        # 3. Execute publication
        exec_res = self.client.post(reverse('hub:import_execute', kwargs={'batch_id': batch.batch_id}), follow=True)
        self.assertEqual(exec_res.status_code, 200)

        # 4. Verify record in database
        legacy_rec = GSTRecord.objects.filter(gst_number='27LEGAC1234F1Z9').first()
        self.assertIsNotNone(legacy_rec)
        self.assertEqual(legacy_rec.legal_name, 'Legacy Enterprise Private Limited')
        self.assertEqual(legacy_rec.canonical_mobile, '9123456780')

    def test_03_empty_file_rejected(self):
        """Case 3: Empty (0 bytes) file is rejected immediately with error message."""
        empty_file = SimpleUploadedFile(
            'empty_file.xlsx',
            b'',
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        res = self.client.post(reverse('hub:gst_upload'), {
            'data_file': empty_file
        }, follow=True)

        self.assertEqual(res.status_code, 200)
        self.assertIn('The selected file is empty (0 bytes)', res.content.decode('utf-8'))
        self.assertEqual(ImportBatch.objects.filter(file_name='empty_file.xlsx').count(), 0)

    def test_04_corrupted_file_rejected(self):
        """Case 4: Corrupted file is rejected with clear error message and no batch created."""
        corrupted_file = SimpleUploadedFile(
            'damaged_spreadsheet.xlsx',
            b'PK\x03\x04NOT_A_VALID_ZIP_PACKAGE_COMPLETELY_DAMAGED',
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        res = self.client.post(reverse('hub:gst_upload'), {
            'data_file': corrupted_file
        }, follow=True)

        self.assertEqual(res.status_code, 200)
        self.assertIn('Corrupted or invalid Excel file', res.content.decode('utf-8'))
        self.assertEqual(ImportBatch.objects.filter(file_name='damaged_spreadsheet.xlsx').count(), 0)

    def test_05_wrong_file_type_rejected(self):
        """Case 5: Unsupported file types (.pdf, .txt, .docx) are rejected."""
        pdf_file = SimpleUploadedFile(
            'invoice.pdf',
            b'%PDF-1.4 sample pdf content',
            content_type="application/pdf"
        )
        res = self.client.post(reverse('hub:gst_upload'), {
            'data_file': pdf_file
        }, follow=True)

        self.assertEqual(res.status_code, 200)
        self.assertIn('Unsupported file format (.pdf)', res.content.decode('utf-8'))
        self.assertEqual(ImportBatch.objects.filter(file_name='invoice.pdf').count(), 0)

    def test_06_mca_file_uploaded_in_gst_module_rejected(self):
        """Case 6: Schema isolation - uploading an MCA corporate file in GST module is rejected."""
        mca_headers = ['cin', 'company_name', 'entity_type', 'roc', 'director_name', 'director_mobile']
        mca_rows = [
            ['U72200MH2026PTC123456', 'MCA Corporate Tech Ltd', 'company', 'ROC-Mumbai', 'Director A', '9820011223']
        ]
        mca_file = self._create_xlsx_file(mca_headers, mca_rows, filename="mca_dataset.xlsx")

        res = self.client.post(reverse('hub:gst_upload'), {
            'data_file': mca_file
        }, follow=True)

        self.assertEqual(res.status_code, 200)
        self.assertIn('MCA file detected', res.content.decode('utf-8'))
        self.assertIn('cannot be imported into the GST module', res.content.decode('utf-8'))
        self.assertEqual(ImportBatch.objects.filter(file_name='mca_dataset.xlsx').count(), 0)

    def test_07_missing_required_gst_columns_prevents_import(self):
        """Case 7: Mandatory GST column (gst_number/gstin) missing prevents import progression."""
        headers = ['legal_name', 'trade_name', 'mobile', 'email', 'city', 'state']
        rows = [['Legal Without GST', 'Trade Without GST', '9820011223', 'contact@nogst.in', 'Mumbai', 'Maharashtra']]
        uploaded_file = self._create_xlsx_file(headers, rows)

        upload_res = self.client.post(reverse('hub:gst_upload'), {'data_file': uploaded_file})
        self.assertEqual(upload_res.status_code, 302)
        batch = ImportBatch.objects.filter(dataset_type='GST').order_by('-created_at').first()

        # Submit mapping without gst_number
        map_res = self.client.post(
            reverse('hub:import_map', kwargs={'batch_id': batch.batch_id}),
            {
                'map_legal_name': 'legal_name',
                'map_trade_name': 'trade_name',
                'map_mobile': 'mobile',
                'date_format': '%Y-%m-%d'
            }
        )
        self.assertEqual(map_res.status_code, 200)
        content_str = map_res.content.decode('utf-8')
        self.assertIn("Mandatory field", content_str)
        self.assertIn("is missing from column mapping", content_str)
        batch.refresh_from_db()
        self.assertEqual(batch.status, 'UPLOADED')
        self.assertEqual(GSTRecord.objects.count(), 0)

    def test_08_duplicate_gstin_values_detected_and_handled(self):
        """Case 8: In-file duplicate GSTIN values are detected, flagged, and handled safely."""
        headers = ['gst_number', 'legal_name', 'trade_name', 'mobile', 'city', 'state']
        rows = [
            ['27DUP1234567Z1', 'Initial Legal Name', 'Initial Trade Name', '9820011223', 'Mumbai', 'Maharashtra'],
            ['27DUP1234567Z1', 'Updated Legal Name', 'Updated Trade Name', '9820011223', 'Mumbai', 'Maharashtra']
        ]
        uploaded_file = self._create_xlsx_file(headers, rows)

        self.client.post(reverse('hub:gst_upload'), {'data_file': uploaded_file})
        batch = ImportBatch.objects.filter(dataset_type='GST').order_by('-created_at').first()

        map_data = {f"map_{h}": h for h in headers}
        map_data['date_format'] = '%Y-%m-%d'
        self.client.post(reverse('hub:import_map', kwargs={'batch_id': batch.batch_id}), map_data)

        batch.refresh_from_db()
        self.assertEqual(batch.conflict_count, 1)
        self.assertEqual(batch.proposed_new_records, 1)
        self.assertEqual(batch.proposed_updates, 1)

        # Publish and verify single record with updated trade name exists
        self.client.post(reverse('hub:import_execute', kwargs={'batch_id': batch.batch_id}))
        self.assertEqual(GSTRecord.objects.filter(gst_number='27DUP1234567Z1').count(), 1)
        dup_rec = GSTRecord.objects.get(gst_number='27DUP1234567Z1')
        self.assertEqual(dup_rec.trade_name, 'Updated Trade Name')

    def test_09_invalid_gstin_pan_email_mobile_and_date_cleaning(self):
        """Case 9: Cleans invalid GSTIN, PAN, email, mobile, and date values safely."""
        headers = [
            'gst_number', 'legal_name', 'trade_name', 'pan', 'mobile',
            'email', 'address', 'city', 'state', 'pincode', 'status', 'registration_date'
        ]
        rows = [
            [
                '27CLEAN1234F1Z5', '  Dirty Text Entity  ', 'Dirty Trade  ',
                'INVALID_PAN', '12345', 'not-a-valid-email', '  Address With Whitespace  ',
                'Mumbai', 'Maharashtra', '400051', 'ACT', 'invalid-date'
            ]
        ]
        uploaded_file = self._create_xlsx_file(headers, rows)

        self.client.post(reverse('hub:gst_upload'), {'data_file': uploaded_file})
        batch = ImportBatch.objects.filter(dataset_type='GST').order_by('-created_at').first()

        map_data = {f"map_{h}": h for h in headers}
        map_data['date_format'] = '%Y-%m-%d'
        self.client.post(reverse('hub:import_map', kwargs={'batch_id': batch.batch_id}), map_data)

        # Staged validation should flag mobile
        staged = StagedRow.objects.filter(batch=batch).first()
        self.assertEqual(staged.cleaned_data.get('phone_validity'), 'invalid')
        self.assertIsNone(staged.cleaned_data.get('canonical_mobile'))
        self.assertEqual(staged.cleaned_data.get('legal_name'), 'Dirty Text Entity')
        self.assertEqual(staged.cleaned_data.get('trade_name'), 'Dirty Trade')

        # Execute and verify QualityIssue was created
        self.client.post(reverse('hub:import_execute', kwargs={'batch_id': batch.batch_id}))
        record = GSTRecord.objects.get(gst_number='27CLEAN1234F1Z5')
        self.assertEqual(record.legal_name, 'Dirty Text Entity')
        self.assertIsNone(record.canonical_mobile)
        self.assertEqual(record.raw_mobile, '12345')
        self.assertEqual(record.contact_validity, 'invalid')
        self.assertIsNone(record.registration_date)

        issue = QualityIssue.objects.filter(batch=batch, entity_identifier='27CLEAN1234F1Z5').first()
        self.assertIsNotNone(issue)
        self.assertEqual(issue.error_code, 'INVALID_PHONE')

    def test_10_large_excel_file(self):
        """Case 10: Multi-row Excel file ingestion and 50MB file size limit."""
        # 1. Multi-row Excel file ingestion
        headers = ['gst_number', 'legal_name', 'trade_name', 'mobile', 'state']
        rows = [
            [f'27TEST{i:07d}Z1', f'Company {i}', f'Trade {i}', f'98200{i:05d}', 'Maharashtra']
            for i in range(1, 35)
        ]
        uploaded_file = self._create_xlsx_file(headers, rows)
        self.client.post(reverse('hub:gst_upload'), {'data_file': uploaded_file})
        batch = ImportBatch.objects.filter(dataset_type='GST').order_by('-created_at').first()

        map_data = {f"map_{h}": h for h in headers}
        self.client.post(reverse('hub:import_map', kwargs={'batch_id': batch.batch_id}), map_data)
        self.client.post(reverse('hub:import_execute', kwargs={'batch_id': batch.batch_id}))
        self.assertEqual(GSTRecord.objects.count(), 34)

        # 2. Rejection of file over 50MB
        oversized_file = SimpleUploadedFile(
            "huge_dataset.xlsx",
            b"header,content\n1,2",
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        with patch('django.core.files.uploadedfile.UploadedFile.size', new_callable=PropertyMock) as mock_size:
            mock_size.return_value = 55 * 1024 * 1024  # 55 MB
            res = self.client.post(reverse('hub:gst_upload'), {
                'data_file': oversized_file
            }, follow=True)
            self.assertIn('File size exceeds 50 MB limit', res.content.decode('utf-8'))

    def test_11_network_api_failure_atomic_rollback(self):
        """Case 11: Transaction atomicity - failure during publication rolls back all changes."""
        headers = ['gst_number', 'legal_name', 'trade_name', 'mobile', 'state']
        rows = [
            ['27ATOMIC1234Z1', 'Atomic Enterprise 1', 'Trade 1', '9820011223', 'Maharashtra'],
            ['27ATOMIC1234Z2', 'Atomic Enterprise 2', 'Trade 2', '9820011224', 'Maharashtra']
        ]
        uploaded_file = self._create_xlsx_file(headers, rows)
        self.client.post(reverse('hub:gst_upload'), {'data_file': uploaded_file})
        batch = ImportBatch.objects.filter(dataset_type='GST').order_by('-created_at').first()

        map_data = {f"map_{h}": h for h in headers}
        self.client.post(reverse('hub:import_map', kwargs={'batch_id': batch.batch_id}), map_data)

        # Simulate unexpected DB failure on record creation
        with patch.object(GSTRecord.objects, 'create', side_effect=RuntimeError("Database Connection Severed")):
            res = self.client.post(reverse('hub:import_execute', kwargs={'batch_id': batch.batch_id}), follow=True)
            self.assertEqual(res.status_code, 200)
            self.assertIn('Import failed', res.content.decode('utf-8'))

        # Verify no partial records remain
        self.assertEqual(GSTRecord.objects.filter(gst_number__startswith='27ATOMIC').count(), 0)
        batch.refresh_from_db()
        self.assertEqual(batch.status, 'FAILED')

    def test_12_unauthorized_user(self):
        """Case 12: Anonymous or unauthorized users cannot access GST upload endpoint."""
        self.client.logout()
        res = self.client.post(reverse('hub:gst_upload'), {})
        self.assertEqual(res.status_code, 302)
        self.assertIn('/login/', res.url)

    def test_13_repeated_upload_of_same_file(self):
        """Case 13: Repeated upload of the same file performs idempotent safe updates."""
        headers = ['gst_number', 'legal_name', 'trade_name', 'mobile', 'state']
        rows = [
            ['27REPEAT1234Z1', 'Repeat Enterprise', 'Trade Initial', '9820011223', 'Maharashtra']
        ]
        uploaded_file_1 = self._create_xlsx_file(headers, rows)

        # First upload and import
        self.client.post(reverse('hub:gst_upload'), {'data_file': uploaded_file_1})
        batch1 = ImportBatch.objects.filter(dataset_type='GST').order_by('-created_at').first()
        map_data = {f"map_{h}": h for h in headers}
        self.client.post(reverse('hub:import_map', kwargs={'batch_id': batch1.batch_id}), map_data)
        self.client.post(reverse('hub:import_execute', kwargs={'batch_id': batch1.batch_id}))

        self.assertEqual(GSTRecord.objects.filter(gst_number='27REPEAT1234Z1').count(), 1)
        rec = GSTRecord.objects.get(gst_number='27REPEAT1234Z1')
        self.assertEqual(rec.revision, 1)

        # Second upload of same file (with updated trade name)
        rows_updated = [
            ['27REPEAT1234Z1', 'Repeat Enterprise', 'Trade Updated', '9820011223', 'Maharashtra']
        ]
        uploaded_file_2 = self._create_xlsx_file(headers, rows_updated)
        self.client.post(reverse('hub:gst_upload'), {'data_file': uploaded_file_2})
        batch2 = ImportBatch.objects.filter(dataset_type='GST').order_by('-created_at').first()
        self.client.post(reverse('hub:import_map', kwargs={'batch_id': batch2.batch_id}), map_data)

        batch2.refresh_from_db()
        self.assertEqual(batch2.proposed_updates, 1)
        self.assertEqual(batch2.proposed_new_records, 0)

        self.client.post(reverse('hub:import_execute', kwargs={'batch_id': batch2.batch_id}))

        # Still only 1 record, but updated and revision incremented
        self.assertEqual(GSTRecord.objects.filter(gst_number='27REPEAT1234Z1').count(), 1)
        rec.refresh_from_db()
        self.assertEqual(rec.trade_name, 'Trade Updated')
        self.assertEqual(rec.revision, 2)


class ColumnFastSearchTestCase(TestCase):
    """
    Verifies per-column fast search filtering across MCA and GST modules:
    - Case-insensitive partial matching on Company Name and other text fields
    - Direct company name searching without requiring industry category
    - Preserving text on phone numbers and identifiers (leading zeros)
    - Multiple combined filters with strict AND logic
    - Individual filter clearing and clear-all behaviour
    - Preservation during pagination and export synchronization
    - Debounced AJAX partial response structure
    """
    def setUp(self):
        self.client = Client()
        self.user = User.objects.create_superuser('admin_search', 'admin_search@test.com', 'adminpass123')
        self.client.force_login(self.user)

        # Create representative MCA companies
        self.c1 = MCACompany.objects.create(
            entity_id='U72200MH2021PTC123456',
            entity_type='company',
            name='ABC Technologies Private Limited',
            state='Maharashtra',
            city='Pune',
            roc='ROC Pune',
            email='info@abctech.com'
        )
        MCADirectorEntry.objects.create(
            company=self.c1,
            director_name='Rajesh Sharma',
            raw_mobile='09822012345',
            canonical_mobile='9822012345',
            contact_validity='valid'
        )

        self.c2 = MCACompany.objects.create(
            entity_id='U72900DL2022PTC654321',
            entity_type='company',
            name='Tech Solutions Global Limited',
            state='Delhi',
            city='New Delhi',
            roc='ROC Delhi',
            email='contact@techsolutions.com'
        )
        MCADirectorEntry.objects.create(
            company=self.c2,
            director_name='Amit Patel',
            raw_mobile='9811099887',
            canonical_mobile='9811099887',
            contact_validity='valid'
        )

        self.c3 = MCACompany.objects.create(
            entity_id='AAA-9876',
            entity_type='llp',
            name='Sunshine Agro Foods LLP',
            state='Maharashtra',
            city='Pune',
            roc='ROC Pune',
            email='sunshine@agro.com'
        )

        # Create representative GST records
        self.g1 = GSTRecord.objects.create(
            gst_number='27AAACH1234F1Z1',
            legal_name='CHITRA INFOTECH PRIVATE LIMITED',
            trade_name='Chitra Cloud Solutions',
            status='ACT',
            source_type='REG',
            pan='AAACH1234F',
            raw_mobile='09822012345',
            canonical_mobile='9822012345',
            contact_validity='valid',
            state='Maharashtra',
            city='Pune',
            address='101 Tech Park, Hinjawadi, Pune'
        )

        self.g2 = GSTRecord.objects.create(
            gst_number='07AABCE5678G2Z2',
            legal_name='EXPRESS COURIER SERVICES',
            trade_name='Delhi Speed Express',
            status='ACT',
            source_type='REG',
            pan='AABCE5678G',
            raw_mobile='9811054321',
            canonical_mobile='9811054321',
            contact_validity='valid',
            state='Delhi',
            city='New Delhi',
            address='Connaught Place, New Delhi'
        )

        self.g3 = GSTRecord.objects.create(
            gst_number='27CANCL9999K1Z9',
            legal_name='CANCELLED TRADERS LLP',
            trade_name='Cancelled Enterprise',
            status='CAN',
            source_type='REG',
            pan='CANCL9999K',
            state='Maharashtra',
            city='Mumbai',
            address='Nariman Point, Mumbai'
        )

    def test_mca_name_partial_match(self):
        """Typing 'tech' matches 'ABC Technologies' and 'Tech Solutions', excludes 'Sunshine Agro'."""
        res = self.client.get(reverse('hub:mca_list'), {'f_name': 'tech'})
        self.assertEqual(res.status_code, 200)
        matched_ids = [c.entity_id for c in res.context['page_obj']]
        self.assertIn('U72200MH2021PTC123456', matched_ids)
        self.assertIn('U72900DL2022PTC654321', matched_ids)
        self.assertNotIn('AAA-9876', matched_ids)
        self.assertEqual(res.context['total_count'], 2)

    def test_mca_combined_name_and_city_and_logic(self):
        """Combined filters: Name contains 'tech' AND City contains 'Pune' matches only ABC Technologies."""
        res = self.client.get(reverse('hub:mca_list'), {'f_name': 'tech', 'f_city': 'Pune'})
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.context['total_count'], 1)
        matched = res.context['page_obj'][0]
        self.assertEqual(matched.name, 'ABC Technologies Private Limited')
        self.assertEqual(matched.city, 'Pune')

    def test_mca_identifier_filter_preserves_text(self):
        """Identifier search treats entity_id as text preserving characters/leading digits."""
        res = self.client.get(reverse('hub:mca_list'), {'f_entity_id': 'U72200'})
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.context['total_count'], 1)
        self.assertEqual(res.context['page_obj'][0].entity_id, 'U72200MH2021PTC123456')

    def test_mca_director_filter(self):
        """Director filter matches director name and mobile number as text."""
        res = self.client.get(reverse('hub:mca_list'), {'f_director': 'Sharma'})
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.context['total_count'], 1)
        self.assertEqual(res.context['page_obj'][0].entity_id, 'U72200MH2021PTC123456')

        # Mobile number text matching
        res_phone = self.client.get(reverse('hub:mca_list'), {'f_phone': '9822012345'})
        self.assertEqual(res_phone.status_code, 200)
        self.assertEqual(res_phone.context['total_count'], 1)

    def test_mca_ajax_partial_response(self):
        """AJAX request with format=partial returns structured JSON with rendered html and total_count."""
        res = self.client.get(reverse('hub:mca_list'), {'f_name': 'tech', 'format': 'partial'}, HTTP_X_REQUESTED_WITH='XMLHttpRequest')
        self.assertEqual(res.status_code, 200)
        import json
        data = json.loads(res.content)
        self.assertTrue(data.get('success'))
        self.assertEqual(data.get('total_count'), 2)
        self.assertIn('ABC Technologies', data.get('html'))
        self.assertIn('Tech Solutions', data.get('html'))
        self.assertNotIn('Sunshine Agro', data.get('html'))

    def test_mca_filtered_exports(self):
        """Filtered exports contain ONLY records matching the active column filters."""
        # JSON export with f_name=tech AND f_city=Pune
        res_json = self.client.get(reverse('hub:mca_export_json'), {'f_name': 'tech', 'f_city': 'Pune'})
        self.assertEqual(res_json.status_code, 200)
        import json
        data = json.loads(res_json.content)
        self.assertEqual(len(data), 1)
        self.assertEqual(data[0]['name'], 'ABC Technologies Private Limited')

        # Excel export with scope=filtered
        res_excel = self.client.get(reverse('hub:mca_export_excel'), {'scope': 'filtered', 'f_name': 'tech', 'f_city': 'Pune'})
        self.assertEqual(res_excel.status_code, 200)
        self.assertIn('spreadsheetml', res_excel.get('Content-Type'))

    def test_gst_column_filters_and_combined_and_logic(self):
        """GST filters support trade name, legal name, status, city, state and combined AND logic."""
        # 1. Trade name partial matching
        res_trade = self.client.get(reverse('hub:gst_list'), {'f_trade_name': 'cloud'})
        self.assertEqual(res_trade.status_code, 200)
        self.assertEqual(res_trade.context['total_count'], 1)
        self.assertEqual(res_trade.context['page_obj'][0].gst_number, '27AAACH1234F1Z1')

        # 2. Combined: Status='ACT' AND State='Maharashtra'
        res_combined = self.client.get(reverse('hub:gst_list'), {'f_status': 'ACT', 'f_state': 'Maharashtra'})
        self.assertEqual(res_combined.status_code, 200)
        self.assertEqual(res_combined.context['total_count'], 1)
        self.assertEqual(res_combined.context['page_obj'][0].gst_number, '27AAACH1234F1Z1')

        # 3. Status='CAN'
        res_can = self.client.get(reverse('hub:gst_list'), {'f_status': 'CAN'})
        self.assertEqual(res_can.status_code, 200)
        self.assertEqual(res_can.context['total_count'], 1)
        self.assertEqual(res_can.context['page_obj'][0].gst_number, '27CANCL9999K1Z9')

    def test_gst_filtered_exports(self):
        """GST exports contain ONLY records matching active filters."""
        # CSV Export with scope=filtered
        res_csv = self.client.get(reverse('hub:gst_export_csv'), {'scope': 'filtered', 'f_status': 'ACT', 'f_state': 'Maharashtra'})
        self.assertEqual(res_csv.status_code, 200)
        content = res_csv.content.decode('utf-8')
        self.assertIn('27AAACH1234F1Z1', content)
        self.assertNotIn('27CANCL9999K1Z9', content)
        self.assertNotIn('07AABCE5678G2Z2', content)

        # JSON Export
        res_json = self.client.get(reverse('hub:gst_export_json'), {'f_status': 'ACT', 'f_city': 'Pune'})
        self.assertEqual(res_json.status_code, 200)
        import json
        data = json.loads(res_json.content)
        self.assertEqual(len(data), 1)
        self.assertEqual(data[0]['gst_number'], '27AAACH1234F1Z1')

    def test_mca_filtered_excel_export_detailed(self):
        """Excel export contains only companies matching active filters, text format '@', and matching count."""
        # Active filter: Name='tech' AND City='Pune' (matches ABC Technologies only)
        res = self.client.get(reverse('hub:mca_export_excel'), {
            'scope': 'filtered',
            'f_name': 'tech',
            'f_city': 'Pune'
        })
        self.assertEqual(res.status_code, 200)
        self.assertIn('spreadsheetml', res.get('Content-Type'))
        self.assertIn('attachment', res.get('Content-Disposition'))

        wb = openpyxl.load_workbook(io.BytesIO(res.content))
        ws = wb.active
        rows = list(ws.iter_rows(values_only=False))
        # 1 header row + 1 data row
        self.assertEqual(len(rows), 2)

        # Header check
        headers = [cell.value for cell in rows[0]]
        self.assertIn("entity_id", headers)
        self.assertIn("name", headers)
        self.assertIn("canonical_mobile", headers)

        # Data check
        cin_idx = headers.index("entity_id")
        name_idx = headers.index("name")
        phone_idx = headers.index("canonical_mobile")

        data_row = rows[1]
        self.assertEqual(data_row[cin_idx].value, 'U72200MH2021PTC123456')
        self.assertEqual(data_row[name_idx].value, 'ABC Technologies Private Limited')
        self.assertIn('9822012345', str(data_row[phone_idx].value))

        # Check cell number_format is '@' (text) to preserve leading zeros
        self.assertEqual(data_row[cin_idx].number_format, '@')
        self.assertEqual(data_row[phone_idx].number_format, '@')

    def test_mca_filtered_excel_export_all_when_no_filter(self):
        """When no filters are active, exports all authorized records across all pages."""
        res = self.client.get(reverse('hub:mca_export_excel'), {'scope': 'filtered'})
        self.assertEqual(res.status_code, 200)
        wb = openpyxl.load_workbook(io.BytesIO(res.content))
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        # 1 header + 3 companies
        self.assertEqual(len(rows), 4)

    def test_mca_filtered_excel_export_zero_records(self):
        """When 0 records match filters, returns 400 for AJAX and redirects with flash message for standard GET."""
        # 1. Standard GET
        res_standard = self.client.get(reverse('hub:mca_export_excel'), {
            'scope': 'filtered',
            'f_name': 'nonexistentcompany999'
        }, follow=True)
        self.assertEqual(res_standard.status_code, 200)
        messages = list(res_standard.context.get('messages', []))
        self.assertTrue(any("No matching records to export." in str(m) for m in messages))

        # 2. AJAX GET
        res_ajax = self.client.get(reverse('hub:mca_export_excel'), {
            'scope': 'filtered',
            'f_name': 'nonexistentcompany999'
        }, HTTP_X_REQUESTED_WITH='XMLHttpRequest')
        self.assertEqual(res_ajax.status_code, 400)
        import json
        data = json.loads(res_ajax.content)
        self.assertEqual(data.get('message'), "No matching records to export.")

    def test_gst_filtered_excel_export_detailed(self):
        """GST Excel export contains only records matching filters, with leading zeros and '@' format."""
        # Active filter: Status='ACT' AND State='Maharashtra' (matches 27AAACH1234F1Z1 only)
        res = self.client.get(reverse('hub:gst_export_excel'), {
            'scope': 'filtered',
            'f_status': 'ACT',
            'f_state': 'Maharashtra'
        })
        self.assertEqual(res.status_code, 200)
        self.assertIn('spreadsheetml', res.get('Content-Type'))

        wb = openpyxl.load_workbook(io.BytesIO(res.content))
        ws = wb.active
        rows = list(ws.iter_rows(values_only=False))
        # 1 header + 1 data row
        self.assertEqual(len(rows), 2)

        headers = [cell.value for cell in rows[0]]
        gstin_idx = headers.index("gst_number")
        phone_idx = headers.index("canonical_mobile")

        data_row = rows[1]
        self.assertEqual(data_row[gstin_idx].value, '27AAACH1234F1Z1')
        self.assertEqual(str(data_row[phone_idx].value), '9822012345')
        self.assertEqual(data_row[phone_idx].number_format, '@')
        self.assertEqual(data_row[gstin_idx].number_format, '@')

    def test_excel_export_preserves_leading_zeros(self):
        """Identifiers and phone numbers with leading zeros are preserved as strings with '@' format."""
        # Create an MCA company with leading zeros in entity_id and phone
        c_zeros = MCACompany.objects.create(
            entity_id='00123456789012345678',
            entity_type='company',
            name='Zero Lead Tech Pvt Ltd',
            state='Maharashtra',
            city='Pune'
        )
        MCADirectorEntry.objects.create(
            company=c_zeros,
            director_name='Zero Director',
            raw_mobile='0123456789',
            canonical_mobile='0123456789',
            contact_validity='valid'
        )
        res = self.client.get(reverse('hub:mca_export_excel'), {
            'scope': 'filtered',
            'f_name': 'Zero Lead Tech'
        })
        self.assertEqual(res.status_code, 200)
        wb = openpyxl.load_workbook(io.BytesIO(res.content))
        ws = wb.active
        rows = list(ws.iter_rows(values_only=False))
        self.assertEqual(len(rows), 2)
        headers = [c.value for c in rows[0]]
        cin_idx = headers.index("entity_id")
        phone_idx = headers.index("canonical_mobile")

        row = rows[1]
        self.assertEqual(row[cin_idx].value, '00123456789012345678')
        self.assertTrue(row[cin_idx].value.startswith('00'))
        self.assertEqual(row[cin_idx].number_format, '@')

        self.assertEqual(row[phone_idx].value, '0123456789')
        self.assertTrue(row[phone_idx].value.startswith('0'))
        self.assertEqual(row[phone_idx].number_format, '@')

    def test_gst_filtered_excel_export_all_when_no_filter(self):
        """When no filters are active, GST Excel export outputs all authorized unarchived records."""
        res = self.client.get(reverse('hub:gst_export_excel'), {'scope': 'filtered'})
        self.assertEqual(res.status_code, 200)
        wb = openpyxl.load_workbook(io.BytesIO(res.content))
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        # 1 header + 3 GST records
        self.assertEqual(len(rows), 4)

    def test_gst_filtered_excel_export_zero_records(self):
        """When 0 records match, returns 400 for AJAX and redirects with flash message for standard GET."""
        # 1. Standard GET
        res_standard = self.client.get(reverse('hub:gst_export_excel'), {
            'scope': 'filtered',
            'f_city': 'NonExistentCity99'
        }, follow=True)
        self.assertEqual(res_standard.status_code, 200)
        messages = list(res_standard.context.get('messages', []))
        self.assertTrue(any("No matching records to export." in str(m) for m in messages))

        # 2. AJAX GET
        res_ajax = self.client.get(reverse('hub:gst_export_excel'), {
            'scope': 'filtered',
            'f_city': 'NonExistentCity99'
        }, HTTP_X_REQUESTED_WITH='XMLHttpRequest')
        self.assertEqual(res_ajax.status_code, 400)
        import json
        data = json.loads(res_ajax.content)
        self.assertEqual(data.get('message'), "No matching records to export.")

    def test_excel_export_preserves_sorting(self):
        """Export preserves sort order (e.g. name_desc vs name_asc)."""
        res_desc = self.client.get(reverse('hub:mca_export_excel'), {
            'scope': 'filtered',
            'f_name': 'tech',
            'sort': 'name_desc'
        })
        self.assertEqual(res_desc.status_code, 200)
        wb = openpyxl.load_workbook(io.BytesIO(res_desc.content))
        rows = list(wb.active.iter_rows(values_only=True))
        headers = list(rows[0])
        name_idx = headers.index("name")
        names_desc = [r[name_idx] for r in rows[1:]]
        # In name_desc: 'Tech Solutions Global Limited' comes before 'ABC Technologies Private Limited'
        self.assertEqual(names_desc, ['Tech Solutions Global Limited', 'ABC Technologies Private Limited'])

        res_asc = self.client.get(reverse('hub:mca_export_excel'), {
            'scope': 'filtered',
            'f_name': 'tech',
            'sort': 'name_asc'
        })
        wb_asc = openpyxl.load_workbook(io.BytesIO(res_asc.content))
        rows_asc = list(wb_asc.active.iter_rows(values_only=True))
        names_asc = [r[name_idx] for r in rows_asc[1:]]
        self.assertEqual(names_asc, ['ABC Technologies Private Limited', 'Tech Solutions Global Limited'])

    def test_excel_export_across_all_pages_ignoring_page_param(self):
        """Export exports all matching records even if a page param is passed in."""
        res = self.client.get(reverse('hub:mca_export_excel'), {
            'scope': 'filtered',
            'f_name': 'tech',
            'page': '2'
        })
        self.assertEqual(res.status_code, 200)
        wb = openpyxl.load_workbook(io.BytesIO(res.content))
        rows = list(wb.active.iter_rows(values_only=True))
        # Both matching companies exported, not sliced by pagination
        self.assertEqual(len(rows), 3)  # 1 header + 2 companies


class ImportHistoryClearAllTests(TestCase):
    """
    Automated verification of Import History 'Clear All' feature:
    - Button placement beside Import History heading with matching styling
    - Exact confirmation dialog: 'Clear all import history? This will remove the history records.'
    - Inclusion of Cancel and Clear All buttons
    - Empty state: 'No import history available.'
    - Disabling Clear All when no eligible records exist or clearing is in progress
    - Clearing across pagination and filters
    - Scoping within user's authorized scope (other users' history unaffected)
    - Retention of active / queued import jobs with clear reporting of retained count
    - Preservation of files, cleaned outputs, saved mappings, and application data
    - Atomic persistence after refresh
    - Failure handling preserving displayed records
    """

    def setUp(self):
        from hub.models import MappingTemplate
        self.MappingTemplate = MappingTemplate
        self.client = Client()
        self.admin = User.objects.create_superuser('admin_user', 'admin@example.com', 'AdminPass123!')
        self.user1 = User.objects.create_user('alice', 'alice@example.com', 'AlicePass123!', is_staff=True, is_superuser=False)
        self.user2 = User.objects.create_user('bob', 'bob@example.com', 'BobPass123!', is_staff=True, is_superuser=False)

    def test_ui_presence_confirmation_and_styling(self):
        """Clear All button is beside Import History heading, matches UI, and modal has exact confirmation text."""
        self.client.login(username='admin_user', password='AdminPass123!')
        ImportBatch.objects.create(
            batch_id='BATCH_TEST_001',
            dataset_type='MCA',
            file_name='test.xlsx',
            file_path='/tmp/test.xlsx',
            file_checksum='abc1',
            file_format='XLSX',
            status='COMPLETED',
            created_by=self.admin
        )

        res = self.client.get(reverse('hub:import_history'))
        self.assertEqual(res.status_code, 200)
        html = res.content.decode('utf-8')

        # Heading and button presence
        self.assertIn('Import History', html)
        self.assertIn('id="clearAllHistoryBtn"', html)
        self.assertIn('Clear All', html)
        self.assertNotIn('disabled aria-disabled="true"', html)

        # Modal and exact confirmation prompt
        self.assertIn('id="clearHistoryModal"', html)
        self.assertIn('Clear all import history? This will remove the history records.', html)
        self.assertIn('id="cancelClearHistoryBtn"', html)
        self.assertIn('Cancel', html)
        self.assertIn('id="confirmClearHistoryBtn"', html)

    def test_ui_disabled_and_empty_state_when_no_records(self):
        """Button is disabled and empty state 'No import history available.' is shown when no records exist."""
        self.client.login(username='admin_user', password='AdminPass123!')
        res = self.client.get(reverse('hub:import_history'))
        self.assertEqual(res.status_code, 200)
        html = res.content.decode('utf-8')

        self.assertIn('No import history available.', html)
        self.assertIn('disabled', html)
        self.assertIn('id="clearAllHistoryBtn"', html)

    def test_ui_disabled_when_only_active_jobs_exist(self):
        """When only active/queued jobs exist, eligible count is 0 so Clear All is disabled."""
        self.client.login(username='admin_user', password='AdminPass123!')
        ImportBatch.objects.create(
            batch_id='BATCH_ACTIVE_001',
            dataset_type='MCA',
            file_name='active.xlsx',
            file_path='/tmp/active.xlsx',
            file_checksum='act1',
            file_format='XLSX',
            status='IMPORTING',
            created_by=self.admin
        )

        res = self.client.get(reverse('hub:import_history'))
        html = res.content.decode('utf-8')
        # Button should be disabled since eligible_count is 0
        self.assertIn('disabled', html)

        # Attempting to POST clear-all should retain the active job and report it
        post_res = self.client.post(reverse('hub:import_history_clear_all'), follow=True)
        self.assertEqual(post_res.status_code, 200)
        self.assertTrue(ImportBatch.objects.filter(batch_id='BATCH_ACTIVE_001').exists())
        messages_text = [m.message for m in post_res.context['messages']]
        self.assertTrue(any('active or queued' in m for m in messages_text))

    def test_clear_all_across_pagination_and_filters_persists_after_refresh(self):
        """Clearing clears all eligible records across pagination and filters, persisting on refresh."""
        self.client.login(username='admin_user', password='AdminPass123!')

        # Create 20 batches across MCA and GST (page size is 15)
        for i in range(12):
            ImportBatch.objects.create(
                batch_id=f'BATCH_MCA_{i:03d}',
                dataset_type='MCA',
                file_name=f'mca_{i}.csv',
                file_path=f'/tmp/mca_{i}.csv',
                file_checksum=f'cs_mca_{i}',
                file_format='CSV',
                status='COMPLETED',
                created_by=self.admin
            )
        for i in range(8):
            ImportBatch.objects.create(
                batch_id=f'BATCH_GST_{i:03d}',
                dataset_type='GST',
                file_name=f'gst_{i}.csv',
                file_path=f'/tmp/gst_{i}.csv',
                file_checksum=f'cs_gst_{i}',
                file_format='CSV',
                status='FAILED',
                created_by=self.admin
            )

        self.assertEqual(ImportBatch.objects.count(), 20)

        # Clear while filtered or paginated
        res = self.client.post(reverse('hub:import_history_clear_all'), follow=True)
        self.assertEqual(res.status_code, 200)
        self.assertEqual(ImportBatch.objects.count(), 0)

        # Refresh / subsequent GET
        refresh_res = self.client.get(reverse('hub:import_history'))
        self.assertEqual(refresh_res.status_code, 200)
        self.assertContains(refresh_res, 'No import history available.')

    def test_active_and_queued_jobs_are_retained_and_reported(self):
        """Active and queued jobs are never deleted or cancelled, and retained count is clearly reported."""
        self.client.login(username='admin_user', password='AdminPass123!')

        # 2 completed, 1 importing, 1 queued
        ImportBatch.objects.create(
            batch_id='BATCH_DONE_1', dataset_type='MCA', file_name='f1.csv',
            file_path='/tmp/f1.csv', file_checksum='c1', file_format='CSV',
            status='COMPLETED', created_by=self.admin
        )
        ImportBatch.objects.create(
            batch_id='BATCH_DONE_2', dataset_type='GST', file_name='f2.csv',
            file_path='/tmp/f2.csv', file_checksum='c2', file_format='CSV',
            status='FAILED', created_by=self.admin
        )
        ImportBatch.objects.create(
            batch_id='BATCH_RUNNING_1', dataset_type='MCA', file_name='f3.csv',
            file_path='/tmp/f3.csv', file_checksum='c3', file_format='CSV',
            status='IMPORTING', created_by=self.admin
        )
        ImportBatch.objects.create(
            batch_id='BATCH_QUEUED_1', dataset_type='GST', file_name='f4.csv',
            file_path='/tmp/f4.csv', file_checksum='c4', file_format='CSV',
            status='QUEUED', created_by=self.admin
        )

        res = self.client.post(reverse('hub:import_history_clear_all'), follow=True)
        self.assertEqual(res.status_code, 200)

        # Check DB
        self.assertFalse(ImportBatch.objects.filter(batch_id__in=['BATCH_DONE_1', 'BATCH_DONE_2']).exists())
        self.assertTrue(ImportBatch.objects.filter(batch_id='BATCH_RUNNING_1').exists())
        self.assertTrue(ImportBatch.objects.filter(batch_id='BATCH_QUEUED_1').exists())

        # Feedback report
        messages_text = [m.message for m in res.context['messages']]
        self.assertTrue(any('cleared 2' in m.lower() and 'retained 2 active' in m.lower() for m in messages_text))

    def test_user_scope_isolation_does_not_affect_other_users(self):
        """Standard user clearing history does not affect another user's import history."""
        # Alice's batch
        b_alice = ImportBatch.objects.create(
            batch_id='BATCH_ALICE', dataset_type='MCA', file_name='alice.csv',
            file_path='/tmp/alice.csv', file_checksum='calice', file_format='CSV',
            status='COMPLETED', created_by=self.user1
        )
        # Bob's batch
        b_bob = ImportBatch.objects.create(
            batch_id='BATCH_BOB', dataset_type='GST', file_name='bob.csv',
            file_path='/tmp/bob.csv', file_checksum='cbob', file_format='CSV',
            status='COMPLETED', created_by=self.user2
        )

        # Login as Alice and clear
        self.client.login(username='alice', password='AlicePass123!')
        res = self.client.post(reverse('hub:import_history_clear_all'), follow=True)
        self.assertEqual(res.status_code, 200)

        # Alice's batch is cleared
        self.assertFalse(ImportBatch.objects.filter(batch_id='BATCH_ALICE').exists())
        # Bob's batch is preserved
        self.assertTrue(ImportBatch.objects.filter(batch_id='BATCH_BOB').exists())

    def test_preserves_uploaded_files_and_other_application_data(self):
        """Clearing import history does NOT delete files on disk, mappings, companies, GST, or quality issues."""
        self.client.login(username='admin_user', password='AdminPass123!')

        # Create temporary source file
        with tempfile.NamedTemporaryFile(delete=False, suffix='.csv') as tmp:
            tmp.write(b"col1,col2\nval1,val2")
            tmp_path = tmp.name

        try:
            batch = ImportBatch.objects.create(
                batch_id='BATCH_PRESERVE_TEST', dataset_type='MCA', file_name='dummy.csv',
                file_path=tmp_path, file_checksum='cdum', file_format='CSV',
                status='COMPLETED', created_by=self.admin
            )
            # Create mapping template
            template = self.MappingTemplate.objects.create(
                name='MCA Preserve Test Template',
                dataset_type='MCA',
                mapping_config={'col1': 'name', 'col2': 'cin'},
                created_by=self.admin
            )
            # Create MCA company
            company = MCACompany.objects.create(
                entity_id='U12345DL2020PTC000001',
                name='Preserved Test Corp',
                entity_type='company'
            )
            # Create GST record
            gst = GSTRecord.objects.create(
                gst_number='07AAAAA0000A1Z5',
                legal_name='Preserved GST Corp',
                trade_name='Preserved GST',
                status='ACT'
            )
            # Create Quality issue
            quality = QualityIssue.objects.create(
                dataset_type='MCA',
                batch=batch,
                record_type='COMPANY',
                entity_identifier='U12345DL2020PTC000001',
                field_name='phone',
                raw_value='123',
                error_code='INVALID_PHONE',
                reason_description='Preserve test'
            )

            # Clear import history
            res = self.client.post(reverse('hub:import_history_clear_all'), follow=True)
            self.assertEqual(res.status_code, 200)

            # Verify history record removed
            self.assertFalse(ImportBatch.objects.filter(batch_id='BATCH_PRESERVE_TEST').exists())

            # Verify file on disk still exists
            self.assertTrue(os.path.exists(tmp_path))

            # Verify other application data intact
            self.assertTrue(self.MappingTemplate.objects.filter(id=template.id).exists())
            self.assertTrue(MCACompany.objects.filter(id=company.id).exists())
            self.assertTrue(GSTRecord.objects.filter(id=gst.id).exists())
            self.assertTrue(QualityIssue.objects.filter(id=quality.id).exists())

        finally:
            if os.path.exists(tmp_path):
                os.remove(tmp_path)

    def test_failure_preserves_records_and_reports_error(self):
        """When an unexpected error occurs during clear-all, records are preserved and error is reported."""
        self.client.login(username='admin_user', password='AdminPass123!')
        ImportBatch.objects.create(
            batch_id='BATCH_FAIL_TEST', dataset_type='MCA', file_name='fail.csv',
            file_path='/tmp/fail.csv', file_checksum='cfail', file_format='CSV',
            status='COMPLETED', created_by=self.admin
        )

        from django.db.models.query import QuerySet
        with patch.object(QuerySet, 'delete', side_effect=RuntimeError("Simulated Failure")):
            res = self.client.post(reverse('hub:import_history_clear_all'), follow=True)
            self.assertEqual(res.status_code, 200)

        # Batch preserved
        self.assertTrue(ImportBatch.objects.filter(batch_id='BATCH_FAIL_TEST').exists())

    def test_get_method_rejected(self):
        """Clear All only accepts POST requests."""
        self.client.login(username='admin_user', password='AdminPass123!')
        res = self.client.get(reverse('hub:import_history_clear_all'), follow=True)
        self.assertEqual(res.status_code, 200)
        messages_text = [m.message for m in res.context['messages']]
        self.assertTrue(any('Invalid request method' in m for m in messages_text))


class ImportLiveProgressTrackingTests(TestCase):
    """
    Test suite for live progress tracking during upload, validation, and cleaning:
    - Clear processing stages (Uploading, Reading headers, Waiting for mapping, Queued, Validating and cleaning, Preparing output, Completed/Failed/Cancelled)
    - Real metrics (bytes, rows, validation failures, elapsed time, real ETA)
    - Bounded-memory file counting
    - Multi-job isolation and concurrency
    - Refresh persistence from persistent storage
    - Safety of active jobs during history clearing
    """

    def setUp(self):
        self.user = User.objects.create_user(username='progresstester', password='TestPassword123!')
        self.admin = User.objects.create_superuser(username='progressadmin', password='AdminPassword123!')
        self.client = Client()

    def test_bounded_row_counting_csv(self):
        """Bounded line counting counts rows correctly without reading full file into memory."""
        from hub.progress import count_file_rows_bounded
        import tempfile

        with tempfile.NamedTemporaryFile(mode='w+', delete=False, suffix='.csv') as f:
            f.write("CIN,Name,ROC\n")
            for i in range(150):
                f.write(f"U12345DL2020PTC{i:06d},Company {i},ROC Delhi\n")
            f_path = f.name

        try:
            count = count_file_rows_bounded(f_path, file_format='CSV', header_row=0)
            self.assertEqual(count, 150)
        finally:
            if os.path.exists(f_path):
                os.remove(f_path)

    def test_bounded_row_counting_missing_file(self):
        """Returns None safely for non-existent file path."""
        from hub.progress import count_file_rows_bounded
        self.assertIsNone(count_file_rows_bounded('/non/existent/path.csv'))

    def test_progress_stages_and_real_percentages(self):
        """Tracks exact stages and calculates percentage strictly from real row counts."""
        from hub.progress import init_progress, update_progress, get_progress

        batch_id = 'BATCH-TEST-STAGE-01'
        # Stage 1: Queued
        init_progress(batch_id=batch_id, dataset_type='MCA', stage='Queued for processing', total_rows=500)
        p = get_progress(batch_id)
        self.assertEqual(p['stage'], 'Queued for processing')
        self.assertEqual(p['total_rows'], 500)
        self.assertEqual(p['processed_rows'], 0)
        self.assertEqual(p['percentage'], 0.0)

        # Stage 2: Validating and cleaning with partial progress
        update_progress(
            batch_id=batch_id,
            stage='Validating and cleaning',
            status='PROCESSING',
            processed_rows=325,
            total_rows=500,
            validation_failures=12
        )
        p = get_progress(batch_id)
        self.assertEqual(p['stage'], 'Validating and cleaning')
        self.assertEqual(p['processed_rows'], 325)
        self.assertEqual(p['percentage'], 65.0)
        self.assertEqual(p['validation_failures'], 12)
        self.assertIn("65%", p['message'])
        self.assertIn("325 of 500 rows processed", p['message'])

        # Total rows unknown: percentage must be None (no fake fabricated percentage)
        batch_id_unknown = 'BATCH-TEST-UNKNOWN'
        init_progress(batch_id=batch_id_unknown, dataset_type='GST', stage='Validating and cleaning', total_rows=None)
        update_progress(batch_id=batch_id_unknown, processed_rows=42)
        p_unk = get_progress(batch_id_unknown)
        self.assertIsNone(p_unk['percentage'])
        self.assertEqual(p_unk['processed_rows'], 42)
        self.assertIn("42 rows processed", p_unk['message'])

    def test_duration_and_eta_calculation(self):
        """ETA is calculated only when sufficient rows processed and time elapsed."""
        import time
        from hub.progress import init_progress, update_progress, get_progress, _PROGRESS_REGISTRY, _PROGRESS_LOCK

        batch_id = 'BATCH-ETA-TEST'
        init_progress(batch_id=batch_id, total_rows=1000)
        # Mock start time to 10 seconds ago
        with _PROGRESS_LOCK:
            _PROGRESS_REGISTRY[batch_id]['start_time'] = time.time() - 10.0

        update_progress(batch_id=batch_id, processed_rows=200, total_rows=1000)
        p = get_progress(batch_id)
        self.assertIsNotNone(p['eta_seconds'])
        self.assertTrue(len(p['eta_formatted']) > 0)
        self.assertTrue(p['elapsed_seconds'] >= 9.9)

    def test_progress_persistence_and_refresh_restoration(self):
        """When in-memory registry is cleared (simulating refresh/restart), state is restored from DB."""
        from hub.progress import init_progress, update_progress, get_progress, _PROGRESS_REGISTRY, _PROGRESS_LOCK

        batch = ImportBatch.objects.create(
            batch_id='BATCH-REFRESH-TEST',
            dataset_type='MCA',
            file_name='test.csv',
            file_path='/tmp/test.csv',
            file_checksum='c1234567',
            file_format='CSV',
            status='PROCESSING',
            total_input_rows=1000,
            created_by=self.user
        )

        init_progress(batch_id=batch.batch_id, dataset_type='MCA', stage='Validating and cleaning', total_rows=1000)
        update_progress(
            batch_id=batch.batch_id,
            stage='Validating and cleaning',
            status='PROCESSING',
            processed_rows=450,
            total_rows=1000,
            validation_failures=18,
            force_db_save=True
        )

        # Clear in-memory cache to simulate page refresh / new worker
        with _PROGRESS_LOCK:
            _PROGRESS_REGISTRY.pop(batch.batch_id, None)

        # get_progress must reconstruct from DB snapshot
        restored = get_progress(batch.batch_id)
        self.assertEqual(restored['batch_id'], batch.batch_id)
        self.assertEqual(restored['stage'], 'Validating and cleaning')
        self.assertEqual(restored['processed_rows'], 450)
        self.assertEqual(restored['total_rows'], 1000)
        self.assertEqual(restored['percentage'], 45.0)
        self.assertEqual(restored['validation_failures'], 18)

    def test_concurrent_import_jobs_isolation(self):
        """Multiple imports running at the same time have separate, isolated progress states."""
        from hub.progress import init_progress, update_progress, get_progress

        batch1 = 'BATCH-CONCURRENT-01'
        batch2 = 'BATCH-CONCURRENT-02'

        init_progress(batch_id=batch1, dataset_type='MCA', stage='Validating and cleaning', total_rows=500)
        init_progress(batch_id=batch2, dataset_type='GST', stage='Validating and cleaning', total_rows=800)

        update_progress(batch1, processed_rows=100, validation_failures=2)
        update_progress(batch2, processed_rows=400, validation_failures=15)

        p1 = get_progress(batch1)
        p2 = get_progress(batch2)

        self.assertEqual(p1['processed_rows'], 100)
        self.assertEqual(p1['percentage'], 20.0)
        self.assertEqual(p1['validation_failures'], 2)

        self.assertEqual(p2['processed_rows'], 400)
        self.assertEqual(p2['percentage'], 50.0)
        self.assertEqual(p2['validation_failures'], 15)

    def test_progress_cancellation(self):
        """Cancelling a job marks status as CANCELLED and flags cancel token."""
        from hub.progress import init_progress, request_cancel, is_cancelled, get_progress

        batch = ImportBatch.objects.create(
            batch_id='BATCH-CANCEL-TEST',
            dataset_type='GST',
            file_name='gst.csv',
            file_path='/tmp/gst.csv',
            file_checksum='c999',
            file_format='CSV',
            status='PROCESSING',
            created_by=self.user
        )

        init_progress(batch_id=batch.batch_id, stage='Validating and cleaning')
        self.assertFalse(is_cancelled(batch.batch_id))

        request_cancel(batch.batch_id)
        self.assertTrue(is_cancelled(batch.batch_id))

        p = get_progress(batch.batch_id)
        self.assertTrue(p['is_cancelled'])
        self.assertEqual(p['status'], 'CANCELLED')

    def test_active_progress_api_view_and_clear_history_safety(self):
        """Active jobs are reported in active-progress API and never deleted by history clear-all."""
        from hub.progress import init_progress

        active_batch = ImportBatch.objects.create(
            batch_id='BATCH-ACTIVE-STAY',
            dataset_type='MCA',
            file_name='active.csv',
            file_path='/tmp/active.csv',
            file_checksum='cact',
            file_format='CSV',
            status='PROCESSING',
            created_by=self.admin
        )
        completed_batch = ImportBatch.objects.create(
            batch_id='BATCH-COMPLETED-GONE',
            dataset_type='MCA',
            file_name='done.csv',
            file_path='/tmp/done.csv',
            file_checksum='cdone',
            file_format='CSV',
            status='COMPLETED',
            created_by=self.admin
        )

        init_progress(batch_id=active_batch.batch_id, dataset_type='MCA', stage='Validating and cleaning', total_rows=100)

        self.client.login(username='progressadmin', password='AdminPassword123!')

        # Active progress API returns active batch
        res = self.client.get(reverse('hub:import_active_progress_api'))
        self.assertEqual(res.status_code, 200)
        data = res.json()
        self.assertIn('active_batches', data)
        self.assertIn(active_batch.batch_id, data['active_batches'])

        # Clear history must preserve active batch
        res_clear = self.client.post(reverse('hub:import_history_clear_all'))
        self.assertEqual(res_clear.status_code, 302)

        self.assertTrue(ImportBatch.objects.filter(batch_id=active_batch.batch_id).exists())
        self.assertFalse(ImportBatch.objects.filter(batch_id=completed_batch.batch_id).exists())


class ColumnSelectionAndExclusionTestCase(TestCase):
    """
    Tests for column-selection checkboxes in MCA/GST mapping:
    - Include checkboxes and Select All / Clear Selection controls.
    - Omission of excluded columns from parsing, cleaning, validation, staging, DB writes, and exports.
    - Required field retention and dependency validation.
    - Safe database updates preserving omitted fields.
    - Export column customization.
    """
    def setUp(self):
        self.client = Client()
        self.user = User.objects.create_superuser('coladmin', 'col@example.com', 'AdminPassword123!')
        self.client.login(username='coladmin', password='AdminPassword123!')
        self.temp_files = []

    def tearDown(self):
        for f in self.temp_files:
            if os.path.exists(f):
                try:
                    os.remove(f)
                except OSError:
                    pass

    def test_mapping_view_renders_column_items_and_selection_defaults(self):
        """Mapping view renders column items, defaults mapped targets to checked, and skip to unchecked."""
        csv_file = tempfile.NamedTemporaryFile(suffix=".csv", mode='w', newline='', encoding='utf-8', delete=False)
        writer = csv.writer(csv_file)
        writer.writerow(["Sr. No.", "CIN", "Company Name", "Extra Remarks"])
        writer.writerow(["1", "U12345MH2020PTC123456", "ALPHA TECH PVT LTD", "Some comment"])
        csv_file.close()
        self.temp_files.append(csv_file.name)

        batch = ImportBatch.objects.create(
            batch_id="TEST-COL-MAP-01",
            dataset_type="MCA",
            file_name="test_col_map.csv",
            file_path=csv_file.name,
            file_checksum="chk_col_map_01",
            file_format="CSV",
            status="UPLOADED",
            created_by=self.user,
            mapping_config={'delimiter': ',', 'encoding': 'utf-8'}
        )

        res = self.client.get(reverse('hub:import_map', args=[batch.batch_id]))
        self.assertEqual(res.status_code, 200)
        self.assertIn('column_items', res.context)
        col_items = res.context['column_items']

        # Check items: CIN is required and selected; Sr. No. is skip and initially unselected
        item_by_header = {c['header']: c for c in col_items}
        self.assertIn('CIN', item_by_header)
        self.assertTrue(item_by_header['CIN']['is_required'])
        self.assertTrue(item_by_header['CIN']['is_selected'])

        self.assertIn('Sr. No.', item_by_header)
        self.assertTrue(item_by_header['Sr. No.']['is_skip'])
        self.assertFalse(item_by_header['Sr. No.']['is_selected'])

        content = res.content.decode('utf-8')
        self.assertIn("Include", content)
        self.assertIn("Select All", content)
        self.assertIn("Clear Selection", content)
        self.assertIn("columns selected", content)
        self.assertIn("col_0", content)

    def test_post_mapping_persists_selected_and_excluded_columns(self):
        """Submitting mapping saves selected_columns and excluded_columns in mapping_config."""
        csv_file = tempfile.NamedTemporaryFile(suffix=".csv", mode='w', newline='', encoding='utf-8', delete=False)
        writer = csv.writer(csv_file)
        writer.writerow(["CIN", "Company Name", "State", "City"])
        writer.writerow(["U12345MH2020PTC123456", "ALPHA TECH PVT LTD", "Maharashtra", "Mumbai"])
        csv_file.close()
        self.temp_files.append(csv_file.name)

        batch = ImportBatch.objects.create(
            batch_id="TEST-COL-MAP-02",
            dataset_type="MCA",
            file_name="test_col_map2.csv",
            file_path=csv_file.name,
            file_checksum="chk_col_map_02",
            file_format="CSV",
            status="UPLOADED",
            created_by=self.user,
            mapping_config={'delimiter': ',', 'encoding': 'utf-8'}
        )

        # Submit only CIN and Company Name; State and City are excluded
        post_data = {
            'selected_columns': ['CIN', 'Company Name'],
            'map_CIN': 'entityId',
            'map_Company Name': 'name',
            'map_State': 'state',
            'map_City': 'city',
            'date_format': '%m/%d/%Y'
        }
        res = self.client.post(reverse('hub:import_map', args=[batch.batch_id]), post_data, follow=True)
        self.assertEqual(res.status_code, 200)

        batch.refresh_from_db()
        self.assertEqual(batch.mapping_config.get('selected_columns'), ['CIN', 'Company Name'])
        self.assertEqual(batch.mapping_config.get('excluded_columns'), ['State', 'City'])

    def test_pipeline_completely_excludes_unchecked_columns(self):
        """Unchecked columns are not cleaned, validated, in extra_data, or in StagedRow raw_data."""
        csv_file = tempfile.NamedTemporaryFile(suffix=".csv", mode='w', newline='', encoding='utf-8', delete=False)
        writer = csv.writer(csv_file)
        # Note: Mobile column has an invalid phone number.
        writer.writerow(["CIN", "Company Name", "Director Mobile", "Secret Notes"])
        writer.writerow(["U12345MH2020PTC123456", "BETA SYSTEMS PVT LTD", "INVALID-PHONE-999", "Sensitive Info"])
        csv_file.close()
        self.temp_files.append(csv_file.name)

        # Batch where Director Mobile and Secret Notes are NOT selected
        batch = ImportBatch.objects.create(
            batch_id="TEST-COL-PIPE-01",
            dataset_type="MCA",
            file_name="test_pipe.csv",
            file_path=csv_file.name,
            file_checksum="chk_col_pipe_01",
            file_format="CSV",
            status="UPLOADED",
            created_by=self.user,
            mapping_config={
                'delimiter': ',',
                'encoding': 'utf-8',
                'columns': {
                    'CIN': 'entityId',
                    'Company Name': 'name',
                    'Director Mobile': 'directorMobile',
                    'Secret Notes': ''
                },
                'selected_columns': ['CIN', 'Company Name'],
                'excluded_columns': ['Director Mobile', 'Secret Notes'],
                'date_format': '%m/%d/%Y'
            }
        )

        stage_res = stage_and_validate_batch(batch)
        self.assertEqual(stage_res['total_rows'], 1)

        staged = StagedRow.objects.get(batch=batch, row_number=1)
        # Excluded columns must not be in raw_data, cleaned_data, or extra_data
        self.assertIn('CIN', staged.raw_data)
        self.assertIn('Company Name', staged.raw_data)
        self.assertNotIn('Director Mobile', staged.raw_data)
        self.assertNotIn('Secret Notes', staged.raw_data)

        self.assertNotIn('Secret Notes', staged.cleaned_data.get('extra_data', {}))
        self.assertNotIn('directorMobile', staged.cleaned_data)
        self.assertEqual(staged.canonical_mobile, '')

        # No QualityIssue logged because Director Mobile was excluded!
        self.assertEqual(QualityIssue.objects.filter(batch=batch).count(), 0)

        # Publish to database
        execute_import_publication(batch, user=self.user)
        company = MCACompany.objects.get(entity_id="U12345MH2020PTC123456")
        self.assertEqual(company.name, "BETA SYSTEMS PVT LTD")
        # No director entry created
        self.assertEqual(MCADirectorEntry.objects.filter(company=company).count(), 0)
        # Extra data does not contain excluded column
        self.assertNotIn('Secret Notes', company.extra_data)

    def test_safe_database_update_preserves_omitted_fields(self):
        """Updating existing record when fields are omitted leaves existing DB values untouched."""
        existing = MCACompany.objects.create(
            entity_type='company',
            entity_id='U99999MH2021PTC999999',
            name='EXISTING OLD NAME',
            state='Maharashtra',
            city='Pune',
            paid_up_capital=750000.00
        )

        csv_file = tempfile.NamedTemporaryFile(suffix=".csv", mode='w', newline='', encoding='utf-8', delete=False)
        writer = csv.writer(csv_file)
        # Import file only contains CIN and updated Name
        writer.writerow(["CIN", "Company Name"])
        writer.writerow(["U99999MH2021PTC999999", "NEW UPDATED NAME PVT LTD"])
        csv_file.close()
        self.temp_files.append(csv_file.name)

        batch = ImportBatch.objects.create(
            batch_id="TEST-SAFE-UPDATE",
            dataset_type="MCA",
            file_name="safe_update.csv",
            file_path=csv_file.name,
            file_checksum="chk_safe_update",
            file_format="CSV",
            status="UPLOADED",
            created_by=self.user,
            mapping_config={
                'columns': {'CIN': 'entityId', 'Company Name': 'name'},
                'selected_columns': ['CIN', 'Company Name'],
                'date_format': '%m/%d/%Y'
            }
        )

        stage_and_validate_batch(batch)
        execute_import_publication(batch, user=self.user)

        existing.refresh_from_db()
        # Name was updated
        self.assertEqual(existing.name, "NEW UPDATED NAME PVT LTD")
        # Omitted fields remain completely preserved!
        self.assertEqual(existing.state, "Maharashtra")
        self.assertEqual(existing.city, "Pune")
        self.assertEqual(float(existing.paid_up_capital), 750000.00)

    def test_export_selection_omits_excluded_fields(self):
        """Exports omit columns that were excluded from import and support custom column selection."""
        csv_file = tempfile.NamedTemporaryFile(suffix=".csv", mode='w', newline='', encoding='utf-8', delete=False)
        writer = csv.writer(csv_file)
        writer.writerow(["CIN", "Company Name", "State"])
        writer.writerow(["U88888MH2022PTC888888", "DELTA TECH CORP", "Karnataka"])
        csv_file.close()
        self.temp_files.append(csv_file.name)

        # Batch where State was excluded
        batch = ImportBatch.objects.create(
            batch_id="TEST-EXP-OMIT",
            dataset_type="MCA",
            file_name="exp_omit.csv",
            file_path=csv_file.name,
            file_checksum="chk_exp_omit",
            file_format="CSV",
            status="UPLOADED",
            created_by=self.user,
            mapping_config={
                'columns': {'CIN': 'entityId', 'Company Name': 'name', 'State': 'state'},
                'selected_columns': ['CIN', 'Company Name'],
                'excluded_columns': ['State'],
                'date_format': '%m/%d/%Y'
            }
        )

        stage_and_validate_batch(batch)
        export_path = export_cleaned_master_data(batch, file_format='CSV')
        self.temp_files.append(export_path)

        with open(export_path, 'r', encoding='utf-8') as f:
            reader = csv.reader(f)
            header_row = next(reader)
            # Entity ID and Company Name are present
            self.assertIn("Entity ID", header_row)
            self.assertIn("Company Name", header_row)
            # Excluded 'State' is NOT in export headers!
            self.assertNotIn("State", header_row)

        # Custom export selection test
        custom_export_path = export_cleaned_master_data(batch, file_format='CSV', columns=['Entity ID'])
        self.temp_files.append(custom_export_path)
        with open(custom_export_path, 'r', encoding='utf-8') as f:
            reader = csv.reader(f)
            custom_headers = next(reader)
            self.assertEqual(custom_headers, ["Entity ID"])
class ImportPreviewRegressionTestCase(TestCase):
    """
    Regression test suite for Django import-preview rendering and optional columns support.
    Verifies:
    1. Neither directorEmail nor email present (exact bug condition from screenshot).
    2. Only directorEmail present.
    3. Only email present (labeled as company/general email).
    4. Both present (both rendered, company labeled).
    5. Empty-string and None values for contact fields.
    6. Company-only MCA data with no director fields.
    7. GST data with optional contact fields excluded.
    8. Fully mapped record.
    9. Other missing optional fields (canonical_mobile, phone_validity, etc.).
    10. Excluded fields are not reintroduced into publishing payloads and existing DB values preserved.
    """
    def setUp(self):
        self.client = Client()
        self.user = User.objects.create_superuser('prevadmin', 'prev@example.com', 'AdminPass123!')
        self.client.login(username='prevadmin', password='AdminPass123!')

    def _create_mca_batch(self, batch_id='BATCH-REG-MCA-01'):
        return ImportBatch.objects.create(
            batch_id=batch_id,
            dataset_type='MCA',
            file_name='test_mca.xlsx',
            file_path='/tmp/test_mca.xlsx',
            file_checksum='chk_mca_reg_01',
            file_format='XLSX',
            status='STAGED',
            created_by=self.user,
            total_input_rows=1,
            proposed_new_records=1,
            mapping_config={'selected_columns': ['Entity ID', 'Company Name'], 'columns': {'CIN': 'entityId', 'Company': 'name'}}
        )

    def _create_gst_batch(self, batch_id='BATCH-REG-GST-01'):
        return ImportBatch.objects.create(
            batch_id=batch_id,
            dataset_type='GST',
            file_name='test_gst.xlsx',
            file_path='/tmp/test_gst.xlsx',
            file_checksum='chk_gst_reg_01',
            file_format='XLSX',
            status='STAGED',
            created_by=self.user,
            total_input_rows=1,
            proposed_new_records=1,
            mapping_config={'selected_columns': ['GSTIN', 'Trade Name'], 'columns': {'GSTIN': 'gst_number', 'Trade': 'trade_name'}}
        )

    def test_neither_director_email_nor_company_email_renders_200(self):
        """Case 1: Neither directorEmail nor email key present in cleaned_data (screenshot condition)."""
        batch = self._create_mca_batch('BATCH-REG-NO-EMAIL')
        StagedRow.objects.create(
            batch=batch,
            row_number=1,
            entity_identifier='U46620DL2025PTC440609',
            cleaned_data={
                'name': 'YOGSHA TRADEX (I) PRIVATE LIMITED',
                'entityId': 'U46620DL2025PTC440609',
                'entityType': 'company',
                'extra_data': {},
                'source_row': 1,
                'source_file': 'Data1.xlsx',
                'company_name': 'YOGSHA TRADEX (I) PRIVATE LIMITED',
                'source_sheet': 'Sheet1'
            },
            proposed_action='CREATE',
            validation_status='VALID'
        )
        url = reverse('hub:import_preview', args=[batch.batch_id])
        res = self.client.get(url)
        self.assertEqual(res.status_code, 200)
        self.assertContains(res, 'YOGSHA TRADEX (I) PRIVATE LIMITED')
        self.assertContains(res, 'U46620DL2025PTC440609')
        # Director name displays '-'
        self.assertContains(res, '<div>-</div>')

    def test_only_director_email_present(self):
        """Case 2: Only directorEmail present; should render director email without company label."""
        batch = self._create_mca_batch('BATCH-REG-DIR-EMAIL')
        StagedRow.objects.create(
            batch=batch,
            row_number=1,
            entity_identifier='U12345MH2020PTC123456',
            cleaned_data={
                'name': 'ALPHA SOLUTIONS PVT LTD',
                'entityId': 'U12345MH2020PTC123456',
                'entityType': 'company',
                'directorName': 'Jane Director',
                'directorEmail': 'jane.director@alpha.com'
            },
            proposed_action='CREATE',
            validation_status='VALID'
        )
        url = reverse('hub:import_preview', args=[batch.batch_id])
        res = self.client.get(url)
        self.assertEqual(res.status_code, 200)
        self.assertContains(res, 'Jane Director')
        self.assertContains(res, 'jane.director@alpha.com')
        # Does NOT label director email as (Company)
        self.assertNotContains(res, '(Company)')

    def test_only_company_email_present_labeled_as_company(self):
        """Case 3: Only company email present; should render labeled as (Company) email."""
        batch = self._create_mca_batch('BATCH-REG-COMP-EMAIL')
        StagedRow.objects.create(
            batch=batch,
            row_number=1,
            entity_identifier='U12345MH2020PTC123456',
            cleaned_data={
                'name': 'BETA ENTERPRISES PVT LTD',
                'entityId': 'U12345MH2020PTC123456',
                'entityType': 'company',
                'email': 'info@betaenterprises.com'
            },
            proposed_action='CREATE',
            validation_status='VALID'
        )
        url = reverse('hub:import_preview', args=[batch.batch_id])
        res = self.client.get(url)
        self.assertEqual(res.status_code, 200)
        self.assertContains(res, 'info@betaenterprises.com')
        self.assertContains(res, '(Company)')

    def test_both_emails_present_distinguished(self):
        """Case 4: Both directorEmail and company email present; both rendered and distinguished."""
        batch = self._create_mca_batch('BATCH-REG-BOTH-EMAILS')
        StagedRow.objects.create(
            batch=batch,
            row_number=1,
            entity_identifier='U12345MH2020PTC123456',
            cleaned_data={
                'name': 'GAMMA SERVICES PVT LTD',
                'entityId': 'U12345MH2020PTC123456',
                'entityType': 'company',
                'directorName': 'Bob Gamma',
                'directorEmail': 'bob@gammaservices.com',
                'email': 'contact@gammaservices.com'
            },
            proposed_action='CREATE',
            validation_status='VALID'
        )
        url = reverse('hub:import_preview', args=[batch.batch_id])
        res = self.client.get(url)
        self.assertEqual(res.status_code, 200)
        self.assertContains(res, 'bob@gammaservices.com')
        self.assertContains(res, 'contact@gammaservices.com')
        self.assertContains(res, '(Company)')

    def test_empty_string_and_none_values_do_not_crash(self):
        """Case 5: Empty-string and None values for emails and phone fields."""
        batch = self._create_mca_batch('BATCH-REG-NONE-EMPTY')
        StagedRow.objects.create(
            batch=batch,
            row_number=1,
            entity_identifier='U12345MH2020PTC123456',
            cleaned_data={
                'name': 'DELTA LOGISTICS PVT LTD',
                'entityId': 'U12345MH2020PTC123456',
                'entityType': 'company',
                'directorName': None,
                'directorEmail': '',
                'email': None,
                'original_phone': '',
                'directorMobile': None,
                'cleaned_phone': None,
                'canonical_mobile': '',
                'phone_status': None,
                'phone_validity': ''
            },
            proposed_action='CREATE',
            validation_status='VALID'
        )
        url = reverse('hub:import_preview', args=[batch.batch_id])
        res = self.client.get(url)
        self.assertEqual(res.status_code, 200)
        self.assertContains(res, 'DELTA LOGISTICS PVT LTD')
        self.assertContains(res, 'null / blank')
        self.assertContains(res, 'MISSING')

    def test_company_only_mca_no_director_fields(self):
        """Case 6: Company-only MCA data with no director fields."""
        batch = self._create_mca_batch('BATCH-REG-COMP-ONLY')
        StagedRow.objects.create(
            batch=batch,
            row_number=1,
            entity_identifier='U99999DL2024PTC111111',
            cleaned_data={
                'name': 'SOLO HOLDINGS PVT LTD',
                'entityId': 'U99999DL2024PTC111111',
                'entityType': 'company',
                'state': 'Delhi',
                'city': 'New Delhi'
            },
            proposed_action='CREATE',
            validation_status='VALID'
        )
        url = reverse('hub:import_preview', args=[batch.batch_id])
        res = self.client.get(url)
        self.assertEqual(res.status_code, 200)
        self.assertContains(res, 'SOLO HOLDINGS PVT LTD')
        self.assertContains(res, '<div>-</div>')

    def test_gst_data_with_optional_contact_fields_excluded(self):
        """Case 7: GST data with optional contact fields (phone, email, PAN) excluded."""
        batch = self._create_gst_batch('BATCH-REG-GST-OPT-EXCL')
        StagedRow.objects.create(
            batch=batch,
            row_number=1,
            entity_identifier='07AAAAA0000A1Z5',
            cleaned_data={
                'gst_number': '07AAAAA0000A1Z5',
                'trade_name': 'KAPOOR TRADERS',
                'legal_name': 'KAPOOR TRADING CORP'
            },
            proposed_action='CREATE',
            validation_status='VALID'
        )
        url = reverse('hub:import_preview', args=[batch.batch_id])
        res = self.client.get(url)
        self.assertEqual(res.status_code, 200)
        self.assertContains(res, 'KAPOOR TRADERS')
        self.assertContains(res, '07AAAAA0000A1Z5')
        self.assertContains(res, 'null / blank')
        self.assertContains(res, 'MISSING')

    def test_fully_mapped_record_displays_properly(self):
        """Case 8: Fully mapped MCA record with phone, director, company email, and DIN."""
        batch = self._create_mca_batch('BATCH-REG-FULLY-MAPPED')
        StagedRow.objects.create(
            batch=batch,
            row_number=1,
            entity_identifier='U12345DL2022PTC123456',
            canonical_mobile='9876543210',
            cleaned_data={
                'name': 'ZENITH INFOTECH PVT LTD',
                'entityId': 'U12345DL2022PTC123456',
                'entityType': 'company',
                'directorName': 'Alice Smith',
                'directorEmail': 'alice@zenith.com',
                'email': 'hq@zenith.com',
                'original_phone': '+91 98765 43210',
                'canonical_mobile': '9876543210',
                'cleaned_phone': '9876543210',
                'phone_status': 'valid',
                'phone_validity': 'valid'
            },
            proposed_action='CREATE',
            validation_status='VALID'
        )
        url = reverse('hub:import_preview', args=[batch.batch_id])
        res = self.client.get(url)
        self.assertEqual(res.status_code, 200)
        self.assertContains(res, 'ZENITH INFOTECH PVT LTD')
        self.assertContains(res, 'Alice Smith')
        self.assertContains(res, 'alice@zenith.com')
        self.assertContains(res, 'hq@zenith.com')
        self.assertContains(res, '+91 98765 43210')
        self.assertContains(res, 'VALID (10 DIGIT)')

    def test_missing_phone_validity_and_canonical_mobile_in_batch_detail(self):
        """Case 9: Batch detail page also renders cleanly without canonical_mobile or phone_validity."""
        batch = self._create_mca_batch('BATCH-REG-DETAIL-SAFE')
        StagedRow.objects.create(
            batch=batch,
            row_number=1,
            entity_identifier='U11111DL2025PTC111111',
            cleaned_data={
                'name': 'MINIMAL ENTITY PVT LTD',
                'entityId': 'U11111DL2025PTC111111',
                'entityType': 'company'
            },
            proposed_action='CREATE',
            validation_status='VALID'
        )
        url = reverse('hub:import_batch_detail', args=[batch.batch_id])
        res = self.client.get(url)
        self.assertEqual(res.status_code, 200)
        self.assertContains(res, 'MINIMAL ENTITY PVT LTD')
        self.assertContains(res, 'MISSING')

    def test_excluded_fields_not_reintroduced_in_publishing_payload(self):
        """Case 10: Omitted fields do not overwrite existing DB values upon publication."""
        MCACompany.objects.create(
            entity_id='U77777MH2021PTC777777',
            name='ORIGINAL NAME PVT LTD',
            state='Maharashtra',
            city='Mumbai',
            email='existing_company@original.com'
        )
        batch = self._create_mca_batch('BATCH-REG-PUB-PRESERVE')
        StagedRow.objects.create(
            batch=batch,
            row_number=1,
            entity_identifier='U77777MH2021PTC777777',
            cleaned_data={
                'name': 'UPDATED NAME PVT LTD',
                'entityId': 'U77777MH2021PTC777777',
                'entityType': 'company'
            },
            proposed_action='UPDATE',
            validation_status='VALID'
        )
        execute_import_publication(batch)
        company = MCACompany.objects.get(entity_id='U77777MH2021PTC777777')
        self.assertEqual(company.name, 'UPDATED NAME PVT LTD')
        self.assertEqual(company.state, 'Maharashtra')
        self.assertEqual(company.city, 'Mumbai')
        self.assertEqual(company.email, 'existing_company@original.com')



class LiveProgressWorkflowTestCase(TestCase):
    """
    Regression test suite for live progress tracking, elapsed time,
    smoothed ETA calculation, database publish stats, and recovery.
    """

    def setUp(self):
        self.user = User.objects.create_user(
            username='progresstester',
            email='progress@example.com',
            password='testpassword123'
        )
        self.other_user = User.objects.create_user(
            username='otherprogressuser',
            email='other@example.com',
            password='testpassword123'
        )
        self.batch = ImportBatch.objects.create(
            batch_id='BATCH-PROGRESS-TEST-001',
            dataset_type='MCA',
            file_name='progress_test.xlsx',
            file_path='imports/progress_test.xlsx',
            file_checksum='1234567890abcdef',
            file_format='XLSX',
            status='PROCESSING',
            created_by=self.user,
            total_input_rows=1000
        )

    def test_init_progress_and_get_progress(self):
        from hub.progress import init_progress, get_progress
        init_progress(
            batch_id=self.batch.batch_id,
            dataset_type='MCA',
            stage='Validating and cleaning',
            total_rows=1000,
            unit='rows',
            step=3
        )
        prog = get_progress(self.batch.batch_id)
        self.assertEqual(prog['batch_id'], self.batch.batch_id)
        self.assertEqual(prog['stage'], 'Validating and cleaning')
        self.assertEqual(prog['total_rows'], 1000)
        self.assertEqual(prog['processed_rows'], 0)
        self.assertEqual(prog['percentage'], 0.0)
        self.assertEqual(prog['unit'], 'rows')
        self.assertEqual(prog['step'], 3)
        self.assertEqual(prog['status'], 'PROCESSING')
        self.assertFalse(prog['is_completed'])
        self.assertFalse(prog['is_failed'])

    def test_smoothed_throughput_and_eta_calculation(self):
        from hub.progress import init_progress, update_progress, get_progress
        import time
        init_progress(
            batch_id=self.batch.batch_id,
            dataset_type='MCA',
            stage='Validating and cleaning',
            total_rows=10000
        )
        # First update
        update_progress(
            batch_id=self.batch.batch_id,
            processed_rows=2000,
            total_rows=10000,
            stage='Validating and cleaning'
        )
        # Advance slightly and second update to form speed
        time.sleep(0.05)
        update_progress(
            batch_id=self.batch.batch_id,
            processed_rows=5000,
            total_rows=10000,
            stage='Validating and cleaning'
        )
        prog = get_progress(self.batch.batch_id)
        self.assertEqual(prog['processed_rows'], 5000)
        self.assertEqual(prog['percentage'], 50.0)
        self.assertIn('rows/sec', prog['speed_formatted'])
        self.assertTrue('approx.' in prog['eta_formatted'] or 'approximately' in prog['eta_formatted'] or 'Calculating' in prog['eta_formatted'])

    def test_monotonic_progress_enforcement(self):
        from hub.progress import init_progress, update_progress, get_progress
        init_progress(
            batch_id=self.batch.batch_id,
            dataset_type='MCA',
            stage='Validating',
            total_rows=1000
        )
        update_progress(self.batch.batch_id, processed_rows=500, total_rows=1000)
        # Attempt to regress processed_rows
        update_progress(self.batch.batch_id, processed_rows=300, total_rows=1000)
        prog = get_progress(self.batch.batch_id)
        self.assertEqual(prog['processed_rows'], 500)
        self.assertEqual(prog['percentage'], 50.0)

    def test_indeterminate_progress_when_total_unknown(self):
        from hub.progress import init_progress, update_progress, get_progress
        init_progress(
            batch_id=self.batch.batch_id,
            dataset_type='MCA',
            stage='Scanning stream',
            total_rows=None
        )
        update_progress(self.batch.batch_id, processed_rows=250, total_rows=None)
        prog = get_progress(self.batch.batch_id)
        self.assertIsNone(prog['total_rows'])
        self.assertIsNone(prog['percentage'])
        self.assertEqual(prog['processed_rows'], 250)

    def test_db_stats_tracking_in_publication_step(self):
        from hub.progress import init_progress, update_progress, get_progress
        init_progress(
            batch_id=self.batch.batch_id,
            dataset_type='MCA',
            stage='Committed batch 1 of 5',
            step=4,
            unit='source rows'
        )
        db_stats = {
            'companies_created': 120,
            'companies_updated': 30,
            'directors_created': 210,
            'issues_created': 4
        }
        update_progress(
            batch_id=self.batch.batch_id,
            processed_rows=150,
            total_rows=1000,
            step=4,
            db_stats=db_stats
        )
        prog = get_progress(self.batch.batch_id)
        self.assertEqual(prog['step'], 4)
        self.assertEqual(prog['db_stats']['companies_created'], 120)
        self.assertEqual(prog['db_stats']['companies_updated'], 30)
        self.assertEqual(prog['db_stats']['directors_created'], 210)
        self.assertEqual(prog['db_stats']['issues_created'], 4)

    def test_set_progress_completed(self):
        from hub.progress import init_progress, set_progress_completed, get_progress
        init_progress(
            batch_id=self.batch.batch_id,
            dataset_type='MCA',
            stage='Publishing'
        )
        set_progress_completed(self.batch.batch_id, message='Published successfully')
        prog = get_progress(self.batch.batch_id)
        self.assertTrue(prog['is_completed'])
        self.assertEqual(prog['percentage'], 100.0)
        self.assertEqual(prog['status'], 'COMPLETED')
        self.assertEqual(prog['eta_formatted'], '00:00')

    def test_set_progress_failed(self):
        from hub.progress import init_progress, set_progress_failed, get_progress
        init_progress(
            batch_id=self.batch.batch_id,
            dataset_type='MCA',
            stage='Publishing'
        )
        set_progress_failed(self.batch.batch_id, error_message='Disk I/O failure')
        prog = get_progress(self.batch.batch_id)
        self.assertTrue(prog['is_failed'])
        self.assertEqual(prog['status'], 'FAILED')
        self.assertEqual(prog['error_message'], 'Disk I/O failure')

    def test_progress_persistence_and_database_reconstruction(self):
        from hub.progress import init_progress, _PROGRESS_STORE, get_progress
        self.batch.status = 'COMPLETED'
        self.batch.total_input_rows = 500
        self.batch.imported_records = 500
        self.batch.proposed_new_records = 400
        self.batch.proposed_updates = 100
        self.batch.save()

        # Clear memory store to simulate reboot / cache miss
        _PROGRESS_STORE.clear()

        prog = get_progress(self.batch.batch_id)
        self.assertEqual(prog['batch_id'], self.batch.batch_id)
        self.assertEqual(prog['status'], 'COMPLETED')
        self.assertEqual(prog['percentage'], 100.0)
        self.assertEqual(prog['total_rows'], 500)
        self.assertEqual(prog['processed_rows'], 500)
        self.assertEqual(prog['db_stats']['companies_created'], 400)
        self.assertEqual(prog['db_stats']['companies_updated'], 100)

    def test_import_progress_api_unauthorized_access(self):
        from django.urls import reverse
        self.client.force_login(self.other_user)
        url = reverse('hub:import_progress_api', kwargs={'batch_id': self.batch.batch_id})
        response = self.client.get(url)
        self.assertEqual(response.status_code, 403)

    def test_import_progress_page_view_authorized(self):
        from django.urls import reverse
        self.client.force_login(self.user)
        url = reverse('hub:import_progress_page', kwargs={'batch_id': self.batch.batch_id})
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        self.assertTemplateUsed(response, 'import/progress.html')

    def test_healthcheck_endpoint(self):
        response = self.client.get('/health/')
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(data.get('status'), 'healthy')
        self.assertEqual(data.get('database'), 'ok')

        resp2 = self.client.get('/healthz')
        self.assertEqual(resp2.status_code, 200)
