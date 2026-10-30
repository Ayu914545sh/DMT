"""
Live Progress Tracking & Background Processing Engine for DMT Ingestion Pipeline.
Supports bounded-memory processing for large files (200MB - 1GB), real progress metrics,
elapsed time and accurate ETA calculation, multi-import concurrency, and page refresh persistence.
"""

import os
import time
import threading
import logging
from typing import Optional, Dict, Any, Callable
from django.conf import settings
from django.db import connection, close_old_connections, transaction
from django.urls import reverse
from django.utils import timezone

logger = logging.getLogger(__name__)

import uuid

# Thread-safe in-memory registry and distributed cache integration
_PROGRESS_LOCK = threading.Lock()
_PROGRESS_REGISTRY: Dict[str, Dict[str, Any]] = {}
_PROGRESS_STORE = _PROGRESS_REGISTRY
_CANCEL_SET = set()


def _progress_cache_key(batch_id: str) -> str:
    return f"import_progress:{batch_id}"


def _cancel_cache_key(batch_id: str) -> str:
    return f"import_cancel:{batch_id}"


def _cache_get(key: str) -> Optional[Any]:
    try:
        from django.core.cache import cache
        return cache.get(key)
    except Exception as e:
        logger.debug(f"Cache get failed for {key}: {e}")
        return None


def _cache_set(key: str, value: Any, timeout: int = 86400) -> None:
    try:
        from django.core.cache import cache
        cache.set(key, value, timeout=timeout)
    except Exception as e:
        logger.debug(f"Cache set failed for {key}: {e}")


def _cache_delete(key: str) -> None:
    try:
        from django.core.cache import cache
        cache.delete(key)
    except Exception as e:
        logger.debug(f"Cache delete failed for {key}: {e}")


def clear_progress_cache(batch_id: str) -> None:
    """Clears progress and cancel tokens from distributed cache and local process memory."""
    _cache_delete(_progress_cache_key(batch_id))
    _cache_delete(_cancel_cache_key(batch_id))
    with _PROGRESS_LOCK:
        _PROGRESS_REGISTRY.pop(batch_id, None)
        _CANCEL_SET.discard(batch_id)


def format_duration(seconds: Optional[float]) -> str:
    """Formats seconds into MM:SS or HH:MM:SS."""
    if seconds is None or seconds < 0:
        return ""
    total_secs = int(round(seconds))
    mins, secs = divmod(total_secs, 60)
    hours, mins = divmod(mins, 60)
    if hours > 0:
        return f"{hours:02d}:{mins:02d}:{secs:02d}"
    return f"{mins:02d}:{secs:02d}"


def count_file_rows_bounded(file_path: str, file_format: str = 'CSV', header_row: int = 0, sheet_name: Optional[str] = None) -> Optional[int]:
    """
    Counts total data rows with bounded memory usage (buffer <= 1MB).
    Does not load the file into memory. Returns None if line count cannot be reliably obtained.
    """
    if not os.path.exists(file_path):
        return None

    if file_format == 'CSV':
        total_lines = 0
        try:
            with open(file_path, 'rb') as f:
                # Read in 1MB chunks
                for chunk in iter(lambda: f.read(1024 * 1024), b''):
                    total_lines += chunk.count(b'\n')
            # Subtract header line(s)
            return max(0, total_lines - (header_row + 1))
        except Exception as e:
            logger.warning(f"Could not count CSV rows in bounded mode: {e}")
            return None

    elif file_format in ('XLSX', 'XLS'):
        ext = file_path.lower().split('.')[-1]
        if ext == 'xlsx':
            try:
                import openpyxl
                # read_only=True streams XML without loading entire workbook
                wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
                target_sheet = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active
                max_r = target_sheet.max_row
                wb.close()
                if max_r:
                    return max(0, max_r - (header_row + 1))
            except Exception as e:
                logger.warning(f"Could not read xlsx max_row in bounded mode: {e}")
        return None

    return None


def init_progress(
    batch_id: str,
    dataset_type: str = 'MCA',
    stage: str = 'Queued for processing',
    total_rows: Optional[int] = None,
    step: int = 3,
    unit: str = 'rows'
) -> Dict[str, Any]:
    """Initializes a new tracking record for a batch."""
    now = time.time()
    info = {
        'batch_id': batch_id,
        'dataset_type': dataset_type,
        'stage': stage,
        'status': 'QUEUED' if stage.lower().startswith('queued') else 'PROCESSING',
        'step': step,
        'unit': unit,
        'total_rows': total_rows,
        'processed_rows': 0,
        'percentage': 0.0 if total_rows and total_rows > 0 else None,
        'speed_rate': 0.0,
        'speed_formatted': '',
        'validation_failures': 0,
        'missing_invalid': 0,
        'quarantined': 0,
        'proposed_new': 0,
        'proposed_updates': 0,
        'unchanged': 0,
        'db_stats': {},
        'created_at': now,
        'active_start_time': None,
        'queue_time_seconds': 0.0,
        'elapsed_seconds': 0.0,
        'eta_seconds': None,
        'elapsed_formatted': '00:00',
        'eta_formatted': '',
        'duration_formatted': '',
        'message': f"{stage}...",
        'error_message': '',
        'start_time': now,
        'last_db_persist': now,
        'output_ready': False,
        'is_completed': False,
        'is_failed': False,
        'is_cancelled': False,
        'samples': [(now, 0)],
        'timestamp': now,
        'download_report_url': reverse('hub:download_cleaning_report', kwargs={'batch_id': batch_id}),
        'download_cleaned_url': reverse('hub:download_cleaned_file', kwargs={'batch_id': batch_id}),
        'preview_url': reverse('hub:import_preview', kwargs={'batch_id': batch_id}),
        'batch_detail_url': reverse('hub:import_batch_detail', kwargs={'batch_id': batch_id}),
    }
    with _PROGRESS_LOCK:
        _PROGRESS_REGISTRY[batch_id] = info
        _CANCEL_SET.discard(batch_id)
    _cache_delete(_cancel_cache_key(batch_id))
    _cache_set(_progress_cache_key(batch_id), info, timeout=86400)
    return info


def update_progress(
    batch_id: str,
    stage: Optional[str] = None,
    status: Optional[str] = None,
    step: Optional[int] = None,
    processed_rows: Optional[int] = None,
    total_rows: Optional[int] = None,
    unit: Optional[str] = None,
    validation_failures: Optional[int] = None,
    proposed_new: Optional[int] = None,
    proposed_updates: Optional[int] = None,
    unchanged: Optional[int] = None,
    quarantined: Optional[int] = None,
    missing_invalid: Optional[int] = None,
    db_stats: Optional[Dict[str, Any]] = None,
    error_message: Optional[str] = None,
    output_ready: Optional[bool] = None,
    force_db_save: bool = False
) -> Dict[str, Any]:
    """Updates progress state in memory with smoothed throughput, honest ETA, and throttles DB writes."""
    now = time.time()
    with _PROGRESS_LOCK:
        info = _PROGRESS_REGISTRY.get(batch_id)
        if not info:
            info = init_progress(
                batch_id,
                stage=stage or 'Validating and cleaning',
                total_rows=total_rows,
                step=step or 3,
                unit=unit or 'rows'
            )

        if stage is not None:
            info['stage'] = stage
        if status is not None:
            info['status'] = status
        if step is not None:
            info['step'] = step
        if unit is not None:
            info['unit'] = unit
        if total_rows is not None:
            info['total_rows'] = total_rows
        if processed_rows is not None:
            # Monotonicity guard: do not regress processed_rows
            cur_p = info.get('processed_rows', 0)
            if processed_rows >= cur_p or (total_rows and processed_rows == total_rows):
                info['processed_rows'] = processed_rows
        if validation_failures is not None:
            info['validation_failures'] = validation_failures
        if proposed_new is not None:
            info['proposed_new'] = proposed_new
        if proposed_updates is not None:
            info['proposed_updates'] = proposed_updates
        if unchanged is not None:
            info['unchanged'] = unchanged
        if quarantined is not None:
            info['quarantined'] = quarantined
        if missing_invalid is not None:
            info['missing_invalid'] = missing_invalid
        if db_stats is not None:
            cur_db = info.get('db_stats') or {}
            cur_db.update(db_stats)
            info['db_stats'] = cur_db
        if error_message is not None:
            info['error_message'] = error_message
        if output_ready is not None:
            info['output_ready'] = output_ready

        # Distinguish queue wait time from active processing time
        cur_status = info.get('status', 'QUEUED')
        cur_stage = info.get('stage', '')
        if cur_status in ('PROCESSING', 'IMPORTING', 'RUNNING') and 'Queue' not in cur_stage:
            if info.get('active_start_time') is None:
                info['active_start_time'] = now
                info['queue_time_seconds'] = max(0.0, now - info.get('created_at', now))

        active_start = info.get('active_start_time') or info.get('start_time', now)
        elapsed = max(0.0, now - active_start)
        info['elapsed_seconds'] = elapsed
        info['elapsed_formatted'] = format_duration(elapsed)

        tot = info.get('total_rows')
        proc = info.get('processed_rows', 0)
        cur_unit = info.get('unit', 'rows')

        # Smoothed throughput calculation using rolling sliding window (last 8-10 seconds)
        samples = info.setdefault('samples', [])
        samples.append((now, proc))
        # Keep samples within the last 10 seconds, maximum 20 points
        samples[:] = [s for s in samples if (now - s[0]) <= 10.0][-20:]

        smoothed_rate = 0.0
        if len(samples) >= 2:
            dt = samples[-1][0] - samples[0][0]
            dp = samples[-1][1] - samples[0][1]
            window_rate = (dp / dt) if dt > 0.01 else 0.0
            overall_rate = (proc / elapsed) if elapsed > 0.01 else 0.0
            if window_rate > 0.0 and overall_rate > 0.0:
                smoothed_rate = (0.7 * window_rate) + (0.3 * overall_rate)
            elif window_rate > 0.0:
                smoothed_rate = window_rate
            elif overall_rate > 0.0:
                smoothed_rate = overall_rate
        elif elapsed > 0.5 and proc > 0:
            smoothed_rate = proc / elapsed

        info['speed_rate'] = round(smoothed_rate, 1)
        if smoothed_rate >= 1.0:
            info['speed_formatted'] = f"{int(round(smoothed_rate)):,} {cur_unit}/sec"
        elif smoothed_rate > 0.0:
            info['speed_formatted'] = f"{smoothed_rate:.1f} {cur_unit}/sec"
        else:
            info['speed_formatted'] = ""

        # Honest percentage and estimated time remaining calculation
        if tot and tot > 0:
            pct = round(min(100.0, (proc / tot) * 100), 1)
            info['percentage'] = pct
            if pct >= 100.0 or info.get('is_completed'):
                info['eta_seconds'] = 0.0
                info['eta_formatted'] = '00:00'
            elif proc < 50 or elapsed < 1.5:
                info['eta_seconds'] = None
                info['eta_formatted'] = 'Calculating estimate...'
            elif smoothed_rate > 0.0:
                remaining = max(0, tot - proc)
                eta_val = remaining / smoothed_rate
                info['eta_seconds'] = round(eta_val, 1)
                info['eta_formatted'] = f"approximately {format_duration(eta_val)}"
            else:
                info['eta_seconds'] = None
                info['eta_formatted'] = 'Calculating estimate...'
        else:
            # Unknown total work: honest indeterminate representation
            info['percentage'] = None
            info['eta_seconds'] = None
            info['eta_formatted'] = ''

        # Dynamic human-readable status message conforming to specification:
        # e.g.: "Validating and cleaning — 45% (45,000 / 100,000 rows processed)"
        if info.get('is_failed'):
            info['message'] = f"Failed: {info.get('error_message') or 'An error occurred during processing.'}"
        elif info.get('is_completed'):
            dur = info.get('duration_formatted') or format_duration(elapsed)
            info['message'] = f"{cur_stage}: Completed in {dur} — {proc:,} {cur_unit} processed."
        elif info.get('percentage') is not None and tot:
            info['message'] = f"{cur_stage} — {int(info['percentage'])}% ({proc:,} of {tot:,} {cur_unit} processed)"
        elif proc > 0:
            info['message'] = f"{cur_stage} — {proc:,} {cur_unit} processed"
        else:
            info['message'] = f"{cur_stage}..."

        info['timestamp'] = now
        # Always mirror to distributed cache for multi-worker and multi-container access
        _cache_set(_progress_cache_key(batch_id), info, timeout=86400)

        # Throttle DB persistence: at most once every 1.5 seconds unless forced or terminal
        should_persist = force_db_save or (now - info.get('last_db_persist', 0) >= 1.5)

    if should_persist:
        _persist_progress_to_db(batch_id, info)
        with _PROGRESS_LOCK:
            info['last_db_persist'] = now

    return info


def _persist_progress_to_db(batch_id: str, info: Dict[str, Any]) -> None:
    """Saves progress snapshot into the ImportBatch record in persistent storage."""
    try:
        from hub.models import ImportBatch
        batch = ImportBatch.objects.filter(batch_id=batch_id).first()
        if not batch:
            return

        summary = batch.cleaning_summary or {}
        summary['progress'] = {
            'stage': info.get('stage'),
            'status': info.get('status'),
            'step': info.get('step', 3),
            'unit': info.get('unit', 'rows'),
            'total_rows': info.get('total_rows'),
            'processed_rows': info.get('processed_rows'),
            'percentage': info.get('percentage'),
            'speed_rate': info.get('speed_rate', 0.0),
            'speed_formatted': info.get('speed_formatted', ''),
            'validation_failures': info.get('validation_failures'),
            'missing_invalid': info.get('missing_invalid'),
            'quarantined': info.get('quarantined'),
            'proposed_new': info.get('proposed_new'),
            'proposed_updates': info.get('proposed_updates'),
            'unchanged': info.get('unchanged'),
            'db_stats': info.get('db_stats') or {},
            'elapsed_seconds': info.get('elapsed_seconds'),
            'elapsed_formatted': info.get('elapsed_formatted'),
            'duration_formatted': info.get('duration_formatted'),
            'eta_formatted': info.get('eta_formatted'),
            'message': info.get('message'),
            'output_ready': info.get('output_ready'),
            'is_completed': info.get('is_completed'),
            'is_failed': info.get('is_failed'),
            'is_cancelled': info.get('is_cancelled'),
            'error_message': info.get('error_message', ''),
            'updated_at': time.time(),
        }
        batch.cleaning_summary = summary

        # Keep batch.status in sync with active status if processing
        cur_status = info.get('status')
        update_fields = ['cleaning_summary']
        if cur_status and cur_status in ('QUEUED', 'PROCESSING', 'IMPORTING', 'FAILED', 'CANCELLED', 'COMPLETED', 'PREVIEWED'):
            batch.status = cur_status
            update_fields.append('status')

        if info.get('total_rows') and (not batch.total_input_rows or batch.total_input_rows == 0):
            batch.total_input_rows = info['total_rows']
            update_fields.append('total_input_rows')

        if info.get('error_message') and not batch.error_message:
            batch.error_message = info['error_message']
            update_fields.append('error_message')

        batch.save(update_fields=list(set(update_fields)))
    except Exception as e:
        logger.warning(f"Error persisting progress to DB for batch {batch_id}: {e}")


def get_progress(batch_id: str) -> Dict[str, Any]:
    """
    Returns the current progress state with multi-worker / multi-container synchronization.
    Checks process memory (L1), distributed cache (L2), and database snapshot (L3) with stage integrity.
    """
    with _PROGRESS_LOCK:
        local_info = dict(_PROGRESS_REGISTRY[batch_id]) if batch_id in _PROGRESS_REGISTRY else None

    cached_info = _cache_get(_progress_cache_key(batch_id))

    info: Optional[Dict[str, Any]] = None
    if local_info and cached_info:
        if cached_info.get('timestamp', 0) >= local_info.get('timestamp', 0):
            info = dict(cached_info)
            with _PROGRESS_LOCK:
                _PROGRESS_REGISTRY[batch_id] = dict(cached_info)
        else:
            info = dict(local_info)
            _cache_set(_progress_cache_key(batch_id), local_info, timeout=86400)
    elif cached_info:
        info = dict(cached_info)
        with _PROGRESS_LOCK:
            _PROGRESS_REGISTRY[batch_id] = dict(cached_info)
    elif local_info:
        info = dict(local_info)
        _cache_set(_progress_cache_key(batch_id), local_info, timeout=86400)

    try:
        from hub.models import ImportBatch
        batch = ImportBatch.objects.filter(batch_id=batch_id).first()
        if not batch:
            if info:
                return info
            return {
                'batch_id': batch_id,
                'stage': 'Unknown',
                'status': 'UNKNOWN',
                'step': 3,
                'unit': 'rows',
                'percentage': 0,
                'message': 'Job not found',
                'is_completed': False,
                'is_failed': True,
                'is_cancelled': False,
                'output_ready': False,
                'error_message': 'Import job record not found.'
            }

        status = batch.status
        is_completed = (status == 'COMPLETED' or status == 'PREVIEWED')
        is_failed = (status == 'FAILED')
        is_cancelled = (status == 'CANCELLED')

        # If batch reached terminal status in DB, DB state is authoritative
        if status in ('COMPLETED', 'PREVIEWED', 'FAILED', 'CANCELLED'):
            summary = batch.cleaning_summary or {}
            saved_progress = summary.get('progress') or {}

            if status == 'COMPLETED':
                stage = saved_progress.get('stage') or 'Database Publishing Complete'
                step = 4
                unit = 'source rows committed'
                pct = 100.0
                output_ready = True
            elif status == 'PREVIEWED':
                stage = saved_progress.get('stage') or 'Validation and Cleaning Complete'
                step = 3
                unit = 'rows'
                pct = 100.0
                output_ready = True
            elif status == 'FAILED':
                stage = saved_progress.get('stage') or 'Failed'
                step = saved_progress.get('step', 3)
                unit = saved_progress.get('unit', 'rows')
                pct = None
                output_ready = False
            else:  # CANCELLED
                stage = 'Cancelled'
                step = saved_progress.get('step', 3)
                unit = saved_progress.get('unit', 'rows')
                pct = None
                output_ready = False

            total = batch.total_input_rows
            processed = saved_progress.get('processed_rows', total if is_completed else 0)
            if is_completed and total:
                processed = total

            dur = saved_progress.get('duration_formatted') or saved_progress.get('elapsed_formatted', '')
            final_info = {
                'batch_id': batch.batch_id,
                'dataset_type': batch.dataset_type,
                'stage': saved_progress.get('stage', stage),
                'status': status,
                'step': saved_progress.get('step', step),
                'unit': saved_progress.get('unit', unit),
                'total_rows': total,
                'processed_rows': processed,
                'percentage': pct,
                'speed_rate': 0.0,
                'speed_formatted': '',
                'validation_failures': saved_progress.get('validation_failures', batch.missing_invalid_values + batch.quarantined_rows),
                'missing_invalid': batch.missing_invalid_values,
                'quarantined': batch.quarantined_rows,
                'proposed_new': batch.proposed_new_records,
                'proposed_updates': batch.proposed_updates,
                'unchanged': batch.unchanged_records,
                'db_stats': saved_progress.get('db_stats') or ({
                    'companies_created': batch.proposed_new_records or 0,
                    'companies_updated': batch.proposed_updates or 0,
                    'directors_created': 0,
                    'issues_created': 0
                } if (status == 'COMPLETED' and batch.dataset_type == 'MCA') else {
                    'gst_created': batch.proposed_new_records or 0,
                    'gst_updated': batch.proposed_updates or 0,
                    'issues_created': 0
                } if (status == 'COMPLETED' and batch.dataset_type == 'GST') else {
                    'companies_created': 0,
                    'companies_updated': 0,
                    'directors_created': 0,
                    'gst_created': 0,
                    'gst_updated': 0,
                    'issues_created': 0
                }),
                'elapsed_seconds': saved_progress.get('elapsed_seconds', 0.0),
                'eta_seconds': 0.0 if is_completed else None,
                'elapsed_formatted': saved_progress.get('elapsed_formatted', dur or '00:00'),
                'duration_formatted': dur,
                'eta_formatted': '00:00' if is_completed else '',
                'message': saved_progress.get('message', f"{stage}"),
                'error_message': batch.error_message or '',
                'output_ready': output_ready,
                'is_completed': is_completed,
                'is_failed': is_failed,
                'is_cancelled': is_cancelled,
                'timestamp': time.time(),
                'download_report_url': reverse('hub:download_cleaning_report', kwargs={'batch_id': batch.batch_id}),
                'download_cleaned_url': reverse('hub:download_cleaned_file', kwargs={'batch_id': batch.batch_id}),
                'preview_url': reverse('hub:import_preview', kwargs={'batch_id': batch.batch_id}),
                'batch_detail_url': reverse('hub:import_batch_detail', kwargs={'batch_id': batch.batch_id}),
            }
            with _PROGRESS_LOCK:
                _PROGRESS_REGISTRY[batch_id] = final_info
            _cache_set(_progress_cache_key(batch_id), final_info, timeout=86400)
            return final_info

        # If in-memory or cached progress exists, ensure stage integrity & watchdog
        if info:
            cur_stage = info.get('stage', '')
            if status in ('QUEUED', 'PROCESSING', 'STAGED') and ('Waiting for column mapping' in cur_stage or info.get('step', 3) < 3):
                use_celery = getattr(settings, 'USE_CELERY', False)
                info['stage'] = 'Queued in Celery worker' if use_celery else 'Queued for processing'
                info['step'] = 3
                if batch.total_input_rows and not info.get('total_rows'):
                    info['total_rows'] = batch.total_input_rows
                with _PROGRESS_LOCK:
                    _PROGRESS_REGISTRY[batch_id] = info
                _cache_set(_progress_cache_key(batch_id), info, timeout=86400)

            # Self-healing watchdog: only if truly stalled without worker activity and not handled by Celery
            staging_locked = bool(_cache_get(f"staging_lock_{batch_id}"))
            publish_locked = bool(_cache_get(f"publish_lock_{batch_id}"))

            if status in ('QUEUED', 'PROCESSING') and 'Queued' in cur_stage and info.get('processed_rows', 0) == 0:
                created_at = info.get('created_at') or info.get('start_time') or time.time()
                if not staging_locked and not getattr(settings, 'USE_CELERY', False) and (time.time() - created_at) >= 60.0 and not info.get('is_completed') and not info.get('is_failed'):
                    if not info.get('_fallback_spawned'):
                        info['_fallback_spawned'] = True
                        with _PROGRESS_LOCK:
                            _PROGRESS_REGISTRY[batch_id] = info
                        _cache_set(_progress_cache_key(batch_id), info, timeout=86400)
                        logger.warning(f"Batch {batch_id} queued without worker pickup; engaging self-healing fallback worker.")
                        update_progress(batch_id, stage='Validating and cleaning', status='PROCESSING', force_db_save=True)
                        start_background_staging_job(batch_id, strict_phone=True, force_thread=True)

            # Self-healing watchdog for Step 4 (Publication):
            if (status == 'IMPORTING' or 'Queued' in cur_stage) and info.get('step', 0) == 4 and info.get('processed_rows', 0) == 0:
                created_at = info.get('start_time') or info.get('created_at') or time.time()
                if not publish_locked and not getattr(settings, 'USE_CELERY', False) and (time.time() - created_at) >= 60.0 and not info.get('is_completed') and not info.get('is_failed'):
                    if not info.get('_pub_fallback_spawned'):
                        info['_pub_fallback_spawned'] = True
                        with _PROGRESS_LOCK:
                            _PROGRESS_REGISTRY[batch_id] = info
                        _cache_set(_progress_cache_key(batch_id), info, timeout=86400)
                        logger.warning(f"Batch {batch_id} publication queued without worker pickup; engaging self-healing publication thread.")
                        start_background_publication_job(batch_id, force_thread=True)

            return info

        # Fallback: reconstruct active/queued job faithfully from DB snapshot
        summary = batch.cleaning_summary or {}
        saved_progress = summary.get('progress') or {}

        if status == 'UPLOADED':
            stage = 'Waiting for column mapping'
            step = 2
            unit = 'rows'
            pct = 0.0
            output_ready = False
        elif status == 'MAPPED':
            stage = 'Queued for processing'
            step = 3
            unit = 'rows'
            pct = 0.0
            output_ready = False
        elif status == 'IMPORTING':
            stage = saved_progress.get('stage') or 'Importing to database'
            step = 4
            unit = 'source rows committed'
            pct = batch.progress_percentage
            output_ready = False
        else:
            stage = saved_progress.get('stage') or 'Validating and cleaning'
            step = 3
            unit = 'rows'
            pct = batch.progress_percentage
            output_ready = False

        total = batch.total_input_rows
        processed = saved_progress.get('processed_rows', 0)

        info = {
            'batch_id': batch.batch_id,
            'dataset_type': batch.dataset_type,
            'stage': saved_progress.get('stage', stage),
            'status': status,
            'step': saved_progress.get('step', step),
            'unit': saved_progress.get('unit', unit),
            'total_rows': total,
            'processed_rows': processed,
            'percentage': saved_progress.get('percentage', pct),
            'speed_rate': saved_progress.get('speed_rate', 0.0),
            'speed_formatted': saved_progress.get('speed_formatted', ''),
            'validation_failures': saved_progress.get('validation_failures', batch.missing_invalid_values + batch.quarantined_rows),
            'missing_invalid': batch.missing_invalid_values,
            'quarantined': batch.quarantined_rows,
            'proposed_new': batch.proposed_new_records,
            'proposed_updates': batch.proposed_updates,
            'unchanged': batch.unchanged_records,
            'db_stats': saved_progress.get('db_stats') or {
                    'companies_created': 0,
                    'companies_updated': 0,
                    'directors_created': 0,
                    'gst_created': 0,
                    'gst_updated': 0,
                    'issues_created': 0,
                },
            'elapsed_seconds': saved_progress.get('elapsed_seconds', 0.0),
            'eta_seconds': saved_progress.get('eta_seconds'),
            'elapsed_formatted': saved_progress.get('elapsed_formatted', '00:00'),
            'duration_formatted': saved_progress.get('duration_formatted', ''),
            'eta_formatted': saved_progress.get('eta_formatted', ''),
            'message': saved_progress.get('message', f"{stage}"),
            'error_message': batch.error_message or '',
            'output_ready': output_ready,
            'is_completed': False,
            'is_failed': False,
            'is_cancelled': False,
            'timestamp': time.time(),
            'download_report_url': reverse('hub:download_cleaning_report', kwargs={'batch_id': batch.batch_id}),
            'download_cleaned_url': reverse('hub:download_cleaned_file', kwargs={'batch_id': batch.batch_id}),
            'preview_url': reverse('hub:import_preview', kwargs={'batch_id': batch.batch_id}),
            'batch_detail_url': reverse('hub:import_batch_detail', kwargs={'batch_id': batch.batch_id}),
        }

        with _PROGRESS_LOCK:
            _PROGRESS_REGISTRY[batch_id] = info
        _cache_set(_progress_cache_key(batch_id), info, timeout=86400)
        return info

    except Exception as e:
        logger.error(f"Error reconstructing progress for batch {batch_id}: {e}", exc_info=True)
        return {
            'batch_id': batch_id,
            'stage': 'Error',
            'status': 'FAILED',
            'step': 3,
            'unit': 'rows',
            'percentage': None,
            'message': str(e),
            'is_completed': False,
            'is_failed': True,
            'is_cancelled': False,
            'output_ready': False,
            'error_message': str(e)
        }


def set_progress_completed(
    batch_id: str,
    summary: Optional[Dict[str, Any]] = None,
    final_status: str = 'PREVIEWED',
    step: Optional[int] = None,
    stage: Optional[str] = None,
    message: Optional[str] = None
) -> Dict[str, Any]:
    """Marks progress as 100% completed with outputs ready, stage duration, and final stats."""
    now = time.time()
    with _PROGRESS_LOCK:
        info = _PROGRESS_REGISTRY.get(batch_id)
        if not info:
            info = init_progress(batch_id)

        active_start = info.get('active_start_time') or info.get('start_time', now)
        elapsed = max(0.0, now - active_start)
        dur = format_duration(elapsed)
        info['elapsed_seconds'] = elapsed
        info['elapsed_formatted'] = dur
        info['duration_formatted'] = dur
        info['percentage'] = 100.0
        info['is_completed'] = True
        info['is_failed'] = False
        info['output_ready'] = True
        info['eta_seconds'] = 0.0
        info['eta_formatted'] = '00:00'
        if message:
            info['message'] = message
            info['stage'] = message
            if 'publish' in message.lower() or 'completed' in message.lower():
                final_status = 'COMPLETED'
        info['status'] = final_status

        if step is not None:
            info['step'] = step
        elif final_status == 'COMPLETED':
            info['step'] = 4
        else:
            info['step'] = 3

        if stage:
            info['stage'] = stage
        elif final_status == 'COMPLETED':
            info['stage'] = 'Database Publishing Complete'
            info['unit'] = 'source rows committed'
        else:
            info['stage'] = 'Validation and Cleaning Complete'
            info['unit'] = 'rows'

        if summary:
            if 'total_rows' in summary:
                info['total_rows'] = summary['total_rows']
                info['processed_rows'] = summary['total_rows']
            if 'proposed_new' in summary:
                info['proposed_new'] = summary['proposed_new']
            if 'proposed_updates' in summary:
                info['proposed_updates'] = summary['proposed_updates']
            if 'unchanged' in summary:
                info['unchanged'] = summary['unchanged']
            if 'quarantined' in summary:
                info['quarantined'] = summary['quarantined']
            if 'missing_invalid' in summary:
                info['missing_invalid'] = summary['missing_invalid']
                info['validation_failures'] = summary.get('missing_invalid', 0) + summary.get('quarantined', 0)
            if 'db_stats' in summary and summary['db_stats']:
                cur_db = info.get('db_stats') or {}
                cur_db.update(summary['db_stats'])
                info['db_stats'] = cur_db

        tot = info.get('total_rows') or info.get('processed_rows', 0)
        cur_unit = info.get('unit', 'rows')
        if final_status == 'COMPLETED':
            info['message'] = f"Published successfully in {dur} — {tot:,} {cur_unit} committed."
        else:
            info['message'] = f"Completed in {dur} — {tot:,} {cur_unit} processed."

        info['timestamp'] = now

    _persist_progress_to_db(batch_id, info)
    return info


def set_progress_failed(batch_id: str, error_message: str, step: Optional[int] = None) -> Dict[str, Any]:
    """Marks progress as failed with actionable error message and final duration."""
    now = time.time()
    with _PROGRESS_LOCK:
        info = _PROGRESS_REGISTRY.get(batch_id)
        if not info:
            info = init_progress(batch_id)

        active_start = info.get('active_start_time') or info.get('start_time', now)
        elapsed = max(0.0, now - active_start)
        info['elapsed_seconds'] = elapsed
        info['elapsed_formatted'] = format_duration(elapsed)
        info['stage'] = 'Failed'
        info['status'] = 'FAILED'
        info['is_completed'] = False
        info['is_failed'] = True
        info['error_message'] = error_message
        info['message'] = f"Failed: {error_message}"
        info['eta_formatted'] = '--'
        info['timestamp'] = now
        if step is not None:
            info['step'] = step

    _persist_progress_to_db(batch_id, info)
    return info


def request_cancel(batch_id: str) -> None:
    """Sets the cancellation flag across process memory, distributed cache, and DB."""
    with _PROGRESS_LOCK:
        _CANCEL_SET.add(batch_id)
        info = _PROGRESS_REGISTRY.get(batch_id)
        if info:
            info['stage'] = 'Cancelled'
            info['status'] = 'CANCELLED'
            info['is_cancelled'] = True
            info['message'] = "Import job cancelled by user."
            info['timestamp'] = time.time()
            _persist_progress_to_db(batch_id, info)

    _cache_set(_cancel_cache_key(batch_id), True, timeout=86400)
    cached = _cache_get(_progress_cache_key(batch_id))
    if cached:
        cached['stage'] = 'Cancelled'
        cached['status'] = 'CANCELLED'
        cached['is_cancelled'] = True
        cached['message'] = "Import job cancelled by user."
        cached['timestamp'] = time.time()
        _cache_set(_progress_cache_key(batch_id), cached, timeout=86400)

    try:
        from hub.models import ImportBatch
        ImportBatch.objects.filter(batch_id=batch_id).update(status='CANCELLED')
    except Exception as e:
        logger.warning(f"Could not persist CANCELLED status to DB for batch {batch_id}: {e}")


def is_cancelled(batch_id: str) -> bool:
    """Checks whether an import job was cancelled via process memory, distributed cache, or DB."""
    with _PROGRESS_LOCK:
        if batch_id in _CANCEL_SET:
            return True
    if _cache_get(_cancel_cache_key(batch_id)) is True:
        return True
    try:
        from hub.models import ImportBatch
        return ImportBatch.objects.filter(batch_id=batch_id, status='CANCELLED').exists()
    except Exception:
        return False


def _has_responsive_celery_worker() -> bool:
    """Checks whether at least one Celery worker is currently online and responsive."""
    try:
        from dmt_project.celery import app
        if not app:
            return False
        inspector = app.control.inspect(timeout=0.8)
        pings = inspector.ping()
        return bool(pings)
    except Exception as e:
        logger.debug(f"Celery inspect ping error: {e}")
        return False


def start_background_staging_job(batch_id: str, strict_phone: bool = True, force_thread: bool = False) -> Any:
    """
    Launches background staging job. Dispatches via Celery worker if available;
    otherwise falls back gracefully to a background daemon thread.
    Safely manages DB connections and updates real progress.
    """
    lock_key = f"staging_lock_{batch_id}"
    if not force_thread and _cache_get(lock_key):
        logger.warning(f"Staging already running for batch {batch_id}; skipping duplicate start.")
        return None

    use_celery = getattr(settings, 'USE_CELERY', False) and not force_thread
    if use_celery:
        try:
            if _has_responsive_celery_worker():
                from hub.tasks import stage_and_validate_batch_task
                update_progress(batch_id, stage='Queued in Celery worker', status='PROCESSING', step=3, unit='rows', force_db_save=True)
                async_res = stage_and_validate_batch_task.delay(batch_id, strict_phone=strict_phone)
                return async_res
            else:
                logger.warning(f"No responsive Celery workers online; activating background daemon staging thread for {batch_id}.")
        except Exception as e:
            logger.warning(f"Could not dispatch to Celery for staging batch {batch_id} ({e}); falling back to background thread.")

    def _worker():
        close_old_connections()
        lock_key = f"staging_lock_{batch_id}"
        _cache_set(lock_key, True, timeout=1800)
        try:
            from hub.models import ImportBatch
            from hub.importers import stage_and_validate_batch

            batch = ImportBatch.objects.filter(batch_id=batch_id).first()
            if not batch:
                set_progress_failed(batch_id, "Import batch record not found.", step=3)
                return

            if is_cancelled(batch_id):
                return

            # Stage 1: Queued for processing
            update_progress(batch_id, stage='Queued for processing', status='PROCESSING', step=3, unit='rows', force_db_save=True)
            time.sleep(0.1)

            # Stage 2: Validating and cleaning (handled inside stage_and_validate_batch via progress callback)
            def _progress_cb(processed: int, total: Optional[int], failures: int, interim_stats: Optional[Dict] = None, stage: Optional[str] = None):
                if is_cancelled(batch_id):
                    raise RuntimeError("Import job cancelled by user.")
                st = interim_stats or {}
                update_progress(
                    batch_id=batch_id,
                    stage=stage or 'Validating and cleaning',
                    status='PROCESSING',
                    step=3,
                    unit='rows',
                    processed_rows=processed,
                    total_rows=total,
                    validation_failures=failures,
                    proposed_new=st.get('proposed_new'),
                    proposed_updates=st.get('proposed_updates'),
                    unchanged=st.get('unchanged'),
                    quarantined=st.get('quarantined'),
                    missing_invalid=st.get('missing_invalid'),
                )

            summary = stage_and_validate_batch(batch, strict_phone=strict_phone, progress_callback=_progress_cb)

            if is_cancelled(batch_id):
                return

            # Stage 3: Preparing output
            update_progress(batch_id, stage='Preparing preview and standard outputs', status='PROCESSING', step=3, unit='rows', force_db_save=True)
            time.sleep(0.2)

            # Stage 4: Completed / Preview Ready with outputs ready
            set_progress_completed(batch_id, summary=summary, final_status='PREVIEWED', step=3)

        except Exception as e:
            logger.error(f"Error during background staging for batch {batch_id}: {e}", exc_info=True)
            if "cancelled" in str(e).lower():
                request_cancel(batch_id)
            else:
                set_progress_failed(batch_id, str(e), step=3)
        finally:
            _cache_delete(lock_key)
            close_old_connections()

    thread = threading.Thread(target=_worker, name=f"StagingWorker-{batch_id}", daemon=True)
    thread.start()
    import sys
    if 'test' in sys.argv:
        thread.join(timeout=30)
    return thread


def start_background_publication_job(batch_id: str, user_id: Optional[int] = None, force_thread: bool = False) -> Any:
    """
    Launches publication job. Dispatches via Celery worker if available;
    otherwise falls back gracefully to a background daemon thread.
    Durably initializes Step 4 progress upfront to prevent stale Step 3 state.
    """
    lock_key = f"publish_lock_{batch_id}"
    if not force_thread and _cache_get(lock_key):
        logger.warning(f"Publication already running for batch {batch_id}; skipping duplicate start.")
        return None

    from hub.models import ImportBatch, StagedRow
    batch = ImportBatch.objects.filter(batch_id=batch_id).first()
    total_source_rows = 1
    dataset_type = 'MCA'
    if batch:
        dataset_type = batch.dataset_type
        total_source_rows = batch.total_input_rows
        if not total_source_rows or total_source_rows == 0:
            total_source_rows = StagedRow.objects.filter(batch=batch).values('row_number').distinct().count() or 1
            batch.total_input_rows = total_source_rows
            batch.save(update_fields=['total_input_rows'])
        if batch.status != 'IMPORTING':
            batch.status = 'IMPORTING'
            batch.save(update_fields=['status'])

    use_celery = getattr(settings, 'USE_CELERY', False) and not force_thread
    stage_name = 'Queued in Celery worker' if use_celery else 'Preparing database publish'

    init_progress(
        batch_id,
        dataset_type=dataset_type,
        stage=stage_name,
        total_rows=total_source_rows,
        step=4,
        unit='source rows committed'
    )
    update_progress(
        batch_id,
        stage=stage_name,
        status='IMPORTING',
        step=4,
        unit='source rows committed',
        processed_rows=batch.imported_records or batch.checkpoint_row or 0,
        total_rows=total_source_rows,
        db_stats={'companies_created': 0, 'companies_updated': 0, 'directors_created': 0, 'gst_created': 0, 'gst_updated': 0, 'issues_created': 0},
        force_db_save=True
    )

    if use_celery:
        try:
            if _has_responsive_celery_worker():
                from hub.tasks import process_import_batch_task
                async_res = process_import_batch_task.delay(batch_id, user_id=user_id)
                return async_res
            else:
                logger.warning(f"No responsive Celery workers online; activating background daemon publication thread for {batch_id}.")
        except Exception as e:
            logger.warning(f"Could not dispatch to Celery for publishing batch {batch_id} ({e}); falling back to background thread.")

    def _worker():
        close_old_connections()
        lock_key = f"publish_lock_{batch_id}"
        _cache_set(lock_key, True, timeout=1800)
        try:
            from hub.models import ImportBatch
            from django.contrib.auth.models import User
            from hub.importers import execute_import_publication

            b = ImportBatch.objects.filter(batch_id=batch_id).first()
            if not b:
                set_progress_failed(batch_id, "Import batch record not found.", step=4)
                return

            user = User.objects.filter(id=user_id).first() if user_id else None

            def _progress_cb(processed: int, total: int, db_stats: Optional[Dict] = None, stage: Optional[str] = None):
                update_progress(
                    batch_id=batch_id,
                    stage=stage or 'Importing to database',
                    status='IMPORTING',
                    step=4,
                    unit='source rows committed',
                    processed_rows=processed,
                    total_rows=total,
                    db_stats=db_stats
                )

            res = execute_import_publication(b, user=user, progress_callback=_progress_cb)

            set_progress_completed(batch_id, summary={
                'total_rows': b.total_input_rows,
                'proposed_new': b.proposed_new_records,
                'proposed_updates': b.proposed_updates,
                'unchanged': b.unchanged_records,
                'quarantined': b.quarantined_rows,
                'missing_invalid': b.missing_invalid_values,
                'db_stats': res.get('db_stats') if isinstance(res, dict) else {},
            }, final_status='COMPLETED', step=4)

        except Exception as e:
            logger.error(f"Error during background publication for batch {batch_id}: {e}", exc_info=True)
            set_progress_failed(batch_id, str(e), step=4)
        finally:
            _cache_delete(lock_key)
            close_old_connections()

    thread = threading.Thread(target=_worker, name=f"PublishWorker-{batch_id}", daemon=True)
    thread.start()
    import sys
    if 'test' in sys.argv:
        thread.join(timeout=30)
    return thread
