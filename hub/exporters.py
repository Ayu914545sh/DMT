"""
Audited and formula-injection-safe data exporter for MCA and GST datasets.
Preserves identifiers and phone numbers as text, prevents Excel macro/formula execution,
and generates both CSV and XLSX output formats with bounded streaming memory.
"""

import os
import csv
import re
import logging

logger = logging.getLogger(__name__)
from typing import List, Dict, Any, Optional, Iterator, Tuple
import openpyxl
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font
from django.conf import settings
from django.utils import timezone

from .models import MCACompany, MCADirectorEntry, GSTRecord, ExportJob, AuditEvent, StagedRow
from .cleaners import sanitize_spreadsheet_value




def _valid_phone_or_blank(phone: Any, validity: Any) -> str:
    """Authoritative phone validator for export: returns 10-digit clean string or empty string."""
    if not phone:
        return ""
    digits = re.sub(r'\D', '', str(phone))
    if len(digits) == 10 and digits[0] in '6789' and not re.match(r'^(\d)\1{9}$', digits):
        if validity in ('valid', 'cleaned', 'multiple_numbers_review'):
            return digits
    return ""

def sanitize_text(val: Any) -> str:
    """Removes invalid XML/Excel characters and applies formula protection."""
    if val is None:
        return ""
    text = str(val)
    text = ILLEGAL_CHARACTERS_RE.sub('', text)
    return sanitize_spreadsheet_value(text)


def stream_rows_to_export(
    headers: List[str],
    row_generator: Iterator[List[Any]],
    file_path: str,
    file_format: str = 'CSV',
    sheet_title: str = 'Export',
    extra_sheets: Optional[List[Tuple[str, List[str], Iterator[List[Any]]]]] = None
) -> str:
    """
    Streams rows directly to disk in CSV or write-only XLSX format with constant memory.
    Automatically formats text cells as '@' to preserve leading zeros and prevent scientific notation.
    Supports splitting worksheets beyond 1,048,500 rows and adding auxiliary sheets (e.g. Phone Review).
    """
    from openpyxl.cell import WriteOnlyCell

    file_format = file_format.upper()
    if file_format == 'CSV':
        with open(file_path, 'w', newline='', encoding='utf-8') as f:
            writer = csv.writer(f)
            writer.writerow(headers)
            for row in row_generator:
                writer.writerow([sanitize_text(v) for v in row])
    else:
        wb = openpyxl.Workbook(write_only=True)
        sheet_part = 1
        ws = wb.create_sheet(title=sheet_title[:31])
        ws.append(headers)

        current_sheet_rows = 1
        max_sheet_rows = 1048500  # Safe threshold below Excel limit 1,048,576

        for row in row_generator:
            if current_sheet_rows >= max_sheet_rows:
                sheet_part += 1
                new_title = f"{sheet_title[:20]} Part {sheet_part}"[:31]
                ws = wb.create_sheet(title=new_title)
                ws.append(headers)
                current_sheet_rows = 1

            row_cells = []
            for v in row:
                c = WriteOnlyCell(ws, value=sanitize_text(v))
                c.number_format = '@'
                row_cells.append(c)
            ws.append(row_cells)
            current_sheet_rows += 1

        if extra_sheets:
            for ex_title, ex_headers, ex_gen in extra_sheets:
                ex_ws = wb.create_sheet(title=ex_title[:31])
                ex_ws.append(ex_headers)
                for ex_row in ex_gen:
                    row_cells = []
                    for v in ex_row:
                        c = WriteOnlyCell(ex_ws, value=sanitize_text(v))
                        c.number_format = '@'
                        row_cells.append(c)
                    ex_ws.append(row_cells)

        wb.save(file_path)
        wb.close()

    return file_path


def export_mca_dataset(queryset, columns: Optional[List[str]] = None, file_format: str = 'CSV', export_dir: Optional[str] = None) -> str:
    """
    Exports MCA Companies along with related director entries in streaming mode.
    """
    export_dir = export_dir or os.path.join(settings.MEDIA_ROOT, 'exports')
    os.makedirs(export_dir, exist_ok=True)
    
    timestamp_str = timezone.now().strftime('%Y%m%d_%H%M%S')
    filename = f"mca_export_{timestamp_str}.{file_format.lower()}"
    file_path = os.path.join(export_dir, filename)
    
    if not columns:
        columns = [
            'entity_id', 'entity_type', 'name', 'state', 'city', 'roc', 'date_of_incorporation',
            'paid_up_capital', 'email', 'director_name', 'canonical_mobile', 'contact_validity'
        ]
        
    export_chunk_size = getattr(settings, 'EXPORT_CHUNK_SIZE', 5000)
    qs = queryset.prefetch_related('directors').order_by('id')

    def _row_gen():
        for company in qs.iterator(chunk_size=export_chunk_size):
            directors = list(company.directors.all())
            row = []
            for col in columns:
                if col == 'director_name':
                    row.append("; ".join(d.director_name for d in directors if d.director_name))
                elif col in ('director_email',):
                    row.append("; ".join(d.director_email for d in directors if d.director_email))
                elif col in ('raw_mobile', 'original_phone'):
                    row.append("; ".join(d.raw_mobile for d in directors if d.raw_mobile))
                elif col in ('canonical_mobile', 'cleaned_phone'):
                    row.append("; ".join(d.canonical_mobile for d in directors if d.canonical_mobile and d.contact_validity in ('valid', 'cleaned', 'multiple_numbers_review')))
                elif col in ('contact_validity', 'phone_status'):
                    row.append("; ".join(d.contact_validity for d in directors if d.contact_validity))
                elif col == 'designation':
                    row.append("; ".join(d.designation for d in directors if d.designation))
                elif col == 'din':
                    row.append("; ".join(d.din for d in directors if d.din))
                else:
                    val = getattr(company, col, "")
                    row.append(str(val) if val is not None else "")
            yield row

    return stream_rows_to_export(columns, _row_gen(), file_path, file_format=file_format, sheet_title="MCA Export")


def export_gst_dataset(queryset, columns: Optional[List[str]] = None, file_format: str = 'CSV', export_dir: Optional[str] = None) -> str:
    """
    Exports GST Registrations with streaming memory and formula sanitization.
    """
    export_dir = export_dir or os.path.join(settings.MEDIA_ROOT, 'exports')
    os.makedirs(export_dir, exist_ok=True)
    
    timestamp_str = timezone.now().strftime('%Y%m%d_%H%M%S')
    filename = f"gst_export_{timestamp_str}.{file_format.lower()}"
    file_path = os.path.join(export_dir, filename)
    
    if not columns:
        columns = [
            'gst_number', 'legal_name', 'trade_name', 'status', 'source_type',
            'registration_date', 'pan', 'canonical_mobile', 'contact_validity',
            'email', 'state', 'city', 'pincode'
        ]
        
    export_chunk_size = getattr(settings, 'EXPORT_CHUNK_SIZE', 5000)
    qs = queryset.order_by('id')

    def _row_gen():
        for rec in qs.iterator(chunk_size=export_chunk_size):
            row = []
            for col in columns:
                if col in ('canonical_mobile', 'cleaned_phone'):
                    row.append(rec.canonical_mobile if (rec.canonical_mobile and rec.contact_validity in ('valid', 'cleaned', 'multiple_numbers_review')) else "")
                else:
                    row.append(getattr(rec, col, ""))
            yield row

    return stream_rows_to_export(columns, _row_gen(), file_path, file_format=file_format, sheet_title="GST Export")


def run_export_job(job: ExportJob) -> str:
    """
    Executes an ExportJob instance, creates the output file,
    records audit details, and updates the job record.
    """
    job.status = 'PROCESSING'
    job.save(update_fields=['status'])
    
    try:
        if job.dataset_type == 'MCA':
            qs = MCACompany.objects.filter(is_archived=False)
            filters = job.filters or {}
            if 'state' in filters and filters['state']:
                qs = qs.filter(state__iexact=filters['state'])
            if 'entity_type' in filters and filters['entity_type']:
                qs = qs.filter(entity_type=filters['entity_type'])
            if 'roc' in filters and filters['roc']:
                qs = qs.filter(roc__iexact=filters['roc'])
            if 'search' in filters and filters['search']:
                s = filters['search']
                qs = qs.filter(name__icontains=s) | qs.filter(entity_id__icontains=s)
                
            path = export_mca_dataset(qs, job.columns, job.file_format)
            job.row_count = qs.count()
        else:
            qs = GSTRecord.objects.filter(is_archived=False)
            filters = job.filters or {}
            if 'state' in filters and filters['state']:
                qs = qs.filter(state__iexact=filters['state'])
            if 'status' in filters and filters['status']:
                qs = qs.filter(status__iexact=filters['status'])
            if 'search' in filters and filters['search']:
                s = filters['search']
                qs = qs.filter(legal_name__icontains=s) | qs.filter(trade_name__icontains=s) | qs.filter(gst_number__icontains=s)
                
            path = export_gst_dataset(qs, job.columns, job.file_format)
            job.row_count = qs.count()
            
        job.file_path = path
        job.status = 'COMPLETED'
        job.completed_at = timezone.now()
        job.save()
        
        AuditEvent.objects.create(
            user=job.requester,
            action='EXPORT',
            dataset_type=job.dataset_type,
            model_name='MCACompany' if job.dataset_type == 'MCA' else 'GSTRecord',
            object_id=str(job.id),
            object_repr=f"Exported {job.dataset_type} ({job.row_count} records) as {job.file_format}",
            changes={'columns': job.columns, 'filters': job.filters, 'file_path': path}
        )
        
        return path
    except Exception as e:
        job.status = 'FAILED'
        job.save(update_fields=['status'])
        raise e


def export_cleaned_master_data(batch, file_format: str = 'CSV', export_dir: Optional[str] = None, columns: Optional[List[str]] = None) -> str:
    """
    Output A: Cleaned Master Data (Requirement 8A).
    Cleaned source records with company, director/contact, phone, and relationship references.
    Streams directly with bounded memory.
    Only exports columns that were selected during import (or explicitly chosen via `columns`).
    """
    export_dir = export_dir or os.path.join(settings.MEDIA_ROOT, 'exports')
    os.makedirs(export_dir, exist_ok=True)
    
    timestamp_str = timezone.now().strftime('%Y%m%d_%H%M%S')
    filename = f"master_data_{batch.batch_id}_{timestamp_str}.{file_format.lower()}"
    file_path = os.path.join(export_dir, filename)

    export_chunk_size = getattr(settings, 'EXPORT_CHUNK_SIZE', 5000)
    staged_rows = StagedRow.objects.filter(batch=batch).order_by('row_number')

    mapping_config = batch.mapping_config or {}
    selected_cols = mapping_config.get('selected_columns')
    columns_mapping = mapping_config.get('columns') or {}

    if batch.dataset_type == 'MCA':
        field_definitions = [
            ('row_number', 'Row Number', lambda s, c: s.row_number, None),
            ('entityId', 'Entity ID', lambda s, c: c.get('entityId', ''), 'entityId'),
            ('entityType', 'Entity Type', lambda s, c: c.get('entityType', 'company'), 'entityId'),
            ('name', 'Company Name', lambda s, c: c.get('company_name') or c.get('name', ''), 'name'),
            ('state', 'State', lambda s, c: c.get('state', ''), 'state'),
            ('city', 'City', lambda s, c: c.get('city', ''), 'city'),
            ('roc', 'ROC', lambda s, c: c.get('roc', ''), 'roc'),
            ('nicCode', 'Main Division', lambda s, c: c.get('nicCode', ''), 'nicCode'),
            ('nicLabel', 'Main Division Description', lambda s, c: c.get('nicLabel', ''), 'nicLabel'),
            ('classOfCompany', 'Class of Company', lambda s, c: c.get('classOfCompany', ''), 'classOfCompany'),
            ('dateOfIncorporation', 'Date of Incorporation', lambda s, c: c.get('clean_doi') or c.get('dateOfIncorporation', ''), 'dateOfIncorporation'),
            ('authorizedCapital', 'Authorized Capital', lambda s, c: c.get('clean_authorized_capital') or c.get('authorizedCapital', ''), 'authorizedCapital'),
            ('paidUpCapital', 'Paid Up Capital', lambda s, c: c.get('clean_paid_up_capital') or c.get('paidUpCapital', ''), 'paidUpCapital'),
            ('email', 'Company Email', lambda s, c: c.get('email', ''), 'email'),
            ('registeredAddress', 'Registered Address', lambda s, c: c.get('registeredAddress', ''), 'registeredAddress'),
            ('din', 'DIN', lambda s, c: c.get('din') or c.get('clean_din', ''), 'din'),
            ('directorName', 'Director Name', lambda s, c: c.get('directorName', ''), 'directorName'),
            ('designation', 'Designation', lambda s, c: ('Director' if c.get('directorName') else ''), 'directorName'),
            ('directorMobile', 'Normalized Phone', lambda s, c: (s.canonical_mobile if c.get('phone_validity') in ('valid', 'cleaned', 'multiple_numbers_review') else ''), 'directorMobile'),
            ('extension', 'Extension', lambda s, c: c.get('extension', ''), 'directorMobile'),
            ('secondary_phones', 'Secondary Phones', lambda s, c: (", ".join(c.get('secondary_phones', [])) if c.get('secondary_phones') else ""), 'directorMobile'),
            ('phone_validity', 'Phone Validity', lambda s, c: c.get('phone_validity', 'missing'), 'directorMobile'),
            ('phone_reason', 'Phone Validation Reason', lambda s, c: c.get('phone_reason') or c.get('error_reason') or '', 'directorMobile'),
            ('phone_ownership', 'Phone Ownership', lambda s, c: c.get('phone_ownership', 'director'), 'directorMobile'),
            ('directorEmail', 'Director Email', lambda s, c: c.get('directorEmail', ''), 'directorEmail'),
            ('source_file', 'Source File', lambda s, c: c.get('source_file', batch.file_name), 'provenance'),
            ('source_sheet', 'Source Sheet', lambda s, c: c.get('source_sheet', batch.sheet_name), 'provenance'),
            ('source_row', 'Source Row', lambda s, c: c.get('source_row', s.row_number), 'provenance')
        ]
    else:
        field_definitions = [
            ('row_number', 'Row Number', lambda s, c: s.row_number, None),
            ('gst_number', 'GSTIN', lambda s, c: c.get('gst_number', ''), 'gst_number'),
            ('legal_name', 'Legal Name', lambda s, c: c.get('legal_name', ''), 'legal_name'),
            ('trade_name', 'Trade Name', lambda s, c: c.get('trade_name', ''), 'trade_name'),
            ('status', 'Status', lambda s, c: c.get('status', ''), 'status'),
            ('type', 'Taxpayer Type', lambda s, c: c.get('source_type', ''), 'type'),
            ('registration_date', 'Registration Date', lambda s, c: c.get('clean_registration_date') or c.get('registration_date', ''), 'registration_date'),
            ('gst_type', 'GST Type', lambda s, c: c.get('gst_type', ''), 'gst_type'),
            ('pan', 'PAN', lambda s, c: c.get('pan', ''), 'pan'),
            ('mobile', 'Normalized Phone', lambda s, c: (s.canonical_mobile if c.get('phone_validity') in ('valid', 'cleaned', 'multiple_numbers_review') else ''), 'mobile'),
            ('extension', 'Extension', lambda s, c: c.get('extension', ''), 'mobile'),
            ('secondary_phones', 'Secondary Phones', lambda s, c: (", ".join(c.get('secondary_phones', [])) if c.get('secondary_phones') else ""), 'mobile'),
            ('phone_validity', 'Phone Validity', lambda s, c: c.get('phone_validity', 'missing'), 'mobile'),
            ('phone_reason', 'Phone Validation Reason', lambda s, c: c.get('phone_reason') or c.get('error_reason') or '', 'mobile'),
            ('phone_ownership', 'Phone Ownership', lambda s, c: c.get('phone_ownership', 'company'), 'mobile'),
            ('email', 'Email', lambda s, c: c.get('email', ''), 'email'),
            ('address', 'Address', lambda s, c: c.get('address', ''), 'address'),
            ('state', 'State', lambda s, c: c.get('state', ''), 'state'),
            ('city', 'City', lambda s, c: c.get('city', ''), 'city'),
            ('Pincode', 'Pincode', lambda s, c: c.get('clean_pincode') or c.get('Pincode', ''), 'Pincode'),
            ('source_file', 'Source File', lambda s, c: c.get('source_file', batch.file_name), 'provenance'),
            ('source_sheet', 'Source Sheet', lambda s, c: c.get('source_sheet', batch.sheet_name), 'provenance'),
            ('source_row', 'Source Row', lambda s, c: c.get('source_row', s.row_number), 'provenance')
        ]

    # Determine imported canonical fields from selected columns
    if selected_cols is not None:
        selected_set = set(selected_cols)
        imported_targets = {t for src, t in columns_mapping.items() if src in selected_set and t and t != 'skip'}
    else:
        imported_targets = {t for t in columns_mapping.values() if t and t != 'skip'}

    active_fields = []
    col_filters = set(columns) if columns else None

    for field_id, title, val_fn, req_target in field_definitions:
        if col_filters is not None:
            # User explicitly chose columns
            if field_id in col_filters or title in col_filters or req_target in col_filters:
                active_fields.append((title, val_fn))
        else:
            # Default: Omit fields that were excluded during import
            if req_target is None:
                active_fields.append((title, val_fn))
            elif req_target == 'provenance':
                active_fields.append((title, val_fn))
            elif req_target in imported_targets:
                active_fields.append((title, val_fn))

    if not active_fields:
        active_fields = [(title, val_fn) for _, title, val_fn, _ in field_definitions[:5]]

    headers = [t for t, _ in active_fields]

    def _row_gen():
        for s in staged_rows.iterator(chunk_size=export_chunk_size):
            c = s.cleaned_data
            yield [fn(s, c) for _, fn in active_fields]

    extra_sheets = None
    phone_target = 'directorMobile' if batch.dataset_type == 'MCA' else 'mobile'
    if file_format == 'XLSX' and phone_target in imported_targets:
        # Build Phone Review Sheet for invalid/uncertain phone values
        review_headers = [
            'Source Row', 'Entity Identifier', 'Company / Trade Name', 'Contact Name',
            'Original Raw Phone', 'Phone Validity', 'Rejection / Validation Reason', 'Rejected Parts'
        ]
        def _review_gen():
            # Uses indexed validation_status rather than re-scanning entire staged_rows table
            review_qs = StagedRow.objects.filter(
                batch=batch,
                validation_status__in=['WARNING', 'FLAGGED', 'ERROR']
            ).order_by('row_number')
            for s in review_qs.iterator(chunk_size=export_chunk_size):
                c = s.cleaned_data
                pstat = c.get('phone_validity')
                rej = c.get('rejected_phones') or []
                if pstat in ('invalid', 'multiple_numbers_review', 'FLAGGED', 'landline') or rej:
                    eid = c.get('entityId') or c.get('gst_number') or ""
                    cname = c.get('company_name') or c.get('name') or c.get('trade_name') or c.get('legal_name') or ""
                    dname = c.get('directorName') or ""
                    raw_p = c.get('original_phone') or ""
                    reason = c.get('phone_reason') or f"Phone format flagged ({pstat})"
                    rej_text = ", ".join(f"{r.get('number', '')} ({r.get('reason', '')})" for r in rej) if rej else ""
                    yield [s.row_number, eid, cname, dname, raw_p, pstat, reason, rej_text]

        extra_sheets = [('Phone Review', review_headers, _review_gen())]

    return stream_rows_to_export(headers, _row_gen(), file_path, file_format=file_format, sheet_title="Cleaned Master Data", extra_sheets=extra_sheets)


def _build_ready_to_call_row(phone: str, occurrences: List[Any], batch) -> List[Any]:
    """Helper to assemble a single deduplicated Ready-to-Call list row."""
    first_row_num, primary_c = occurrences[0]
    comp_name = primary_c.get('company_name') or primary_c.get('name') or primary_c.get('trade_name') or primary_c.get('legal_name') or ""
    entity_id = primary_c.get('entityId') or primary_c.get('gst_number') or ""
    contact_name = primary_c.get('directorName') or comp_name
    din = primary_c.get('din') or primary_c.get('clean_din') or ""
    designation = 'Director' if primary_c.get('directorName') else 'Authorized Signatory'
    contact_type = primary_c.get('phone_ownership', 'director').capitalize()
    city = primary_c.get('city') or ""
    state = primary_c.get('state') or ""
    extension = primary_c.get('extension') or ""

    alt_numbers = set()
    for _, occ in occurrences:
        for sec in occ.get('secondary_phones', []):
            if sec and sec != phone:
                alt_numbers.add(sec)

    is_shared = len(occurrences) > 1
    shared_str = "Yes" if is_shared else "No"
    
    associated_list = []
    source_refs = []
    for r_num, occ in occurrences:
        c_name = occ.get('company_name') or occ.get('trade_name') or occ.get('legal_name') or ""
        d_name = occ.get('directorName') or ""
        e_id = occ.get('entityId') or occ.get('gst_number') or ""
        desc = f"{c_name} ({e_id})" if c_name else e_id
        if d_name:
            desc += f" - Contact: {d_name}"
        if desc and desc not in associated_list:
            associated_list.append(desc)
        source_refs.append(f"Row {r_num}")

    assoc_str = "; ".join(associated_list) if is_shared else f"{comp_name} ({entity_id})"
    rel_notes = "Delimited primary mobile" if primary_c.get('phone_validity') == 'multiple_numbers_review' else "Direct contact endpoint"
    if is_shared:
        rel_notes += f"; Shared across {len(occurrences)} source records"

    return [
        phone,
        extension,
        primary_c.get('phone_validity', 'valid'),
        contact_type,
        contact_name,
        din,
        designation,
        comp_name,
        entity_id,
        city,
        state,
        shared_str,
        assoc_str,
        ", ".join(sorted(alt_numbers)) if alt_numbers else "",
        rel_notes,
        f"{batch.file_name} ({', '.join(source_refs)})"
    ]


def export_ready_to_call_list(batch, file_format: str = 'CSV', export_dir: Optional[str] = None) -> str:
    """
    Output B: Ready-to-Call List (Requirement 6 & 8B).
    Deduplicated by normalized 10-digit phone with O(1) memory via DB-ordered streaming.
    """
    export_dir = export_dir or os.path.join(settings.MEDIA_ROOT, 'exports')
    os.makedirs(export_dir, exist_ok=True)
    
    timestamp_str = timezone.now().strftime('%Y%m%d_%H%M%S')
    filename = f"ready_to_call_{batch.batch_id}_{timestamp_str}.{file_format.lower()}"
    file_path = os.path.join(export_dir, filename)

    headers = [
        'Normalized Phone', 'Extension', 'Phone Status', 'Contact Type',
        'Director / Contact Name', 'DIN', 'Designation',
        'Company / Legal / Trade Name', 'CIN / GSTIN', 'City', 'State',
        'Shared Number', 'Associated Entities / Contacts', 'Alternative Valid Numbers',
        'Relationship Notes', 'Source References'
    ]

    export_chunk_size = getattr(settings, 'EXPORT_CHUNK_SIZE', 5000)
    # Stream ordered by canonical_mobile with database-level eligibility filtering
    staged_qs = StagedRow.objects.filter(
        batch=batch
    ).exclude(
        canonical_mobile=''
    ).exclude(
        proposed_action='QUARANTINE'
    ).exclude(
        validation_status='ERROR'
    ).order_by('canonical_mobile', 'row_number')

    def _row_gen():
        current_phone = None
        occurrences = []

        for s in staged_qs.iterator(chunk_size=export_chunk_size):
            c = s.cleaned_data
            phone = s.canonical_mobile
            if not phone or c.get('phone_validity') not in ('valid', 'cleaned', 'multiple_numbers_review'):
                continue

            if current_phone is not None and phone != current_phone:
                yield _build_ready_to_call_row(current_phone, occurrences, batch)
                occurrences = []

            current_phone = phone
            occurrences.append((s.row_number, c))

        if current_phone is not None and occurrences:
            yield _build_ready_to_call_row(current_phone, occurrences, batch)

    return stream_rows_to_export(headers, _row_gen(), file_path, file_format=file_format, sheet_title="Ready-To-Call List")


def export_review_report(batch, file_format: str = 'CSV', export_dir: Optional[str] = None) -> str:
    """
    Output C: Review Report (Requirement 8C).
    Invalid phones, multiple numbers, ambiguous associations, and conflicting records.
    Streams directly with bounded memory.
    """
    export_dir = export_dir or os.path.join(settings.MEDIA_ROOT, 'exports')
    os.makedirs(export_dir, exist_ok=True)
    
    timestamp_str = timezone.now().strftime('%Y%m%d_%H%M%S')
    filename = f"review_report_{batch.batch_id}_{timestamp_str}.{file_format.lower()}"
    file_path = os.path.join(export_dir, filename)

    headers = [
        'Source Row', 'Entity Identifier', 'Company Name', 'Contact Name',
        'Field Name', 'Original Value', 'Validation Status', 'Issue Code',
        'Actionable Review Reason', 'Suggested Action', 'Source Reference'
    ]

    mapping_config = batch.mapping_config or {}
    selected_cols = mapping_config.get('selected_columns')
    columns_mapping = mapping_config.get('columns') or {}

    if selected_cols is not None:
        selected_set = set(selected_cols)
        imported_targets = {t for src, t in columns_mapping.items() if src in selected_set and t and t != 'skip'}
    else:
        imported_targets = {t for t in columns_mapping.values() if t and t != 'skip'}

    phone_selected = ('directorMobile' in imported_targets) if batch.dataset_type == 'MCA' else ('mobile' in imported_targets)

    export_chunk_size = getattr(settings, 'EXPORT_CHUNK_SIZE', 5000)
    staged_rows = StagedRow.objects.filter(batch=batch).order_by('row_number')

    def _row_gen():
        for s in staged_rows.iterator(chunk_size=export_chunk_size):
            c = s.cleaned_data
            entity_id = c.get('entityId') or c.get('gst_number') or ""
            comp_name = c.get('company_name') or c.get('name') or c.get('trade_name') or ""
            contact_name = c.get('directorName') or ""
            source_ref = f"{batch.file_name} (Row {s.row_number})"

            if phone_selected:
                phone_stat = c.get('phone_validity')
                rej_phones = c.get('rejected_phones') or []
                if phone_stat in ('invalid', 'multiple_numbers_review', 'FLAGGED', 'landline') or rej_phones:
                    orig_phone = c.get('original_phone', '')
                    issue_code = 'MULTIPLE_NUMBERS' if phone_stat == 'multiple_numbers_review' else ('FLAGGED_PHONE' if phone_stat == 'FLAGGED' else 'INVALID_PHONE_FORMAT')
                    reason = c.get('phone_reason') or f"Phone format flagged ({phone_stat}): {orig_phone}"
                    suggestion = "Verify against official ROC filing or split secondary numbers manually" if phone_stat == 'multiple_numbers_review' else "Check digit count or area code"
                    yield [
                        s.row_number, entity_id, comp_name, contact_name,
                        'Mobile / Phone', orig_phone, s.validation_status, issue_code,
                        reason, suggestion, source_ref
                    ]

            for msg in s.validation_messages:
                if "Mobile flagged" not in msg:
                    yield [
                        s.row_number, entity_id, comp_name, contact_name,
                        'General Field', "", s.validation_status, 'VALIDATION_WARNING',
                        msg, "Review source record", source_ref
                    ]

    return stream_rows_to_export(headers, _row_gen(), file_path, file_format=file_format, sheet_title="Review Report")


def export_audit_output(batch, file_format: str = 'CSV', export_dir: Optional[str] = None) -> str:
    """
    Output D: Audit Output (Requirement 8D).
    Original and cleaned values with the transformations applied.
    Streams directly with bounded memory.
    """
    export_dir = export_dir or os.path.join(settings.MEDIA_ROOT, 'exports')
    os.makedirs(export_dir, exist_ok=True)
    
    timestamp_str = timezone.now().strftime('%Y%m%d_%H%M%S')
    filename = f"audit_output_{batch.batch_id}_{timestamp_str}.{file_format.lower()}"
    file_path = os.path.join(export_dir, filename)

    headers = [
        'Source Row', 'Entity Identifier', 'Field Name',
        'Original (Raw) Value', 'Cleaned Value', 'Transformation Applied', 'Status'
    ]

    export_chunk_size = getattr(settings, 'EXPORT_CHUNK_SIZE', 5000)
    staged_rows = StagedRow.objects.filter(batch=batch).order_by('row_number')

    def _row_gen():
        for s in staged_rows.iterator(chunk_size=export_chunk_size):
            raw = s.raw_data
            c = s.cleaned_data
            entity_id = c.get('entityId') or c.get('gst_number') or ""

            orig_phone = c.get('original_phone', '')
            clean_phone = _valid_phone_or_blank(c.get('canonical_mobile'), c.get('phone_validity'))
            if orig_phone or clean_phone:
                if orig_phone == clean_phone:
                    trans = "Pristine: No transformation required"
                elif clean_phone:
                    trans = "Normalized to 10-digit Indian Mobile; stripped country code / formatting characters"
                else:
                    trans = f"Flagged as {c.get('phone_validity')}; original preserved"
                yield [
                    s.row_number, entity_id, 'Mobile / Phone', orig_phone, clean_phone, trans, c.get('phone_validity', 'missing')
                ]

            orig_name = raw.get('Company Name') or raw.get('legal_name') or raw.get('trade_name') or ""
            clean_comp = c.get('company_name', '')
            if orig_name or clean_comp:
                trans = "Normalized whitespace and quotes" if orig_name != clean_comp else "Preserved exact value"
                yield [
                    s.row_number, entity_id, 'Company Name', orig_name, clean_comp, trans, 'Cleaned'
                ]

            if 'Director First Name' in raw or 'Director Second Name' in raw:
                first = raw.get('Director First Name', '')
                sec = raw.get('Director Second Name', '')
                raw_dir = f"{first} {sec}".strip()
                clean_dir = c.get('directorName', '')
                trans = "Combined First Name and Second Name; filtered dot markers"
                yield [
                    s.row_number, entity_id, 'Director Name', raw_dir, clean_dir, trans, 'Cleaned'
                ]

            orig_date = raw.get('Date of Incorporation') or raw.get('registration_date') or ""
            clean_date_val = c.get('clean_doi') or c.get('clean_registration_date') or ""
            if orig_date or clean_date_val:
                trans = "Standardized to ISO format (YYYY-MM-DD)" if orig_date != clean_date_val else "Preserved"
                yield [
                    s.row_number, entity_id, 'Registration/Inc Date', orig_date, clean_date_val, trans, 'Cleaned'
                ]

    return stream_rows_to_export(headers, _row_gen(), file_path, file_format=file_format, sheet_title="Audit Transformations")


def export_processing_summary(batch) -> Dict[str, Any]:
    """
    Output E: Processing Summary (Requirement 8E).
    Returns comprehensive count and multiplicity explanation dictionary.
    """
    summary = batch.cleaning_summary or {}
    total_rows = summary.get('total_rows', batch.total_input_rows)
    companies = summary.get('unique_companies', 0)
    contacts = summary.get('unique_contacts', 0)
    valid_phones = summary.get('unique_valid_phones', 0)
    shared_phones = summary.get('shared_phones', 0)
    
    return {
        'batch_id': batch.batch_id,
        'dataset_type': batch.dataset_type,
        'file_name': batch.file_name,
        'total_source_rows': total_rows,
        'unique_companies': companies,
        'unique_contacts_directors': contacts,
        'valid_unique_phones': valid_phones,
        'shared_phones': shared_phones,
        'missing_invalid_values': summary.get('missing_invalid_values', batch.missing_invalid_values),
        'quarantined_rows': summary.get('quarantined_rows', batch.quarantined_rows),
        'conflict_count': summary.get('conflict_count', batch.conflict_count),
        'multiplicity_explanation': summary.get(
            'multiplicity_explanation',
            "Each source row represents a company-director relationship. Companies with multiple directors have repeated company rows preserving distinct contact endpoints."
        )
    }


def invalidate_stale_exports_for_batch(batch_id: str, export_dir: Optional[str] = None) -> int:
    """
    Invalidates stale cached/generated export files for a batch when its
    underlying cleaning rules, column mappings, or staged rows change.
    """
    export_dir = export_dir or os.path.join(settings.MEDIA_ROOT, 'exports')
    if not os.path.exists(export_dir):
        return 0
    removed_count = 0
    try:
        for fname in os.listdir(export_dir):
            if batch_id in fname:
                fpath = os.path.join(export_dir, fname)
                if os.path.isfile(fpath):
                    try:
                        os.remove(fpath)
                        removed_count += 1
                    except Exception as e:
                        logger.warning(f"Failed to remove stale export file {fpath}: {e}")
    except Exception as e:
        logger.warning(f"Error scanning export dir for stale batch files {batch_id}: {e}")
    return removed_count
